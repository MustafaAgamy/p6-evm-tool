"""Reusable Word (.docx) template furniture for the Baseline Narrative Report.

This module is a *self-contained* toolbox of low-level ``python-docx`` builders that
lay down the professional "shell" of the deliverable — A4 page geometry, a double
page-border frame, a repeating logo + title header, a centred page-number footer, a
cover page, a live Table-of-Contents field, base paragraph/heading styles, and the
shared numbered-heading / styled-table / date helpers.

It deliberately owns NO document-content logic (that lives in the section renderers);
everything here is furniture a caller assembles into a finished report. Each function
takes an explicit ``document`` / ``section`` so the module holds no global state and is
safe to call in any order.

Contract (SLICE B):
    apply_page_geometry(section)          apply_base_styles(document)
    add_page_border(section)              format_number(parts) -> str
    add_header(document, meta)            heading(document, number_label, title, level=1)
    add_footer(section)                   subheading(document, number_label, text)
    add_cover(document, meta)             styled_table(document, headers, rows, ...)
    add_toc(document)                     full_date(v) -> str
"""
import base64
import io
from datetime import datetime

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

# ── palette ───────────────────────────────────────────────────────────────────
NAVY = RGBColor(0x1F, 0x4E, 0x79)   # heading + table-header navy
INK = RGBColor(0x1A, 0x1D, 0x21)    # body ink
GREY = RGBColor(0x8A, 0x90, 0x99)   # muted / footer grey
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

_NAVY_FILL = '1F4E79'   # navy table-header background (hex, no '#')
_ZEBRA_FILL = 'F2F6F8'  # alternating body-row background
_BORDER_CLR = '000000'  # page-frame colour (black double rule)

_BODY_FONT = 'Calibri'
_HEAD_FONT = 'Calibri Light'
_BODY_PT = 11
_TABLE_PT = 9.5


# ── low-level helpers ─────────────────────────────────────────────────────────
def _img_bytes(data_url):
    """Decode a 'data:image/...;base64,XXXX' URL (or bare base64) to a BytesIO, or None."""
    if not data_url:
        return None
    try:
        b64 = data_url.split(',', 1)[1] if ',' in str(data_url) else data_url
        return io.BytesIO(base64.b64decode(b64))
    except Exception:
        return None


def _set_run_font(run, name, size=None, bold=None, italic=None, color=None):
    """Set a run's font, forcing the OOXML rFonts ascii/hAnsi so themed fonts like
    'Calibri Light' actually apply in Word (the plain .name setter is not always
    honoured for the +body/+headings theme fonts)."""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _set_cell_bg(cell, hex_fill):
    """Shade a table cell (header + zebra striping)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def _page_field_run(paragraph):
    """Append a live Word PAGE field (begin + instrText ' PAGE ' + separate + '1' + end)."""
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t'); placeholder.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, sep, placeholder, end):
        r.append(el)
    return run


# ── page geometry ─────────────────────────────────────────────────────────────
def apply_page_geometry(section):
    """A4 portrait with the report's specified asymmetric margins and 1cm header/footer
    distances. Values per SLICE B contract."""
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.15)
    section.left_margin = Cm(1.34)
    section.right_margin = Cm(1.06)
    section.bottom_margin = Cm(1.69)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    return section


# ── page border ───────────────────────────────────────────────────────────────
def add_page_border(section):
    """Draw a double page-level frame on EVERY page of the section (w:pgBorders).

    Idempotent: any prior ``w:pgBorders`` is removed first, and the new element is
    inserted immediately after ``w:pgMar`` (its correct schema position)."""
    sectPr = section._sectPr
    for existing in sectPr.findall(qn('w:pgBorders')):
        sectPr.remove(existing)
    borders = OxmlElement('w:pgBorders')
    borders.set(qn('w:offsetFrom'), 'page')
    borders.set(qn('w:display'), 'allPages')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'double')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '24')
        el.set(qn('w:color'), _BORDER_CLR)
        borders.append(el)
    pgMar = sectPr.find(qn('w:pgMar'))
    if pgMar is not None:
        pgMar.addnext(borders)
    else:
        sectPr.append(borders)
    return borders


# ── header ────────────────────────────────────────────────────────────────────
def add_header(document, meta):
    """Repeating page header: a 1x3 table of the three party logos
    (owner / consultant / contractor) FOLLOWED BY a centred title band
    (project name + the quoted document title). Both are always present.

    Missing logos leave their cell empty rather than dropping the column."""
    meta = meta or {}
    section = document.sections[0]
    header = section.header
    header.is_linked_to_previous = False

    logos = meta.get('logos') or {}
    table = header.add_table(rows=1, cols=3, width=Inches(6.9))
    table.autofit = True
    for i, key in enumerate(('owner', 'consultant', 'contractor')):
        cell = table.rows[0].cells[i]
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img = _img_bytes(logos.get(key))
        if img:
            try:
                para.add_run().add_picture(img, width=Inches(1.5))
            except Exception:
                pass  # missing / bad logo → empty cell, column kept

    band = header.add_paragraph()
    band.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name = str(meta.get('project_name') or 'Project')
    _set_run_font(band.add_run(name), _HEAD_FONT, size=10.5, bold=True, color=NAVY)
    doc_title = meta.get('document_title') or meta.get('doc_title')
    if doc_title:
        _set_run_font(band.add_run('   —   '), _BODY_FONT, size=9, color=GREY)
        _set_run_font(band.add_run('"%s"' % doc_title), _BODY_FONT, size=9,
                      italic=True, color=GREY)
    return header


# ── footer ────────────────────────────────────────────────────────────────────
def add_footer(section):
    """Centred footer holding ONLY a live PAGE field (no NUMPAGES), Calibri 9pt grey."""
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.text = ''
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_field_run(para)
    for run in para.runs:
        _set_run_font(run, _BODY_FONT, size=9, color=GREY)
    return footer


# ── cover ─────────────────────────────────────────────────────────────────────
def add_cover(document, meta):
    """Title page: project name (large Calibri Light), a location line, the quoted
    document title and a 'REV. <revision>' line, then a page break. Blank/None fields
    are skipped so the cover degrades gracefully."""
    meta = meta or {}
    for _ in range(5):
        document.add_paragraph()

    name_p = document.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(name_p.add_run(str(meta.get('project_name') or 'Project')),
                  _HEAD_FONT, size=30, bold=True, color=INK)

    location = meta.get('location')
    if location:
        loc_p = document.add_paragraph()
        loc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(loc_p.add_run(str(location)), _BODY_FONT, size=12, color=GREY)

    doc_title = meta.get('document_title') or meta.get('doc_title')
    if doc_title:
        dt_p = document.add_paragraph()
        dt_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(dt_p.add_run('"%s"' % doc_title), _HEAD_FONT, size=16, color=NAVY)

    revision = meta.get('revision')
    if revision not in (None, ''):
        rev_p = document.add_paragraph()
        rev_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(rev_p.add_run('REV. %s' % revision), _BODY_FONT, size=12,
                      bold=True, color=INK)

    document.add_page_break()


# ── table of contents ─────────────────────────────────────────────────────────
def add_toc(document):
    """A 'Table of Contents' heading plus a real, updatable Word TOC field, then a
    page break. The field shows a placeholder until the reader updates it in Word."""
    heading(document, '', 'Table of Contents', level=1)

    para = document.add_paragraph()
    run = para.add_run()
    r = run._r
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t'); placeholder.text = 'Right-click to update field.'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, sep, placeholder, end):
        r.append(el)

    document.add_page_break()
    return para


# ── base styles ───────────────────────────────────────────────────────────────
def apply_base_styles(document):
    """Normal = Calibri 11pt ink; Heading 1/2/3 = Calibri Light bold navy at
    16 / 13 / 11.5pt (rFonts ascii/hAnsi forced to 'Calibri Light')."""
    normal = document.styles['Normal']
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(_BODY_PT)
    normal.font.color.rgb = INK

    for level, size in (('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)):
        try:
            st = document.styles[level]
        except KeyError:
            continue
        st.font.name = _HEAD_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        # Force the theme font onto the style's rPr rFonts (ascii + hAnsi).
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), _HEAD_FONT)
        rFonts.set(qn('w:hAnsi'), _HEAD_FONT)
    return document


# ── numbering ─────────────────────────────────────────────────────────────────
def format_number(parts):
    """Format a heading number tuple: numeric levels joined by '.', a trailing alpha
    label kept as-is, always suffixed with ')'.

        (4,)            -> '4)'
        (3, 1)          -> '3.1)'
        (4, 2, 2, 'A')  -> '4.2.2.A)'
    """
    pieces = [str(p) for p in (parts or []) if p is not None and str(p) != '']
    return '.'.join(pieces) + ')'


# ── headings ──────────────────────────────────────────────────────────────────
def heading(document, number_label, title, level=1):
    """A Heading-styled paragraph carrying '<number_label> <title>' in Calibri Light
    bold navy. ``level`` maps to Heading 1/2/3 (sizes 16/13/11.5pt)."""
    size = {1: 16, 2: 13, 3: 11.5}.get(level, 16)
    para = document.add_heading('', level=level)
    text = ('%s %s' % (number_label or '', title or '')).strip()
    _set_run_font(para.add_run(text), _HEAD_FONT, size=size, bold=True, color=NAVY)
    return para


def subheading(document, number_label, text):
    """A second-level Heading-styled paragraph carrying '<number_label> <text>' in
    Calibri Light bold navy."""
    return heading(document, number_label, text, level=2)


# ── tables ────────────────────────────────────────────────────────────────────
def styled_table(document, headers, rows, widths=None, bold_last_row=False):
    """A styled, editable Word table: navy header row (white bold), thin Table-Grid
    borders and zebra shading on even body rows. Returns the table."""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    for i, head in enumerate(headers):
        run = hdr_cells[i].paragraphs[0].add_run('' if head is None else str(head))
        _set_run_font(run, _BODY_FONT, size=_TABLE_PT, bold=True, color=WHITE)
        _set_cell_bg(hdr_cells[i], _NAVY_FILL)

    last = len(rows) - 1
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row):
            if ci >= len(cells):
                break
            run = cells[ci].paragraphs[0].add_run('' if val is None else str(val))
            _set_run_font(run, _BODY_FONT, size=_TABLE_PT,
                          bold=bool(bold_last_row and ri == last))
        if ri % 2 == 0 and not (bold_last_row and ri == last):
            for c in cells:
                _set_cell_bg(c, _ZEBRA_FILL)

    if widths:
        table.autofit = False
        for row in table.rows:
            for i, width in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = width
    return table


# ── dates ─────────────────────────────────────────────────────────────────────
def full_date(v):
    """Render a date as 'DD Month YYYY' (e.g. '29 May 2026'); unparseable/blank input
    is passed through unchanged (blank -> '')."""
    if v in (None, ''):
        return ''
    if isinstance(v, datetime):
        return '%d %s %d' % (v.day, v.strftime('%B'), v.year)
    try:
        dt = datetime.strptime(str(v)[:10], '%Y-%m-%d')
        return '%d %s %d' % (dt.day, dt.strftime('%B'), dt.year)
    except (ValueError, TypeError):
        return str(v)
