"""SLICE D — Word Section-5 (Project Calendars & Holidays) renderer.

Draws the LIVE *P6 Calendar Audit* feature's five pieces into a python-docx
``Document`` from the SLICE A calendar payload (the widened Calendar-Audit
passthrough carried on the narrative's calendar section), EXCLUDING the Baseline
Start / Baseline Finish dates (those are cover meta, not Section-5 content and were
already dropped from the dashboard block by the builder):

    5.1)  Dashboard tiles      — a compact borderless label/value grid
    5.2)  Monthly histogram    — a stacked working(green)/non-working(red) per-month
                                 bar chart (net working days above each bar); when
                                 Chrome is unavailable falls back to a monthly stats
                                 table (Month | Working | Non-working | Holidays | Hours)
    5.3)  National holidays     — Date | Weekday | Reason
    5.4)  Working hours         — Profile | Hours | Hours/Day | Days/Week | Note
    5.5)  Comparison & usage    — Calendar | Hrs/Day | Days/Week | Non-working ahead |
                                 Default?  followed by  Calendar | Activities | % | Role

The module consumes B's template furniture (``docx_template``) and C's chart
rasteriser (``docx_charts.render_svg_png``); it re-uses the shared numbered
sub-heading + styled-table helpers so the section matches the rest of the report.
Every block is guarded on the presence of its payload key, so a partial calendar
payload renders whatever it has and never raises.
"""
import html as _h
import math
import re

from docx.shared import Pt

from p6_narrative import docx_charts, docx_template


# ── local helpers ─────────────────────────────────────────────────────────────
def _esc(x):
    return _h.escape('' if x is None else str(x))


def _num(x):
    return '' if x is None else str(x)


def _days_per_week(sub):
    """Pull the numeric days/week out of an hours-profile ``sub`` string
    (e.g. '5 days/week' -> '5'); pass the string through if no number is found."""
    m = re.search(r'(\d+(?:\.\d+)?)', str(sub or ''))
    return m.group(1) if m else _num(sub)


# Dashboard tiles — the LIVE Calendar Audit labels, EXCLUDING Baseline Start/Finish.
# (label, dashboard-key, format)  format: None | 'date' | 'hrs'
_DASH_TILES = [
    ('Total Calendar Days', 'total_calendar_days', None),
    ('Working Days', 'total_working_days', None),
    ('Non-Working Days', 'total_nonworking_days', None),
    ('Holidays', 'total_holidays', None),
    ('Exceptions', 'total_exceptions', None),
    ('Shutdown Periods', 'shutdown_periods', None),
    ('Avg Working Days / Month', 'avg_working_days_per_month', None),
    ('Avg Working Hours / Day', 'avg_working_hours_per_day', 'hrs'),
    ('Data Date', 'data_date', 'date'),
    ('Forecast Finish', 'project_finish', 'date'),
]


# ── 5.1) dashboard tiles ───────────────────────────────────────────────────────
def _dashboard_tiles(document, dash):
    pairs = []
    for label, key, fmt in _DASH_TILES:
        v = dash.get(key)
        if v is None or v == '':
            continue
        if fmt == 'date':
            v = docx_template.full_date(v)
        elif fmt == 'hrs':
            v = '%s hrs' % v
        pairs.append((label, str(v)))
    if not pairs:
        return None

    per_row = 2                                   # two label/value tiles per row
    nrows = math.ceil(len(pairs) / per_row)
    table = document.add_table(rows=nrows, cols=per_row * 2)
    table.autofit = True
    for idx, (label, val) in enumerate(pairs):
        r, base = divmod(idx, per_row)
        c = base * 2
        lab_cell = table.rows[r].cells[c]
        val_cell = table.rows[r].cells[c + 1]
        lr = lab_cell.paragraphs[0].add_run(label)
        docx_template._set_run_font(lr, 'Calibri', size=8.5, color=docx_template.GREY)
        vr = val_cell.paragraphs[0].add_run(val)
        docx_template._set_run_font(vr, 'Calibri', size=10.5, bold=True,
                                    color=docx_template.NAVY)
    return table


# ── 5.2) monthly histogram (SVG → PNG) with table fallback ─────────────────────
def _monthly_total(m):
    """Total days in the month row (calendar days), from its per-day list when
    present, else working + (holidays/exceptions) as a floor."""
    days = m.get('days') or []
    if days:
        return len(days)
    wd = m.get('working_days', 0) or 0
    return wd + max(m.get('holidays', 0) or 0, m.get('exceptions', 0) or 0)


def _monthly_svg(months):
    """Self-contained stacked working(green)/non-working(red) per-month bar SVG,
    with the NET working-days number above each bar. Returns (svg, width, height)."""
    n = len(months)
    bar_w, gap, left, top, plot_h = 34, 16, 44, 30, 200
    base_y = top + plot_h
    W = left + n * (bar_w + gap) + 24
    H = base_y + 48
    totals = [_monthly_total(m) for m in months]
    maxt = max(totals) or 1

    parts = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d" '
        'viewBox="0 0 %d %d" font-family="Segoe UI,Arial,sans-serif">' % (W, H, W, H),
        '<rect x="0" y="0" width="%d" height="%d" fill="#ffffff"/>' % (W, H),
        '<line x1="%d" y1="%d" x2="%d" y2="%d" stroke="#cbd5e1" stroke-width="1"/>'
        % (left - 6, base_y, W - 12, base_y),
    ]
    for i, m in enumerate(months):
        total = totals[i]
        wd = m.get('working_days', 0) or 0
        nwd = max(total - wd, 0)
        x = left + i * (bar_w + gap)
        h_w = plot_h * (wd / maxt) if maxt else 0
        h_n = plot_h * (nwd / maxt) if maxt else 0
        y_work = base_y - h_w
        y_non = y_work - h_n
        parts.append('<rect x="%d" y="%.1f" width="%d" height="%.1f" fill="#22c55e"/>'
                     % (x, y_work, bar_w, h_w))
        parts.append('<rect x="%d" y="%.1f" width="%d" height="%.1f" fill="#ef4444"/>'
                     % (x, y_non, bar_w, h_n))
        parts.append('<text x="%.1f" y="%.1f" text-anchor="middle" font-size="11" '
                     'font-weight="700" fill="#1a1d21">%s</text>'
                     % (x + bar_w / 2, y_non - 6, wd))
        parts.append('<text x="%.1f" y="%.1f" text-anchor="middle" font-size="9" '
                     'fill="#5b6470">%s</text>'
                     % (x + bar_w / 2, base_y + 16, _esc(m.get('label') or '')))
    ly = H - 12
    parts.append('<rect x="%d" y="%d" width="11" height="11" fill="#22c55e"/>'
                 '<text x="%d" y="%d" font-size="9.5" fill="#1a1d21">Working days</text>'
                 % (left, ly - 9, left + 16, ly))
    parts.append('<rect x="%d" y="%d" width="11" height="11" fill="#ef4444"/>'
                 '<text x="%d" y="%d" font-size="9.5" fill="#1a1d21">Non-working days</text>'
                 % (left + 120, ly - 9, left + 136, ly))
    parts.append('</svg>')
    return ''.join(parts), W, H


def _monthly_table(document, months):
    rows = []
    for m in months:
        total = _monthly_total(m)
        wd = m.get('working_days', 0) or 0
        nwd = max(total - wd, 0)
        rows.append([m.get('label'), wd, nwd,
                     m.get('holidays', 0), m.get('working_hours', 0)])
    return docx_template.styled_table(
        document, ['Month', 'Working', 'Non-working', 'Holidays', 'Hours'], rows)


def _monthly(document, months, chrome):
    png = None
    try:
        svg, w, h = _monthly_svg(months)
        if svg and chrome:
            png = docx_charts.render_svg_png(svg, w, h, chrome)
    except Exception:
        png = None
    if png:
        from docx.shared import Inches
        # scale to the text column width, preserving the SVG aspect ratio
        target = Inches(6.2)
        try:
            document.add_picture(docx_charts.stream(png), width=target)
        except Exception:
            _monthly_table(document, months)
    else:
        _monthly_table(document, months)


# ── 5.3) national holidays ─────────────────────────────────────────────────────
def _holidays(document, rows):
    hr = [[docx_template.full_date(h.get('date')) or h.get('display'),
           h.get('weekday'), h.get('reason') or ''] for h in rows]
    return docx_template.styled_table(document, ['Date', 'Weekday', 'Reason'], hr)


# ── 5.4) working hours ─────────────────────────────────────────────────────────
def _hours(document, profiles):
    rows = [[p.get('name'), p.get('hours'), p.get('hours_per_day'),
             _days_per_week(p.get('sub')), p.get('sub') or ''] for p in profiles]
    return docx_template.styled_table(
        document, ['Profile', 'Hours', 'Hours/Day', 'Days/Week', 'Note'], rows)


# ── 5.5) comparison + usage ────────────────────────────────────────────────────
def _comparison(document, rows):
    cr = [[r.get('name'), r.get('hours_per_day'), r.get('days_per_week'),
           r.get('exceptions'), 'Yes' if r.get('is_default') else 'No'] for r in rows]
    return docx_template.styled_table(
        document, ['Calendar', 'Hours/Day', 'Days/Week', 'Non-working ahead',
                   'Default?'], cr)


def _usage(document, rows):
    ur = [[r.get('name'), r.get('activities'), '%s%%' % _num(r.get('pct')),
           r.get('role')] for r in rows]
    return docx_template.styled_table(
        document, ['Calendar', 'Activities', '%', 'Role'], ur)


# ── entry point ────────────────────────────────────────────────────────────────
def render_calendar(document, payload, chrome, number=5):
    """Render Section-``number`` (Project Calendars & Holidays) into ``document`` from
    the SLICE A calendar ``payload``. ``chrome`` is a Chrome path for chart rasterising
    (None → the monthly histogram falls back to a stats table). The five sub-blocks are
    numbered ``number.1 … number.5`` so they follow the ACTUAL parent section number
    (the caller places the parent heading)."""
    p = payload or {}

    # x.1) Dashboard tiles
    dash = p.get('dashboard')
    if dash:
        docx_template.subheading(document, docx_template.format_number((number, 1)),
                                 'Calendar dashboard')
        _dashboard_tiles(document, dash)

    # x.2) Monthly histogram (or fallback table)
    monthly = p.get('monthly')
    if monthly:
        docx_template.subheading(document, docx_template.format_number((number, 2)),
                                 'Working / non-working days by month')
        _monthly(document, monthly, chrome)

    # x.3) National holidays
    holiday_dates = p.get('holiday_dates')
    if holiday_dates:
        docx_template.subheading(document, docx_template.format_number((number, 3)),
                                 'National holidays')
        _holidays(document, holiday_dates)

    # x.4) Working hours
    hours_profiles = p.get('hours_profiles')
    if hours_profiles:
        docx_template.subheading(document, docx_template.format_number((number, 4)),
                                 'Working hours')
        _hours(document, hours_profiles)

    # x.5) Comparison + usage
    comparison = p.get('comparison')
    usage = p.get('usage')
    if comparison or usage:
        docx_template.subheading(document, docx_template.format_number((number, 5)),
                                 'Calendar comparison & usage')
        if comparison:
            _comparison(document, comparison)
        if usage:
            _usage(document, usage)
    return document
