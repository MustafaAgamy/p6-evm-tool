"""Out-of-Sequence — Resolve & Correct.

The planner accepts (and may edit) the proposed correction on an out-of-sequence
finding; this module

  1. **re-validates** — applies the accepted corrections to an in-memory copy of the
     network and re-runs the SAME detection engine (the source of truth). A finding is
     *resolved* only when the out-of-sequence condition genuinely no longer holds; and

  2. **exports a corrected schedule** — writes the accepted relationship changes back to
     a copy of the imported file, in the same format (P6 XML or XER). Only relationship
     Type / Lag / existence is touched — actuals, % complete and dates are never changed.
     P6 reschedules on F9; the tool never invents a date.

Detection is untouched. This layer only *applies* what the planner accepted and lets the
engine re-check it.
"""
import types

from p6_evm.parser import parse_file, lag_day_hours
from p6_audit.graph import ScheduleGraph
from p6_audit.modules.out_of_sequence import run_out_of_sequence

_REL_TYPES = {'FS', 'SS', 'FF', 'SF'}


# ── Accepted-op normalisation ────────────────────────────────────────────────

def _norm_op(op):
    """A client-accepted correction → a stable internal shape.

    action ∈ {'change', 'remove', 'replace', 'add', 'data'}. 'data' (a wrong actual
    date) is never applied — no relationship edit clears it.
    """
    action = (op.get('action') or '').lower()
    lag = op.get('new_lag_days')
    return {
        'finding_id': op.get('finding_id') or '',
        'action': action,
        'pred_id': op.get('pred_id') or '',
        'succ_id': op.get('succ_id') or '',
        'new_type': ((op.get('new_type') or '').upper() or None),
        'new_lag_days': (float(lag) if lag is not None else None),
        'new_pred_id': op.get('new_pred_id') or '',
        'reason': op.get('reason') or '',
    }


def _oid_to_code(data):
    return {oid: (a.get('id') or '') for oid, a in data.activities.items()}


def _code_to_oids(data):
    m = {}
    for oid, a in data.activities.items():
        m.setdefault(a.get('id') or '', []).append(oid)
    return m


# ── Re-validation (in-memory, re-runs the detection engine) ──────────────────

def apply_ops_to_relationships(data, accepted):
    """Return a NEW relationships list with the accepted corrections applied by activity
    CODE — to every duplicate copy of a code pair, matching the file writer so re-validation
    and the exported file always agree. 'data' corrections are no-ops (they stay Open)."""
    oid_to_code = _oid_to_code(data)
    code_to_oids = _code_to_oids(data)
    rels = [dict(r) for r in data.relationships]

    def _match(r, pc, sc):
        return oid_to_code.get(r['pred_id']) == pc and oid_to_code.get(r['succ_id']) == sc

    for raw in accepted:
        o = _norm_op(raw)
        act, pc, sc = o['action'], o['pred_id'], o['succ_id']
        if act in ('data', 'manual') or not sc:
            continue
        if act in ('remove', 'replace'):
            rels = [r for r in rels if not _match(r, pc, sc)]
        if act == 'change':
            for r in rels:
                if _match(r, pc, sc):
                    if o['new_type'] in _REL_TYPES:
                        r['type'] = o['new_type']
                    if o['new_lag_days'] is not None:
                        r['lag_days'] = o['new_lag_days']
        if act in ('add', 'replace'):
            # an 'add' names its predecessor as pred_id (as the file writer reads it); a
            # 'replace' names the NEW predecessor as new_pred_id
            npc = o['new_pred_id'] or (pc if act == 'add' else '')
            if not npc:
                continue   # a replace with no new predecessor = the remove above; nothing to re-add
            p_oids, s_oids = code_to_oids.get(npc), code_to_oids.get(sc)
            if p_oids and s_oids:
                rels.append({
                    'pred_id': p_oids[0], 'succ_id': s_oids[0],
                    'type': o['new_type'] if o['new_type'] in _REL_TYPES else 'FS',
                    'lag_days': o['new_lag_days'] if o['new_lag_days'] is not None else 0.0,
                })
    return rels


def _shim(data, rels):
    return types.SimpleNamespace(activities=data.activities, relationships=rels,
                                 calendars=getattr(data, 'calendars', {}) or {},
                                 project=getattr(data, 'project', None))


_FIX_ACTIONS = ('change', 'remove', 'replace')
_MAX_ROUNDS = 200          # one more out-of-sequence tie of an activity is found per round


def complete_ops(data, config, accepted):
    """The accepted corrections, completed so the corrected schedule is clean in P6 (comment 88:
    after importing the corrected file, P6's schedule log still listed out-of-sequence activities
    and showed activities with no predecessor).

    1. EVERY out-of-sequence tie of an accepted activity is corrected. A finding shows the driving
       tie; an activity that started ahead of twenty unfinished predecessors has twenty such ties,
       and P6 lists it until all are corrected. Each further tie gets the engine's own
       recommendation - the same rule the planner accepted for the first one.
    2. The network stays closed. An activity that had a predecessor (successor) before the
       corrections and has none after them is tied to the project commencement (completion)
       milestone, so no correction leaves an open end.

    Returns the full op list: the accepted ops first, the added ones marked ``auto``."""
    acc = [dict(o) for o in accepted]
    norm = [_norm_op(o) for o in acc]
    targets = {o['succ_id'] for o in norm if o['action'] in _FIX_ACTIONS and o['succ_id']}
    seen = {(o['pred_id'], o['succ_id']) for o in norm}
    # a tie the planner chose (a replacement predecessor) is the planner's decision - never re-corrected
    seen |= {(o['new_pred_id'], o['succ_id']) for o in norm if o['new_pred_id']}
    if not targets:
        return acc
    # 1) the corrections the table SHOWS for every driving tie of an accepted activity
    first = run_out_of_sequence(ScheduleGraph(_shim(data, apply_ops_to_relationships(data, acc))), config)
    for f in first['findings']:
        if f.get('activity_id') not in targets:
            continue
        for x in f.get('pred_fixes') or []:
            pair = (x.get('pred_id') or '', f['activity_id'])
            if x.get('action') not in _FIX_ACTIONS or pair in seen:
                continue
            seen.add(pair)
            acc.append({'finding_id': '', 'pred_id': pair[0], 'succ_id': pair[1], 'action': x['action'],
                        'new_type': x.get('new_type'), 'new_lag_days': x.get('new_lag_days'),
                        'new_pred_id': x.get('new_pred_id') or '', 'auto': 'tie'})
    # 2) anything still out of sequence after them (a tie that only shows once another changed)
    for _round in range(_MAX_ROUNDS):
        fresh = run_out_of_sequence(ScheduleGraph(_shim(data, apply_ops_to_relationships(data, acc))), config)
        new = []
        for f in fresh['findings']:
            res = f.get('pred_resolution') or {}
            pair = (f.get('pred_id') or '', f.get('activity_id') or '')
            if f.get('activity_id') not in targets or res.get('action') not in _FIX_ACTIONS or pair in seen:
                continue
            seen.add(pair)
            new.append({'finding_id': '', 'pred_id': pair[0], 'succ_id': pair[1], 'action': res['action'],
                        'new_type': res.get('new_type'), 'new_lag_days': res.get('new_lag_days'),
                        'new_pred_id': (res.get('new_pred_id') or '') if res['action'] == 'replace' else '',
                        'auto': 'tie'})
        if not new:
            break
        acc += new
    return acc + _closing_ops(data, acc)


def _closing_ops(data, acc):
    """'add' ops that tie an activity the corrections left without any predecessor (successor)
    to the project commencement (completion) milestone."""
    from p6_audit.modules.out_of_sequence import _find_commencement, _find_completion
    acts = data.activities

    def ends(rels):
        ok = [r for r in rels if r['pred_id'] in acts and r['succ_id'] in acts]
        return {r['succ_id'] for r in ok}, {r['pred_id'] for r in ok}
    had_pred, had_succ = ends(data.relationships)
    now_pred, now_succ = ends(apply_ops_to_relationships(data, acc))
    graph = ScheduleGraph(_shim(data, data.relationships))
    start, finish = _find_commencement(graph), _find_completion(graph)
    out = []
    for oid in sorted(had_pred - now_pred, key=str):
        code = acts[oid].get('id') or ''
        if start and start[0] and start[0] != code and code:
            out.append({'finding_id': '', 'pred_id': start[0], 'succ_id': code, 'action': 'add',
                        'new_type': 'FS', 'new_lag_days': 0.0, 'new_pred_id': '', 'auto': 'close'})
    for oid in sorted(had_succ - now_succ, key=str):
        code = acts[oid].get('id') or ''
        if finish and finish[0] and finish[0] != code and code:
            out.append({'finding_id': '', 'pred_id': code, 'succ_id': finish[0], 'action': 'add',
                        'new_type': 'FS', 'new_lag_days': 0.0, 'new_pred_id': '', 'auto': 'close'})
    return out


def _rel_text(t, lag):
    lag = round(float(lag or 0), 1)
    lag = int(lag) if float(lag).is_integer() else lag
    return f"{t}({lag})" if lag else t


def changed_ties(data, full):
    """{activity code: [every relationship tie of that activity the corrections changed]} - both
    ends of each tie (comment 88 follow-up: highlight the relationship changed after the fix, on the
    predecessor or the successor side). Each entry: side ('pred' = a predecessor tie of the
    activity, 'succ' = a successor tie), other (the related activity's ID), before, after, auto."""
    code = _oid_to_code(data)
    name = {a.get('id') or '': a.get('name') or '' for a in data.activities.values()}
    base = {}
    for r in data.relationships:
        base.setdefault((code.get(r['pred_id']), code.get(r['succ_id'])), (r.get('type', 'FS'), r.get('lag_days', 0.0)))
    out = {}
    for raw in full:
        o = _norm_op(raw)
        act, pc, sc = o['action'], o['pred_id'], o['succ_id']
        if act not in ('change', 'remove', 'replace', 'add') or not sc:
            continue
        cur = base.get((pc, sc))
        before = _rel_text(*cur) if cur else ''
        if act == 'change':
            after = _rel_text(o['new_type'] or (cur or ('FS', 0))[0],
                              o['new_lag_days'] if o['new_lag_days'] is not None else (cur or ('FS', 0))[1])
            ties = [(pc, sc, before, after)]
        elif act == 'remove':
            ties = [(pc, sc, before, 'Removed')]
        elif act == 'add':
            ties = [(pc, sc, '', 'Added ' + _rel_text(o['new_type'] or 'FS', o['new_lag_days'] or 0))]
        else:   # replace = the offending tie removed + a new predecessor
            ties = [(pc, sc, before, 'Removed')]
            if o['new_pred_id']:
                ties.append((o['new_pred_id'], sc, '', 'Added ' + _rel_text(o['new_type'] or 'FS', o['new_lag_days'] or 0)))
        for p, s, b, a in ties:
            auto = raw.get('auto') or ''
            out.setdefault(s, []).append({'side': 'pred', 'other': p, 'other_name': name.get(p, ''),
                                          'before': b, 'after': a, 'auto': auto})
            out.setdefault(p, []).append({'side': 'succ', 'other': s, 'other_name': name.get(s, ''),
                                          'before': b, 'after': a, 'auto': auto})
    return out


def revalidate(data, config, accepted):
    """Apply the accepted corrections to a graph copy, re-run the detection engine, and
    report which accepted findings are now genuinely resolved. Returns the fresh (post-
    correction) findings plus the resolved finding_ids and the recomputed KPIs. The corrections
    are completed first (complete_ops): every out-of-sequence tie of an accepted activity, and no
    open end left behind - the same op list the corrected file is written from."""
    full = complete_ops(data, config, accepted)
    new_rels = apply_ops_to_relationships(data, full)
    fresh = run_out_of_sequence(ScheduleGraph(_shim(data, new_rels)), config)
    fresh_ids = {f['finding_id'] for f in fresh['findings']}
    still_oos_acts = {f['activity_id'] for f in fresh['findings']}
    # A finding is resolved only when its out-of-sequence condition genuinely no longer holds. For a
    # 'replace' the finding_id changes with the predecessor, so a vanished id is NOT proof on its own:
    # if the activity is still out of sequence via the new predecessor the finding is not resolved —
    # honest accounting that never credits a "fix" that merely moved the problem.
    resolved = []
    for o in (_norm_op(x) for x in accepted):
        fid = o['finding_id']
        if not fid or fid in fresh_ids:
            continue
        if o['action'] == 'replace' and o['succ_id'] and o['succ_id'] in still_oos_acts:
            continue
        resolved.append(fid)
    resolved = sorted(set(resolved))
    return {
        'findings': fresh['findings'],
        'fresh_ids': sorted(fresh_ids),
        'resolved': resolved,
        'kpis': fresh['kpis'],
        # the whole check as it stands after the corrections (comment 89): the Executive Dashboard,
        # the distribution by WBS and the impact assessment are redrawn from it at once
        'module': {k: fresh.get(k) for k in ('score', 'grade', 'pct', 'kpis', 'wbs_summary')},
        # every tie the corrections changed, per activity - highlighted on screen, in the report and Excel
        'changes': changed_ties(data, full),
        # what Apply did beyond the ties shown: further out-of-sequence ties of the same
        # activities, and ties added so no activity is left without a predecessor / successor
        'auto': {'ties': sum(1 for o in full if o.get('auto') == 'tie'),
                 'closed': sum(1 for o in full if o.get('auto') == 'close')},
    }


def revalidate_from_path(path, config, accepted):
    return revalidate(parse_file(path), config, accepted)


# ── Corrected-file export ────────────────────────────────────────────────────

def _lag_day_hours(data, pred_code, succ_code):
    """Hours per day of the link's LAG calendar (the project's 'Calendar for scheduling
    Relationship Lag' - predecessor by default) - the same basis the parser used to turn the
    file's lag hours into lag_days (finding P16), so days -> hours round-trips exactly."""
    by_code = {a.get('id'): oid for oid, a in data.activities.items()}
    return lag_day_hours(data, by_code.get(pred_code), by_code.get(succ_code))


def _current_type(data, pc, sc):
    oid_to_code = _oid_to_code(data)
    for r in data.relationships:
        if oid_to_code.get(r['pred_id']) == pc and oid_to_code.get(r['succ_id']) == sc:
            return r.get('type', 'FS')
    return 'FS'


def to_file_ops(accepted, data):
    """Translate accepted corrections → the file-writer op vocabulary shared with
    ``p6_compare.revert``: set_rel / remove_rel / add_rel, with the lag already converted
    to hours (P6 stores lag in hours in both XML and XER)."""
    ops = []
    for raw in accepted:
        o = _norm_op(raw)
        act, pc, sc = o['action'], o['pred_id'], o['succ_id']
        if act in ('data', 'manual') or not sc:
            continue
        lag_days = o['new_lag_days'] if o['new_lag_days'] is not None else 0.0
        new_pc = o['new_pred_id'] if act == 'replace' and o['new_pred_id'] else pc
        lag_hours = float(lag_days) * _lag_day_hours(data, new_pc, sc)
        typ = o['new_type'] or _current_type(data, pc, sc)
        if act == 'remove':
            ops.append({'kind': 'remove_rel', 'pred_code': pc, 'succ_code': sc})
        elif act == 'change':
            ops.append({'kind': 'set_rel', 'pred_code': pc, 'succ_code': sc,
                        'type': typ, 'lag_hours': lag_hours})
        elif act == 'add':
            ops.append({'kind': 'add_rel', 'pred_code': pc, 'succ_code': sc,
                        'type': typ, 'lag_hours': lag_hours})
        elif act == 'replace':
            ops.append({'kind': 'remove_rel', 'pred_code': pc, 'succ_code': sc})
            npc = o['new_pred_id']
            if npc:   # no new predecessor → behave as a plain remove (never re-add the same tie)
                ops.append({'kind': 'add_rel', 'pred_code': npc, 'succ_code': sc,
                            'type': typ, 'lag_hours': lag_hours})
    return _repoint(ops)


def _repoint(ops):
    """Turn each 'remove a tie' + 'add a tie' that share an activity into ONE 'repoint_rel': the
    existing relationship keeps its identity in P6 and only changes its predecessor (or
    successor), type and lag. P6 applies an edit of an existing relationship on import every time;
    a relationship that is new to the file was not created there (15-Oct: the replacement
    predecessor showed on screen but never reached P6). Adds with no removed partner stay adds."""
    removes = [o for o in ops if o['kind'] == 'remove_rel']
    used = set()
    out = []
    for o in ops:
        if o['kind'] != 'add_rel':
            continue
        # same successor (a new predecessor), else same predecessor (a new successor)
        partner = next((i for i, r in enumerate(removes) if i not in used and r['succ_code'] == o['succ_code']), None)
        if partner is None:
            partner = next((i for i, r in enumerate(removes) if i not in used and r['pred_code'] == o['pred_code']), None)
        if partner is None:
            out.append(o)
            continue
        used.add(partner)
        r = removes[partner]
        out.append({'kind': 'repoint_rel', 'pred_code': r['pred_code'], 'succ_code': r['succ_code'],
                    'new_pred_code': o['pred_code'], 'new_succ_code': o['succ_code'],
                    'type': o.get('type', 'FS'), 'lag_hours': o.get('lag_hours', 0.0)})
    kept = [o for o in ops if o['kind'] not in ('add_rel', 'remove_rel')]
    kept += [r for i, r in enumerate(removes) if i not in used]
    return kept + out


_CORRECTION_NOTE = ('Out-of-Sequence — corrected schedule. Relationship logic revised to '
                    'match actual execution (accepted corrections only). Actuals and dates '
                    'untouched — reschedule (F9) in P6.')


def write_corrected(source_path, accepted, out_path, note=_CORRECTION_NOTE, config=None):
    """Write a corrected copy of the imported schedule with the accepted relationship
    corrections applied, in the same format as the source (P6 XML or XER). With ``config`` the
    corrections are completed first (complete_ops), exactly as the on-screen re-check does."""
    data = parse_file(source_path)
    if config is not None:
        from p6_audit.engine import _enrich
        _enrich(data, config)
        accepted = complete_ops(data, config, accepted)
    ops = to_file_ops(accepted, data)
    low = source_path.lower()
    if low.endswith('.xml'):
        from p6_compare.revert import write_corrected_xml
        return write_corrected_xml(source_path, ops, out_path, note=note)
    if low.endswith('.xer'):
        return write_corrected_xer(source_path, ops, out_path, note=note)
    raise ValueError('Corrected schedule export supports P6 XML (.xml) or XER (.xer) only.')


# ── The corrections as an Excel workbook (Ibrahim, 15-Oct) ──────────────────────
_P6_TYPE = {'FS': 'PR_FS', 'SS': 'PR_SS', 'FF': 'PR_FF', 'SF': 'PR_SF'}


def write_corrections_xlsx(source_path, accepted, out_path, config):
    """The same corrections as the corrected schedule, as a workbook the planner can check and
    import into P6 (File > Import > Spreadsheet):

      * 'TASKPRED'      - P6's spreadsheet-import layout for relationships: row 1 the P6 field
                          codes, row 2 the titles, then one row per relationship to update, add or
                          delete ('Delete This Row' = Y);
      * 'Changes'       - every change exactly as the screen lists it: activity, related activity,
                          side, before -> after, why;
      * 'After the fix' - every relationship of each corrected activity after the corrections, to
                          compare with the corrected XML / XER and with P6 after importing.
    Returns {'rows', 'changes', 'out_path'}."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    data = parse_file(source_path)
    from p6_audit.engine import _enrich
    _enrich(data, config)
    full = complete_ops(data, config, accepted)
    fops = to_file_ops(full, data)
    code = _oid_to_code(data)
    name = {a.get('id') or '': a.get('name') or '' for a in data.activities.values()}
    hours = {}
    for r in data.relationships:
        hours[(code.get(r['pred_id']), code.get(r['succ_id']))] = (r.get('type', 'FS'), r.get('lag_hours'))

    def cur(pc, sc):
        t, h = hours.get((pc, sc), ('FS', 0.0))
        return _P6_TYPE.get(t, 'PR_FS'), (h if h is not None else 0.0)

    rows = []
    for o in fops:
        k, pc, sc = o['kind'], o.get('pred_code'), o.get('succ_code')
        if k == 'set_rel':
            rows.append([sc, pc, _P6_TYPE.get(o.get('type'), 'PR_FS'), round(o.get('lag_hours') or 0.0, 2), '', 'Changed'])
        elif k == 'remove_rel':
            t, h = cur(pc, sc)
            rows.append([sc, pc, t, round(h, 2), 'Y', 'Removed'])
        elif k == 'add_rel':
            rows.append([sc, pc, _P6_TYPE.get(o.get('type'), 'PR_FS'), round(o.get('lag_hours') or 0.0, 2), '', 'Added'])
        elif k == 'repoint_rel':
            t, h = cur(pc, sc)
            rows.append([sc, pc, t, round(h, 2), 'Y', 'Removed (replaced)'])
            rows.append([o['new_succ_code'], o['new_pred_code'], _P6_TYPE.get(o.get('type'), 'PR_FS'),
                         round(o.get('lag_hours') or 0.0, 2), '', 'Added (replacement)'])

    wb = Workbook()
    hdr = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='1F4E79')
    amber = PatternFill('solid', fgColor='FEF3C7')

    ws = wb.active
    ws.title = 'TASKPRED'
    ws.append(['task_id', 'pred_task_id', 'pred_type', 'lag_hr_cnt', 'delete_record_flag', 'controlyx_change'])
    ws.append(['Successor', 'Predecessor', 'Relationship Type', 'Lag (h)', 'Delete This Row', 'Change'])
    for r in rows:
        ws.append(r)
    for c in ws[2]:
        c.font, c.fill = hdr, fill

    ch = changed_ties(data, full)
    focus = sorted({f for f in ch})
    wc = wb.create_sheet('Changes')
    wc.append(['Activity ID', 'Activity Name', 'Side', 'Related Activity ID', 'Related Activity Name',
               'Before', 'After', 'Why'])
    n_changes = 0
    seen = set()
    for act in focus:
        for c in ch[act]:
            key = (act, c['side'], c['other'], c['after'])
            if c['side'] != 'pred' or key in seen:
                continue                     # each tie once, on its successor's row (as the screen)
            seen.add(key)
            why = {'tie': 'further out-of-sequence tie of the same activity',
                   'close': 'keeps the activity connected (commencement / completion)'}.get(c.get('auto'), 'driving tie')
            wc.append([act, name.get(act, ''), 'Predecessor', c['other'], c.get('other_name', ''),
                       c.get('before', ''), c.get('after', ''), why])
            n_changes += 1
    for c in wc[1]:
        c.font, c.fill = hdr, fill

    after_rels = apply_ops_to_relationships(data, full)
    touched = {act for act in focus}
    wa = wb.create_sheet('After the fix')
    wa.append(['Predecessor ID', 'Predecessor Name', 'Successor ID', 'Successor Name', 'Type', 'Lag (d)', 'Changed'])
    changed_pairs = {(c['other'], a) for a in ch for c in ch[a] if c['side'] == 'pred'}
    for r in after_rels:
        pc, sc = code.get(r['pred_id']), code.get(r['succ_id'])
        if sc not in touched and pc not in touched:
            continue
        flag = 'Yes' if (pc, sc) in changed_pairs else ''
        wa.append([pc, name.get(pc, ''), sc, name.get(sc, ''), r.get('type', 'FS'), round(r.get('lag_days') or 0.0, 2), flag])
        if flag:
            for c in wa[wa.max_row]:
                c.fill = amber
    for c in wa[1]:
        c.font, c.fill = hdr, fill
    for sheet in (ws, wc, wa):
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
        for row in sheet.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=True)
    wb.save(out_path)
    return {'rows': len(rows), 'changes': n_changes, 'out_path': out_path}


# ── XER writer (in-place TASKPRED edit — no XER writer existed before) ────────

_XER_PRED_TYPE = {'FS': 'PR_FS', 'SS': 'PR_SS', 'FF': 'PR_FF', 'SF': 'PR_SF'}


def read_xer_keeping_encoding(path):
    """(text, encoding, newline) of an XER file, decoded the way the parser reads it - so a
    corrected copy is written back in the SAME encoding and line ending. P6 writes XER in the
    Windows code page; re-saving it as UTF-8 garbles Arabic text and currency symbols on import."""
    with open(path, 'rb') as fh:
        raw = fh.read()
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return text, enc, ('\r\n' if '\r\n' in text else '\n')


def _fmt_hours(v):
    v = float(v or 0.0)
    return str(int(v)) if v == int(v) else repr(v)


def write_corrected_xer(source_path, ops, out_path, note=None):
    """Apply set_rel / remove_rel / add_rel to a copy of the update XER by editing the
    TASKPRED table in place — change a predecessor's type/lag, drop a relationship row,
    or append a new one (cloned from an existing row so every required column is present).
    Every other table is preserved verbatim. TASKPRED matches on internal task ids, so
    ops (keyed by activity code) are expanded to all task-id pairs for that code — duplicate
    codes are handled the same way the XML writer does. Returns {'applied', 'out_path'}."""
    text, encoding, newline = read_xer_keeping_encoding(source_path)
    lines = [ln.rstrip('\r') for ln in text.split('\n')]

    # Pass 1 — map activity code → [task_id] from the TASK table.
    code_to_taskids = {}
    cur, fields = None, []
    for ln in lines:
        parts = ln.split('\t')
        tag = parts[0]
        if tag == '%T':
            cur = parts[1] if len(parts) > 1 else None
            fields = []
        elif tag == '%F':
            fields = parts[1:]
        elif tag == '%R' and cur == 'TASK':
            row = dict(zip(fields, parts[1:]))
            code, tid = row.get('task_code'), row.get('task_id')
            if code and tid:
                code_to_taskids.setdefault(code, []).append(tid)

    # Expand code-pair ops → (pred_task_id, succ_task_id) targets.
    remove_pairs, set_pairs, add_ops, repoint = set(), {}, [], {}
    for op in ops:
        kind = op.get('kind')
        pc, sc = op.get('pred_code'), op.get('succ_code')
        p_ids, s_ids = code_to_taskids.get(pc, []), code_to_taskids.get(sc, [])
        if kind == 'remove_rel':
            for pi in p_ids:
                for si in s_ids:
                    remove_pairs.add((pi, si))
        elif kind == 'set_rel':
            for pi in p_ids:
                for si in s_ids:
                    set_pairs[(pi, si)] = (op.get('type', 'FS'), op.get('lag_hours', 0.0))
        elif kind == 'add_rel':
            if p_ids and s_ids:
                add_ops.append((p_ids[0], s_ids[0], op.get('type', 'FS'), op.get('lag_hours', 0.0)))
        elif kind == 'repoint_rel':
            np_ids = code_to_taskids.get(op.get('new_pred_code'), [])
            ns_ids = code_to_taskids.get(op.get('new_succ_code'), [])
            if np_ids and ns_ids:
                first = True
                for pi in p_ids:
                    for si in s_ids:
                        # the first copy is re-pointed, any other copy of the old tie is dropped
                        repoint[(pi, si)] = (np_ids[0], ns_ids[0], op.get('type', 'FS'),
                                             op.get('lag_hours', 0.0)) if first else None
                        first = False

    # Pass 2 — rewrite. Track TASKPRED fields + the max row id + a template row for adds.
    tp_fields, tp_template, max_pred_id = [], None, 0
    for ln in lines:
        parts = ln.split('\t')
        if parts[0] == '%T':
            cur = parts[1] if len(parts) > 1 else None
        elif parts[0] == '%F' and cur == 'TASKPRED':
            tp_fields = parts[1:]
        elif parts[0] == '%R' and cur == 'TASKPRED':
            if tp_template is None:
                tp_template = parts[1:]
            row = dict(zip(tp_fields, parts[1:]))
            try:
                max_pred_id = max(max_pred_id, int(row.get('task_pred_id') or 0))
            except (TypeError, ValueError):
                pass

    def _idx(name):
        return tp_fields.index(name) if name in tp_fields else -1

    i_pred_id = _idx('task_pred_id')
    i_task = _idx('task_id')
    i_pred_task = _idx('pred_task_id')
    i_type = _idx('pred_type')
    i_lag = _idx('lag_hr_cnt')

    def _new_row(pred_task, succ_task, rel_type, lag_hours):
        nonlocal max_pred_id
        vals = list(tp_template) if tp_template else [''] * len(tp_fields)
        if len(vals) < len(tp_fields):
            vals += [''] * (len(tp_fields) - len(vals))
        max_pred_id += 1
        if i_pred_id >= 0:
            vals[i_pred_id] = str(max_pred_id)
        if i_task >= 0:
            vals[i_task] = succ_task
        if i_pred_task >= 0:
            vals[i_pred_task] = pred_task
        if i_type >= 0:
            vals[i_type] = _XER_PRED_TYPE.get(rel_type, 'PR_FS')
        if i_lag >= 0:
            vals[i_lag] = _fmt_hours(lag_hours)
        return '%R\t' + '\t'.join(vals)

    out_lines, applied = [], 0
    cur = None
    in_taskpred = False
    emitted_adds = False

    def _flush_adds():
        nonlocal applied, emitted_adds
        if add_ops and not emitted_adds and i_task >= 0 and i_pred_task >= 0:
            for (pi, si, t, lag) in add_ops:
                out_lines.append(_new_row(pi, si, t, lag))
                applied += 1
        emitted_adds = True

    for ln in lines:
        parts = ln.split('\t')
        tag = parts[0]
        if tag == '%T':
            # Leaving TASKPRED for a new table → append any pending new rows first.
            if in_taskpred:
                _flush_adds()
            cur = parts[1] if len(parts) > 1 else None
            in_taskpred = (cur == 'TASKPRED')
            out_lines.append(ln)
            continue
        if tag == '%E':
            if in_taskpred:
                _flush_adds()
            in_taskpred = False
            out_lines.append(ln)
            continue
        if in_taskpred and tag == '%R':
            vals = parts[1:]
            pair = (vals[i_pred_task] if i_pred_task >= 0 and i_pred_task < len(vals) else None,
                    vals[i_task] if i_task >= 0 and i_task < len(vals) else None)
            if pair in remove_pairs:
                applied += 1
                continue  # drop the relationship row
            if pair in repoint:
                target = repoint.pop(pair)
                if target is None:
                    continue      # a duplicate copy of the re-pointed tie
                npi, nsi, rel_type, lag_hours = target
                vals[i_pred_task], vals[i_task] = npi, nsi
                if i_type >= 0 and i_type < len(vals):
                    vals[i_type] = _XER_PRED_TYPE.get(rel_type, 'PR_FS')
                if i_lag >= 0 and i_lag < len(vals):
                    vals[i_lag] = _fmt_hours(lag_hours)
                applied += 1
                out_lines.append('%R\t' + '\t'.join(vals))
                continue
            if pair in set_pairs:
                rel_type, lag_hours = set_pairs[pair]
                if i_type >= 0 and i_type < len(vals):
                    vals[i_type] = _XER_PRED_TYPE.get(rel_type, 'PR_FS')
                if i_lag >= 0 and i_lag < len(vals):
                    vals[i_lag] = _fmt_hours(lag_hours)
                applied += 1
                out_lines.append('%R\t' + '\t'.join(vals))
                continue
            out_lines.append(ln)
            continue
        out_lines.append(ln)

    # File with no %E terminator: flush any adds at the very end.
    if in_taskpred:
        _flush_adds()

    if note:
        # XER comments aren't standard; carry the note as an ERMHDR-adjacent comment line
        # that P6 ignores on import is risky — instead leave the file clean. (Kept param for
        # signature parity with the XML writer.)
        pass

    with open(out_path, 'w', encoding=encoding, newline='') as fh:
        fh.write(newline.join(out_lines))
    return {'applied': applied, 'out_path': out_path}
