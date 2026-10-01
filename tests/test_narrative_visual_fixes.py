"""Baseline Narrative — visible defects found in the page-by-page PNG review (NARRFIX) that
the page checker cannot see:

1. Word §11 chevrons broke a word in two ('Column / s'): Word's default chevron notch is
   half the shape's height, narrower text than the layout (and the SVG, notch 14 px) allows.
2. Word Appendix Critical Path: a year in its first month's cell alone wrapped '202 / 5'.
3. Word §13.1: 22 monthly man-hour labels at 10 pt printed over each other.
4. Word + PDF §6: white on-ring labels ran past a slice ('ocurement 17.3%').

python-docx only (no Word needed).
"""
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn

from p6_narrative import docx_native as N
from p6_narrative import docx_template as T
from p6_narrative import docx_writer as W
from p6_narrative import html as H
from p6_narrative.util import chevron_layout, ring_label_fits, ring_label_spot


def _xml(doc, tmp_path, part='word/document.xml'):
    out = str(tmp_path / 'x.docx')
    doc.save(out)
    with zipfile.ZipFile(out) as z:
        return z.read(part).decode('utf-8'), z.namelist(), out


# ── 1. chevrons ────────────────────────────────────────────────────────────────
def test_word_chevron_notch_matches_the_layout_so_no_word_breaks(tmp_path):
    labels = ['Pile Works', 'Foundation', 'Columns', 'Insulation Works', 'SOG', 'Beams & Slab']
    d = Document()
    assert N.add_chevron_flow(d, labels) is not None
    body, _, _ = _xml(d, tmp_path)
    lay = {it['label']: it for row in chevron_layout(labels)['rows'] for it in row}
    shapes = re.findall(r'<wps:cNvPr id="\d+" name="([^"]+)"/>.*?<a:prstGeom prst="(\w+)"><a:avLst>'
                        r'(.*?)</a:avLst>', body)
    assert len(shapes) == len(labels)
    for name, prst, av in shapes:
        m = re.search(r'<a:gd name="adj" fmla="val (\d+)"/>', av)
        assert m, (name, av)                       # the notch is set, not Word's default
        it = lay[name.replace('&amp;', '&')]
        notch_px = int(m.group(1)) / 100000.0 * min(it['w'], it['h'])
        assert abs(notch_px - 14) < 0.1, (name, notch_px)
        # the chevron's text area (w - 2 notches) holds the layout's longest line
        longest = max(len(ln) for ln in it['lines']) * it['font_px'] * 0.52
        room = it['w'] - (2 if prst == 'chevron' else 0.5) * notch_px
        assert room >= longest + 8, (name, room, longest)


# ── 2. critical-path sweep year header ──────────────────────────────────────────
def _critpath(months):
    return {'available': True, 'narrative': 'x', 'months': months,
            'zones': [{'label': 'Silo 3', 'cells': [None] * len(months)}]}


def _months(spec):
    out = []
    for y, ms in spec:
        out += [{'y': y, 'm': m, 'label': 'M%d' % m} for m in ms]
    return out


def test_word_year_header_spans_its_months_centred_and_never_wraps():
    months = _months([(2025, [12]), (2026, range(1, 13)), (2027, [1, 2])])
    d = Document()
    T.apply_base_styles(d)
    W._render_critpath(d, _critpath(months), 17, None)
    t = d.tables[-1]
    head = t.rows[0]._tr
    tcs = head.findall(qn('w:tc'))
    assert len(tcs) == 1 + 3                             # label cell + one cell per year
    spans = []
    for tc in tcs[1:]:
        gs = tc.find(qn('w:tcPr')).find(qn('w:gridSpan'))
        spans.append(int(gs.get(qn('w:val'))) if gs is not None else 1)
        jc = tc.find('.//' + qn('w:jc'))
        assert jc is not None and jc.get(qn('w:val')) == 'center'
    assert spans == [1, 12, 2]
    texts = [''.join(x.text for x in tc.iter(qn('w:t'))) for tc in tcs[1:]]
    # a one-month year (0.36 in) holds '2025' with slim margins; the 12-month one surely does
    assert texts[1] == '2026'
    for txt, sp in zip(texts, spans):
        w_in = (6.9 - 1.5) / len(months) * sp - 0.04
        assert len(txt) * 7.5 * 0.55 <= w_in * 72 + 0.01, (txt, w_in)


def test_word_year_too_narrow_for_four_digits_prints_two():
    months = _months([(2024, [12]), (2025, range(1, 13)), (2026, range(1, 13)), (2027, range(1, 13))])
    assert W._year_label(2024, (6.9 - 1.5) / len(months) - 0.04, 7.5) == '’24'
    assert W._year_label(2025, 0.5, 7.5) == '2025'


# ── 3. bar labels ─────────────────────────────────────────────────────────────
def _chart_xml(tmp_path, values):
    d = Document()
    months = ['M%02d' % i for i in range(len(values))]
    assert N.add_bar_chart(d, months, values, '', data_labels=True, num_fmt='#,##0') is not None
    out = str(tmp_path / 'c.docx')
    d.save(out)
    with zipfile.ZipFile(out) as z:
        name = next(n for n in z.namelist() if n.startswith('word/charts/chart'))
        return z.read(name).decode('utf-8')


def test_word_many_monthly_labels_are_turned_upright_with_headroom(tmp_path):
    vals = [5938, 5913, 6124, 4741, 6141, 6943, 9205, 8361, 11700] + [5000] * 13   # GBT §13.1: 22 months
    xml = _chart_xml(tmp_path, vals)
    dl = re.search(r'<c:dLbls>.*?</c:dLbls>', xml).group(0)
    assert 'rot="-5400000"' in dl and 'sz="750"' in dl
    assert dl.index('<c:txPr>') < dl.index('<c:dLblPos')          # schema order
    ax = re.search(r'<c:valAx>.*?</c:valAx>', xml).group(0)
    mx = float(re.search(r'<c:max val="([\d.]+)"/>', ax).group(1))
    unit = float(re.search(r'<c:majorUnit val="([\d.]+)"/>', ax).group(1))
    assert mx >= 11700 * 1.1 and '<c:min val="0"/>' in ax          # raised, and from 0
    assert (mx / unit) == int(mx / unit) and 4 <= mx / unit <= 8     # round gridlines
    assert ax.index('<c:max ') < ax.index('<c:min ') and ax.index('<c:crossAx') < ax.index('<c:majorUnit')


def test_word_few_bars_keep_the_default_labels(tmp_path):
    xml = _chart_xml(tmp_path, [4151, 22487, 20504, 21023, 33727, 8265, 9594, 4390])   # SG: 8 months
    assert '<c:txPr>' not in re.search(r'<c:dLbls>.*?</c:dLbls>', xml).group(0)
    assert '<c:max ' not in xml


def test_label_fit_steps_down_before_rotating():
    assert N._bar_label_fit([1] * 19, '#,##0') == (None, 0, None, None)
    pt, rot, mx, unit = N._bar_label_fit([12345] * 13, '#,##0')           # 13 bars of '12,345'
    assert rot == 0 and 7 <= pt < 10 and mx is None and unit is None


# ── 4. doughnut on-ring labels ─────────────────────────────────────────────────
SG_ROWS = [('Construction', 82.2), ('Procurement', 17.3), ('Unclassified', 0.5)]


def _angles(rows):
    tot, acc, out = sum(p for _, p in rows), 0.0, []
    for n, p in rows:
        out.append((n, p, acc / tot * 360.0, (acc + p) / tot * 360.0))
        acc += p
    return out


def test_ring_label_spot_puts_every_dominant_label_inside_its_slice():
    for n, p, a0, a1 in _angles(SG_ROWS):
        if p < 15:
            continue
        pct = '%s%%' % H._fmt_pct(p)
        assert not (n == 'Procurement' and ring_label_fits(n, pct, 205, 165, 124, 73, a0, a1))  # the defect
        spot = ring_label_spot(n, pct, 205, 165, 124, 73, a0, a1)
        assert spot is not None, n
        ang, npx, ppx, k = spot
        assert a0 <= ang <= a1 and npx >= 9.5
        assert ring_label_fits(n, pct, 205, 165, 124, 73, a0, a1, npx, ppx, at=ang)


def test_pdf_and_word_doughnut_place_the_label_at_the_same_spot(tmp_path):
    rows = [{'name': n, 'pct': p, 'value': p * 1000} for n, p in SG_ROWS]
    svg = H._doughnut(rows, 'TOTAL', '243,805,397', lambda r: str(r['value']))
    # the Procurement name is drawn smaller, at its fitted spot, not at 15 px on the mid-angle
    m = re.search(r'<text x="([\d.]+)" y="([\d.]+)"[^>]*font-size="([\d.]+)"[^>]*>Procurement</text>', svg)
    assert m and float(m.group(3)) < 15
    n, p, a0, a1 = _angles(SG_ROWS)[1]
    ang, npx, _, k = ring_label_spot('Procurement', '17.3%', 205, 165, 124, 73, a0, a1)
    lx, ly = H._polar(205, 165, (124 + 73) / 2.0, ang)
    assert abs(float(m.group(1)) - lx) < 0.2 and abs(float(m.group(2)) - (ly - 5 * k)) < 0.2
    d = Document()
    assert N.add_doughnut(d, [n for n, _ in SG_ROWS], [p * 1000 for _, p in SG_ROWS],
                          pcts=[p for _, p in SG_ROWS], total=243805397) is not None
    body, _, _ = _xml(d, tmp_path)
    assert 'Procurement' in body and body.count('name="ring-name"') == 2
