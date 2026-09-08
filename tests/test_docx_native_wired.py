"""SLICE C — the Word renderers are wired to NATIVE charts/diagrams, not PNGs.

These tests drive the ``docx_writer`` / ``docx_calendar`` section renderers directly
and assert the produced .docx carries REAL Word chart objects (``word/charts/chartN.xml``)
and native grouped shapes (``wpg:wgp`` / ``wps:wsp``) for value / costbars / cashflow /
wbs_tree / seq / calendar-histogram — with NO rasterised image part (``word/media/*``,
``<a:blip``) for those kinds, and WITHOUT any Chrome (native charts never need one).

Every renderer must also stay None-safe: a degenerate payload falls back to the native
table / outline and the export never raises.
"""
import io
import os
import zipfile

from docx import Document

from p6_narrative import docx_calendar
from p6_narrative import docx_writer as W


# ── helpers ───────────────────────────────────────────────────────────────────
def _sub(document, number=3):
    return W._Sub(document, number)


def _save_reopen(document, tmp_path, name='out.docx'):
    out = os.path.join(str(tmp_path), name)
    document.save(out)
    Document(out)          # must reopen without raising
    return out


def _names(path):
    with zipfile.ZipFile(path) as z:
        return z.namelist()


def _read(path, member):
    with zipfile.ZipFile(path) as z:
        return z.read(member).decode('utf-8')


def _chart_xml(path):
    """Concatenated text of every word/charts/chartN.xml part."""
    return '\n'.join(_read(path, n) for n in _names(path)
                     if n.startswith('word/charts/chart') and n.endswith('.xml'))


def _no_raster(path):
    """No rasterised picture anywhere: no media parts, no blip in the body."""
    names = _names(path)
    assert [n for n in names if n.startswith('word/media/')] == [], names
    assert '<a:blip' not in _read(path, 'word/document.xml')


_VALUE = {'total': 100.0, 'rows': [
    {'name': 'Civil', 'cost': 56.0, 'pct': 56.0},
    {'name': 'Mechanical', 'cost': 28.0, 'pct': 28.0},
    {'name': 'Electrical', 'cost': 16.0, 'pct': 16.0},
]}

_COSTBARS = {'total': 100.0, 'rows': [
    {'name': 'Foundations', 'cost': 60.0, 'pct': 60.0},
    {'name': 'Superstructure', 'cost': 40.0, 'pct': 40.0},
]}

_CASHFLOW = {'total': 30.0, 'monthly': [
    {'label': 'Jun 2025', 'cost': 4.2, 'pct': 14.0},
    {'label': 'Jul 2025', 'cost': 6.8, 'pct': 36.7},
    {'label': 'Aug 2025', 'cost': 7.6, 'pct': 62.0},
    {'label': 'Sep 2025', 'cost': 11.4, 'pct': 100.0},
]}

_TREE = {'name': 'Project', 'children': [
    {'name': 'Civil', 'children': [{'name': 'Foundations'}, {'name': 'Structures'}]},
    {'name': 'Mechanical', 'children': [{'name': 'Piping'}]},
]}

_WBS = {'overview': _TREE, 'branches': [{'root': _TREE}]}

_SEQ = {'worlds': [{'world': 'Construction', 'fronts': [
    {'title': 'Building 01', 'sequence': ['Excavation', 'Foundation', 'Columns', 'Slab'],
     'instances': ['B01', 'B02'],
     'activities': [{'id': 'A1', 'name': 'Excavation B01', 'wbs': 'CON'},
                    {'id': 'A2', 'name': 'Foundation B01', 'wbs': 'CON'}]},
]}]}

_MONTHS = [
    {'label': 'Jan 2025', 'working_days': 22, 'holidays': 1, 'working_hours': 176,
     'days': list(range(31))},
    {'label': 'Feb 2025', 'working_days': 20, 'holidays': 0, 'working_hours': 160,
     'days': list(range(28))},
]


# ── value: native doughnut + the cost table beneath ───────────────────────────
def test_value_is_native_doughnut_plus_table(tmp_path):
    doc = Document()
    W._render_value(doc, _VALUE, _sub(doc), None, None)
    out = _save_reopen(doc, tmp_path)
    xml = _chart_xml(out)
    assert '<c:doughnutChart>' in xml, 'value donut is not a native doughnut chart'
    # the cost table is kept beneath the chart (PDF shows both)
    assert len(Document(out).tables) >= 1
    _no_raster(out)


# ── costbars: native column chart, no image ───────────────────────────────────
def test_costbars_is_native_bar(tmp_path):
    doc = Document()
    W._render_costbars(doc, _COSTBARS, _sub(doc), None, None)
    out = _save_reopen(doc, tmp_path)
    xml = _chart_xml(out)
    assert '<c:barChart>' in xml and '<c:barDir val="col"/>' in xml
    _no_raster(out)


# ── cashflow: native column chart from monthly label/cost, no image ───────────
def test_cashflow_is_native_bar(tmp_path):
    doc = Document()
    W._render_cashflow(doc, _CASHFLOW, _sub(doc), None, None)
    out = _save_reopen(doc, tmp_path)
    xml = _chart_xml(out)
    assert '<c:barChart>' in xml
    assert 'Jun 2025' in xml and 'Sep 2025' in xml
    _no_raster(out)


# ── wbs_tree: native org-chart group shapes (overview + branch), no image ─────
def test_wbs_tree_is_native_org_chart(tmp_path):
    doc = Document()
    W._render_wbs_tree(doc, _WBS, _sub(doc), None, None)
    out = _save_reopen(doc, tmp_path)
    body = _read(out, 'word/document.xml')
    assert 'wpg:wgp' in body and 'wps:wsp' in body
    assert body.count('<w:drawing') >= 2, 'expected overview + branch org-charts'
    assert 'Foundations' in body
    _no_raster(out)


# ── seq: native process chevrons per front + the drill-down activities table ──
def test_seq_is_native_process_plus_table(tmp_path):
    doc = Document()
    W._render_seq(doc, _SEQ, _sub(doc), None, None)
    out = _save_reopen(doc, tmp_path)
    body = _read(out, 'word/document.xml')
    assert 'wpg:wgp' in body and 'prst="chevron"' in body
    assert 'Excavation' in body
    # the drill-down P6 activity table is kept beneath the flow
    tbls = Document(out).tables
    assert any('Activity' in [c.text for c in t.rows[0].cells] for t in tbls)
    _no_raster(out)


# ── calendar histogram: native STACKED column chart, no image ─────────────────
def test_calendar_histogram_is_native_stacked_bar(tmp_path):
    doc = Document()
    docx_calendar._monthly(doc, _MONTHS, None)
    out = _save_reopen(doc, tmp_path)
    xml = _chart_xml(out)
    assert '<c:barChart>' in xml and '<c:grouping val="stacked"/>' in xml
    assert '22C55E' in xml and 'EF4444' in xml       # working green + non-working red
    _no_raster(out)


# ══ fallbacks: a degenerate payload falls back to the native table/outline ════
def test_value_fallback_table_on_bad_rows(tmp_path):
    # rows with non-numeric cost/pct → native pie returns None, cost table still lands
    doc = Document()
    bad = {'total': None, 'rows': [{'name': 'X', 'cost': None, 'pct': None}]}
    W._render_value(doc, bad, _sub(doc), None, None)        # must not raise
    out = _save_reopen(doc, tmp_path)
    assert len(Document(out).tables) >= 1


def test_wbs_tree_fallback_outline(tmp_path):
    # if the org-chart cannot be built, the indented outline still renders the names
    doc = Document()
    W._render_wbs_tree(doc, _WBS, _sub(doc), None, None)     # native path
    # force fallback: monkeypatch add_org_chart to return None
    from p6_narrative import docx_native
    orig = docx_native.add_org_chart
    docx_native.add_org_chart = lambda *a, **k: None
    try:
        doc2 = Document()
        W._render_wbs_tree(doc2, _WBS, _sub(doc2), None, None)
        out = _save_reopen(doc2, tmp_path, 'fallback.docx')
        text = '\n'.join(p.text for p in Document(out).paragraphs)
        assert 'Civil' in text and 'Foundations' in text
    finally:
        docx_native.add_org_chart = orig


def test_seq_fallback_arrow_line(tmp_path):
    doc = Document()
    from p6_narrative import docx_native
    orig = docx_native.add_process
    docx_native.add_process = lambda *a, **k: None
    try:
        W._render_seq(doc, _SEQ, _sub(doc), None, None)
        out = _save_reopen(doc, tmp_path, 'seqfb.docx')
        text = '\n'.join(p.text for p in Document(out).paragraphs)
        assert 'Excavation' in text        # arrow-line fallback carries the steps
    finally:
        docx_native.add_process = orig


def test_calendar_histogram_fallback_table(tmp_path):
    doc = Document()
    from p6_narrative import docx_native
    orig = docx_native.add_bar_chart_stacked
    docx_native.add_bar_chart_stacked = lambda *a, **k: None
    try:
        docx_calendar._monthly(doc, _MONTHS, None)
        out = _save_reopen(doc, tmp_path, 'calfb.docx')
        # the monthly stats table is the fallback (header Month | Working | ...)
        tbls = Document(out).tables
        assert any([c.text for c in t.rows[0].cells][:3] == ['Month', 'Working', 'Non-working']
                   for t in tbls)
    finally:
        docx_native.add_bar_chart_stacked = orig
