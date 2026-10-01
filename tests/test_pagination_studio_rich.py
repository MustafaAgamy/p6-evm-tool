"""[pagination:STUDIO] Reporting Studio Document PDF + Word with a RICH multi-feature
selection (GBT: overview, EVM, every Schedule Health module, Calendar Audit, Update
Analysis, Baseline Revision Comparison, Baseline Narrative - ~60 results).

Findings fixed in this step (each test fails without its fix):

* STUDIO-RICH-1 - a big selection could not be exported: every Chrome print of the Studio
  had a flat 180 s allowance (90 s for a Word section picture); the GBT document with the
  Baseline Revision 'Key Findings' is 6.4 MB / ~1 800 pages and prints in ~210 s a pass,
  so the PDF export stopped with 'did not finish within 180 s' while Chrome was still
  printing normally. The allowance now grows a minute per MB of markup.
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)


# ── STUDIO-RICH-1: the print allowance grows with the document ─────────────────────────
def test_print_timeout_grows_with_the_document():
    from p6_special.pdf_render import print_timeout
    assert print_timeout('', 180) == 180                    # a small document: the base
    assert print_timeout('x' * 1000, 180) == 180
    big = 'x' * 6_400_000                                   # the GBT rich selection
    assert print_timeout(big, 180) >= 180 + 6 * 60          # > the ~210 s a pass it needs
    assert print_timeout(big, 90) >= 90 + 6 * 60            # Word section pictures too
    assert print_timeout(None, 180) == 180


def test_chrome_pdf_passes_the_scaled_allowance(monkeypatch, tmp_path):
    import p6_export.pdf as P
    from p6_special import pdf_render
    seen = {}

    def fake_run_chrome(chrome, args, timeout=180):
        seen['timeout'] = timeout
    monkeypatch.setattr(P, 'run_chrome', fake_run_chrome)
    html = '<html><body>' + ('<div>row</div>' * 400_000) + '</body></html>'   # ~5.6 MB
    pdf_render.chrome_pdf(html, 'chrome.exe', str(tmp_path / 'out.pdf'))
    assert seen['timeout'] > 180 + 5 * 60, seen     # was a flat 180 s


def test_word_section_picture_gets_the_scaled_allowance(monkeypatch):
    import p6_export.pdf as P
    from p6_special import docx_report
    seen = {}

    def fake_run_chrome(chrome, args, timeout=180):
        seen['timeout'] = timeout
        raise RuntimeError('stop here')              # the slicer falls back quietly
    monkeypatch.setattr(P, 'run_chrome', fake_run_chrome)
    frag = '<div class="lanes">' + ('<div class="lane">#1 link</div>' * 120_000) + '</div>'
    assert docx_report._slice_section(frag, '', 'light', 'chrome.exe', 700.0, 600.0) is None
    assert seen['timeout'] > 90 + 3 * 60, seen      # was a flat 90 s


# ── STUDIO-RICH-2: checker - a cell's line is not a heading ───────────────────────────
pymupdf = None
try:
    import pymupdf
except ImportError:                                  # pragma: no cover
    pass

A4 = (595, 842)
_REC = 'Predecessor link Insulation Works -> Insulation Works: change SS to FS if the finish drives the successor'


def _pdf(path, pages):
    doc = pymupdf.open()
    for ops in pages:
        pg = doc.new_page(width=A4[0], height=A4[1])
        for _, x, y, text, size, bold in ops:
            pg.insert_text((x, y), text, fontsize=size, fontname='hebo' if bold else 'helv')
    doc.save(path)
    doc.close()
    return path


def _running(no, n):
    return [('t', 40, 30, 'PROJECT REPORT  Grain Bulk Terminal', 8, False),
            ('t', 280, 825, f'Page {no} of {n}', 8, False)]


def _tall_rows(y, n, first):
    """A Relationship-Types style register: ID + name centred in a tall row whose
    Recommendation cell (x 350) wraps over 4 lines - its FIRST line 'Predecessor link' sits
    above the centred cells (and is the start of a hinted heading text elsewhere)."""
    ops = [('t', 40, y, '#  Activity ID', 9, True), ('t', 150, y, 'Activity Name', 9, True),
           ('t', 350, y, 'Recommendation', 9, True)]
    y += 22
    for i in range(n):
        k = first + i
        ops += [('t', 350, y, 'Predecessor link', 9, False),
                ('t', 40, y + 16, f'{k}  CONS.PL.S{k}.1010', 9, False),
                ('t', 150, y + 16, f'RFT Works For Bored Piles {k}', 9, False),
                ('t', 350, y + 11, 'Insulation Works -> Insulation', 9, False),
                ('t', 350, y + 22, 'Works: change SS to FS if the', 9, False),
                ('t', 350, y + 33, 'finish drives the successor', 9, False)]
        y += 50
    return ops


def _register_pdf(path):
    n = 3
    pages = [
        _running(1, n) + [('t', 40, 70, 'Detailed findings', 14, True)] + _tall_rows(95, 14, 1),
        _running(2, n) + _tall_rows(70, 14, 15),
        _running(3, n) + [('t', 40, 70, 'Next section', 14, True)]
        + [('t', 40, y, f'Narrative line {y} of the next section, plain words.', 10, False)
           for y in range(95, 400, 14)],
    ]
    return _pdf(path, pages)


def test_checker_reads_a_tall_row_register_as_one_table(tmp_path):
    import pytest
    if pymupdf is None:
        pytest.skip('PyMuPDF missing')
    from p6_export import pagination_check as pc
    res = pc.check_pdf(_register_pdf(str(tmp_path / 'reg.pdf')), headings=[_REC])
    assert res['flags'] == [], res['flags']


def test_checker_pre_fix_broke_the_register_into_few_row_tables(tmp_path, monkeypatch):
    import pytest
    if pymupdf is None:
        pytest.skip('PyMuPDF missing')
    from p6_export import pagination_check as pc
    monkeypatch.setattr(pc, '_cell_line', lambda reg, b: False)        # the old reading
    res = pc.check_pdf(_register_pdf(str(tmp_path / 'reg.pdf')), headings=[_REC])
    assert 'table_split_few_rows' in {f['type'] for f in res['flags']}, res['flags']


def test_checker_still_flags_a_real_heading_under_a_table(tmp_path):
    """Control: a heading at the content's LEFT edge right under the table, ending the page,
    is still an orphaned heading (it is not inside a column)."""
    import pytest
    if pymupdf is None:
        pytest.skip('PyMuPDF missing')
    from p6_export import pagination_check as pc
    n = 3
    pages = [
        _running(1, n) + [('t', 40, 70, 'Detailed findings', 14, True)] + _tall_rows(95, 13, 1)
        + [('t', 40, 790, 'Predecessor link summary', 13, True)],
        _running(2, n) + [('t', 40, y, f'Summary line {y}: the links to review, plain words.', 10, False)
                          for y in range(70, 700, 14)],
        _running(3, n) + [('t', 40, 70, 'Next section', 14, True)]
        + [('t', 40, y, f'Narrative line {y} of the next section.', 10, False) for y in range(95, 790, 14)],
    ]
    res = pc.check_pdf(_pdf(str(tmp_path / 'ctl.pdf'), pages), headings=[_REC])
    assert ('orphaned_heading', 1) in {(f['type'], f['page']) for f in res['flags']}, res['flags']


# ── STUDIO-RICH-9: a Word page slice that ends INSIDE a card is not a cut picture ─────────
# The Studio .docx prints a long reused section (a Baseline Revision register / WBS tree in a
# card) as one picture per Word page. Every slice but the last ends inside the card: only the
# card's side borders (or the bottom line of a box closing right at the slice edge) reach the
# picture's bottom edge, and the NEXT picture goes on with the same section. The checker read
# each of those as 'picture_truncated' - 206 of 206 flags on the GBT rich Studio .docx, the
# pages verified as PNG: nothing missing. A slice cut through a text line, or the LAST slice
# of a section ending open, is still flagged.
def _card_slice(h_px, edge='sides', w_px=900, lines=10):
    """A page-sized slice of a card: top border, side borders running past the bottom edge,
    ``lines`` text rows; ``edge``: 'sides' (only the side borders reach the edge), 'box' (a
    band's bottom line on the last pixel row - the band closes there) or 'text' (a text line
    cut by the edge)."""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=w_px, height=h_px)
    grey = (0.75, 0.78, 0.82)
    page.draw_line((20, 8), (w_px - 20, 8), color=grey, width=1.5)
    for x in (20, w_px - 20):
        page.draw_line((x, 8), (x, h_px + 6), color=grey, width=1.5)
    y = 34
    for i in range(lines):
        page.insert_text((36, y), 'CONS.TR28.32.SS.%04d   Install steel beams   Added   Tower %d' % (1100 + i, i),
                         fontsize=15)
        y += 22
    if edge == 'box':                      # a tree band in the right column closes at the edge
        band = pymupdf.Rect(470, h_px - 40, w_px - 40, h_px - 0.75)
        page.draw_rect(band, color=(0.3, 0.3, 0.3), width=1.5)
        page.insert_text((486, h_px - 16), 'Tower 33 - Machine Tower', fontsize=13)
    elif edge == 'text':
        page.insert_text((36, h_px + 7), 'CONS.TR28.32.SS.1170   Phase D > Delivery Bins > Chain', fontsize=15)
    return page.get_pixmap(alpha=False).tobytes('png')


def _docx_sections(sections, path):
    """A .docx with one Heading 1 per section and its pictures (png list) under it."""
    import io
    from docx import Document
    from docx.shared import Inches
    d = Document()
    for i, pics in enumerate(sections, 1):
        d.add_heading('%d Section' % i, level=1)
        for png in pics:
            d.add_picture(io.BytesIO(png), width=Inches(6.3))
    d.save(path)
    return path


def _truncated(path):
    from p6_export import pagination_check as pc
    return sorted(f['detail'].split(':')[0] for f in pc.docx_picture_flags(path)
                  if f['type'] == 'picture_truncated')


def test_word_slice_ending_inside_a_card_that_the_next_slice_continues_is_not_cut(tmp_path):
    path = _docx_sections([
        [_card_slice(300, 'sides'), _card_slice(330, 'box', lines=6), _card_slice_closed()],  # continues
        [_card_slice(300, 'sides')],                                   # the last slice ends open
        [_card_slice(300, 'text', lines=8), _card_slice_closed()],     # a text line cut by the edge
    ], str(tmp_path / 'slices.docx'))
    assert _truncated(path) == ["picture 4 under '2 Section'", "picture 5 under '3 Section'"]


def _card_slice_closed(h_px=300, w_px=900, lines=6):
    """The last slice of a card: the box closes with white below it (a clean end)."""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=w_px, height=h_px)
    grey = (0.75, 0.78, 0.82)
    bottom = 34 + lines * 22
    page.draw_rect(pymupdf.Rect(20, -4, w_px - 20, bottom), color=grey, width=1.5)
    y = 20
    for i in range(lines):
        page.insert_text((36, y), 'CONS.TR33.SS.%04d   Machine tower   Moved' % (1200 + i), fontsize=15)
        y += 22
    return page.get_pixmap(alpha=False).tobytes('png')


def test_word_slice_rule_pre_fix_flagged_every_continued_slice(tmp_path, monkeypatch):
    """Without the rule (STUDIO-RICH-9 reverted) the continued slices are 'cut' as well."""
    from p6_export import pagination_check as pc
    path = _docx_sections([[_card_slice(300, 'sides'), _card_slice(330, 'box', lines=6),
                            _card_slice_closed()]], str(tmp_path / 'pre.docx'))
    assert _truncated(path) == []
    monkeypatch.setattr(pc, '_frame_sides_only', lambda *a, **k: False)
    assert _truncated(path) == ["picture 1 under '1 Section'", "picture 2 under '1 Section'"]


# ── a Word layout that times out stops the Word IT started, never another Word ────────────
def test_word_layout_timeout_stops_only_its_own_word(tmp_path, monkeypatch):
    import subprocess
    from p6_export import pagination_check as pc
    calls = []

    def fake_run(cmd, **kw):
        if '--word-layout' in cmd:
            with open(cmd[-1], 'w', encoding='ascii') as fh:    # the child records its Word
                fh.write('4242')
            raise subprocess.TimeoutExpired(cmd, kw.get('timeout'))
        calls.append(cmd)
        return subprocess.CompletedProcess(cmd, 0)

    monkeypatch.setattr(pc.subprocess, 'run', fake_run)
    monkeypatch.setattr(pc.os, 'name', 'nt')
    import pytest
    with pytest.raises(subprocess.TimeoutExpired):
        pc._word_layout_subprocess(str(tmp_path / 'x.docx'), 1)
    assert calls == [['taskkill', '/PID', '4242', '/F']]
    # no pid recorded (Word never started) -> nothing is stopped
    calls.clear()
    assert pc._stop_pid_in(str(tmp_path / 'missing.pid')) is False and calls == []
