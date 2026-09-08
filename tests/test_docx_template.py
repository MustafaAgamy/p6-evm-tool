"""Tests for p6_narrative.docx_template — the reusable Word template furniture (SLICE B).

Verifies page geometry, the double page border, the 3-cell logo header, the
PAGE-only footer, base styles, the numbered-heading / date helpers and the TOC field.
"""
import base64

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from p6_narrative import docx_template as dt


# 1x1 transparent PNG as a data URL, for the logo header.
_PNG = (
    'data:image/png;base64,'
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk'
    'YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=='
)


def _sectPr(section):
    return section._sectPr


def _approx(a, b, tol=1000):
    """EMU comparison tolerant of python-docx's twip round-tripping."""
    return abs(int(a) - int(b)) <= tol


# ── page geometry ─────────────────────────────────────────────────────────────
def test_apply_page_geometry_a4_and_margins():
    doc = Document()
    section = doc.sections[0]
    dt.apply_page_geometry(section)
    assert _approx(section.page_width, Cm(21.0))
    assert _approx(section.page_height, Cm(29.7))
    assert _approx(section.top_margin, Cm(2.15))
    assert _approx(section.left_margin, Cm(1.34))
    assert _approx(section.right_margin, Cm(1.06))
    assert _approx(section.bottom_margin, Cm(1.69))
    assert _approx(section.header_distance, Cm(1.0))
    assert _approx(section.footer_distance, Cm(1.0))


# ── page border ───────────────────────────────────────────────────────────────
def test_add_page_border_double_sz4():
    doc = Document()
    section = doc.sections[0]
    dt.apply_page_geometry(section)
    dt.add_page_border(section)
    borders = _sectPr(section).findall(qn('w:pgBorders'))
    assert len(borders) == 1
    b = borders[0]
    assert b.get(qn('w:offsetFrom')) == 'page'
    assert b.get(qn('w:display')) == 'allPages'
    for edge in ('top', 'left', 'bottom', 'right'):
        el = b.find(qn('w:' + edge))
        assert el is not None, edge
        assert el.get(qn('w:val')) == 'double'
        assert el.get(qn('w:sz')) == '4'
        assert el.get(qn('w:space')) == '24'
        assert el.get(qn('w:color')) == '000000'


def test_add_page_border_idempotent_and_after_pgMar():
    doc = Document()
    section = doc.sections[0]
    dt.apply_page_geometry(section)
    dt.add_page_border(section)
    dt.add_page_border(section)  # second call must not duplicate
    sectPr = _sectPr(section)
    borders = sectPr.findall(qn('w:pgBorders'))
    assert len(borders) == 1
    # pgBorders sits immediately after pgMar in the sectPr child order.
    children = list(sectPr)
    pgMar = sectPr.find(qn('w:pgMar'))
    assert pgMar is not None
    assert children[children.index(pgMar) + 1] is borders[0]


# ── header ────────────────────────────────────────────────────────────────────
def test_add_header_three_cells_with_logos_and_title():
    doc = Document()
    dt.add_header(doc, {
        'project_name': 'Grand Museum',
        'document_title': 'Baseline Narrative',
        'logos': {'owner': _PNG, 'consultant': _PNG, 'contractor': _PNG},
    })
    header = doc.sections[0].header
    assert header.is_linked_to_previous is False
    assert len(header.tables) == 1
    cells = header.tables[0].rows[0].cells
    assert len(cells) == 3
    # Title band paragraph carries the project name and the quoted document title.
    header_text = '\n'.join(p.text for p in header.paragraphs)
    assert 'Grand Museum' in header_text
    assert '"Baseline Narrative"' in header_text


def test_add_header_missing_logos_keep_three_cells():
    doc = Document()
    dt.add_header(doc, {'project_name': 'No Logos', 'logos': {}})
    cells = doc.sections[0].header.tables[0].rows[0].cells
    assert len(cells) == 3  # empty cells, columns preserved


# ── footer ────────────────────────────────────────────────────────────────────
def test_add_footer_has_page_not_numpages():
    doc = Document()
    section = doc.sections[0]
    dt.add_footer(section)
    footer = section.footer
    assert footer.is_linked_to_previous is False
    xml = footer.paragraphs[0]._p.xml
    instrs = [e.text for e in footer.paragraphs[0]._p.iter(qn('w:instrText'))]
    joined = ' '.join(t or '' for t in instrs)
    assert 'PAGE' in joined
    assert 'NUMPAGES' not in xml
    # Font is Calibri 9pt grey.
    run = footer.paragraphs[0].runs[0]
    assert run.font.name == 'Calibri'
    assert run.font.size == Pt(9)


# ── cover ─────────────────────────────────────────────────────────────────────
def test_add_cover_full_and_graceful():
    doc = Document()
    dt.add_cover(doc, {
        'project_name': 'Grand Museum',
        'location': 'Cairo, Egypt',
        'document_title': 'Baseline Narrative',
        'revision': '3',
    })
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Grand Museum' in text
    assert 'Cairo, Egypt' in text
    assert '"Baseline Narrative"' in text
    assert 'REV. 3' in text


def test_add_cover_blank_fields_degrade():
    doc = Document()
    # Should not raise on missing/None fields.
    dt.add_cover(doc, {})
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Project' in text  # default project name
    assert 'REV.' not in text  # no revision line when absent


# ── toc ───────────────────────────────────────────────────────────────────────
def test_add_toc_real_field():
    doc = Document()
    dt.add_toc(doc)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Table of Contents' in text
    instrs = []
    for p in doc.paragraphs:
        instrs.extend(e.text for e in p._p.iter(qn('w:instrText')))
    assert any(t and t.strip().startswith('TOC') for t in instrs)


# ── base styles ───────────────────────────────────────────────────────────────
def test_apply_base_styles_normal_and_headings():
    doc = Document()
    dt.apply_base_styles(doc)
    normal = doc.styles['Normal']
    assert normal.font.name == 'Calibri'
    assert normal.font.size == Pt(11)
    h1 = doc.styles['Heading 1']
    assert h1.font.color.rgb == dt.NAVY
    # Navy is 1F4E79.
    assert str(h1.font.color.rgb) == '1F4E79'
    assert h1.font.name == 'Calibri Light'
    assert doc.styles['Heading 1'].font.size == Pt(16)
    assert doc.styles['Heading 2'].font.size == Pt(13)
    assert doc.styles['Heading 3'].font.size == Pt(11.5)


# ── numbering ─────────────────────────────────────────────────────────────────
def test_format_number_cases():
    assert dt.format_number((4,)) == '4)'
    assert dt.format_number((3, 1)) == '3.1)'
    assert dt.format_number((4, 2, 2, 'A')) == '4.2.2.A)'


# ── headings ──────────────────────────────────────────────────────────────────
def test_heading_and_subheading():
    doc = Document()
    dt.apply_base_styles(doc)
    h = dt.heading(doc, '4)', 'Sequence of Work', level=1)
    assert h.style.name == 'Heading 1'
    assert h.runs[0].text == '4) Sequence of Work'
    assert h.runs[0].font.color.rgb == dt.NAVY
    sh = dt.subheading(doc, '4.1)', 'Fronts')
    assert sh.style.name == 'Heading 2'
    assert sh.runs[0].text == '4.1) Fronts'


# ── tables ────────────────────────────────────────────────────────────────────
def test_styled_table_header_and_zebra():
    doc = Document()
    table = dt.styled_table(doc, ['A', 'B'], [['1', '2'], ['3', '4']])
    assert len(table.rows) == 3  # header + 2 body rows
    hdr_run = table.rows[0].cells[0].paragraphs[0].runs[0]
    assert hdr_run.bold is True
    assert hdr_run.font.color.rgb == dt.WHITE
    # Header cell carries the navy fill.
    tcPr = table.rows[0].cells[0]._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    assert shd is not None and shd.get(qn('w:fill')) == '1F4E79'


def test_styled_table_bold_last_row():
    doc = Document()
    table = dt.styled_table(doc, ['A'], [['x'], ['total']], bold_last_row=True)
    assert table.rows[-1].cells[0].paragraphs[0].runs[0].bold is True


# ── dates ─────────────────────────────────────────────────────────────────────
def test_full_date():
    assert dt.full_date('2026-05-29') == '29 May 2026'
    assert dt.full_date('') == ''
    assert dt.full_date('not a date') == 'not a date'
