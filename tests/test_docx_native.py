"""SLICE A — NATIVE, EDITABLE Word charts (p6_narrative.docx_native).

These build REAL Word chart objects (c:chartSpace parts wired into the package
with the correct content-type override + relationship + an inline <w:drawing>),
NOT flattened PNG pictures. The tests save a Document, reopen it, unzip the
package and assert the OOXML chart parts, content-type overrides and document
relationships are all present — proof Word will open them as editable charts.

Everything must be None-safe: a bad/empty payload returns None and never raises,
so a chart failure can never crash the export.
"""
import os
import zipfile

import pytest
from docx import Document

from p6_narrative import docx_native as dn


# ── helpers ───────────────────────────────────────────────────────────────────
def _save_reopen(document, tmp_path, name='out.docx'):
    out = os.path.join(str(tmp_path), name)
    document.save(out)
    Document(out)          # must reopen without raising (well-formed package)
    return out


def _names(path):
    with zipfile.ZipFile(path) as z:
        return z.namelist()


def _read(path, member):
    with zipfile.ZipFile(path) as z:
        return z.read(member).decode('utf-8')


# ── a bar + a pie chart produce native chart parts ────────────────────────────
def test_bar_and_pie_produce_native_chart_parts(tmp_path):
    doc = Document()
    doc.add_heading('Native charts', 0)
    r1 = dn.add_bar_chart(doc, ['Jun', 'Jul', 'Aug', 'Sep'],
                          [4.2, 6.8, 7.6, 9.1], 'Planned cost per month')
    r2 = dn.add_pie_chart(doc, ['Civil', 'Mechanical', 'Electrical'],
                          [56, 28, 16], 'Contract value by branch')
    assert r1 is not None and r2 is not None

    out = _save_reopen(doc, tmp_path)
    names = _names(out)

    chart_parts = [n for n in names if n.startswith('word/charts/chart') and n.endswith('.xml')]
    assert len(chart_parts) >= 2, chart_parts

    # [Content_Types].xml must carry the chart+xml override for BOTH parts
    ct = _read(out, '[Content_Types].xml')
    assert ct.count('drawingml.chart+xml') >= 2, ct

    # document relationships must reference the chart parts
    rels = _read(out, 'word/_rels/document.xml.rels')
    assert rels.count('/relationships/chart') >= 2, rels
    assert 'charts/chart' in rels


# ── the bar chart is a native COLUMN chart (barDir=col), not a picture ─────────
def test_bar_chart_is_native_column_chart(tmp_path):
    doc = Document()
    dn.add_bar_chart(doc, ['A', 'B', 'C'], [1, 2, 3], 'T', color='2E75B6')
    out = _save_reopen(doc, tmp_path)
    part = next(n for n in _names(out) if n.startswith('word/charts/chart'))
    xml = _read(out, part)
    assert '<c:barChart>' in xml
    assert '<c:barDir val="col"/>' in xml
    # the value cache carries the data points (native, editable)
    assert xml.count('<c:pt ') >= 3
    # series colour kept via solidFill (bars keep the palette)
    assert '2E75B6' in xml
    # NOT a rasterised picture
    assert '<a:blip' not in xml


# ── the pie chart is a native DOUGHNUT with per-slice palette ─────────────────
def test_pie_chart_is_native_doughnut_with_palette(tmp_path):
    doc = Document()
    dn.add_pie_chart(doc, ['x', 'y', 'z'], [10, 20, 30], 'Donut')
    out = _save_reopen(doc, tmp_path)
    part = next(n for n in _names(out) if n.startswith('word/charts/chart'))
    xml = _read(out, part)
    assert '<c:doughnutChart>' in xml
    assert '<c:holeSize' in xml
    # per-point dPt solidFill so slices keep their colours
    assert xml.count('<c:dPt>') >= 3
    # default palette applied
    assert '1F5FA8' in xml


# ── stacked bar: two native series (green working, red non-working) ───────────
def test_stacked_bar_two_series(tmp_path):
    doc = Document()
    r = dn.add_bar_chart_stacked(
        doc, ['W1', 'W2', 'W3'],
        [{'name': 'Working', 'values': [5, 6, 4], 'color': '22C55E'},
         {'name': 'Non-working', 'values': [2, 1, 3], 'color': 'EF4444'}],
        'Net working vs non-working days')
    assert r is not None
    out = _save_reopen(doc, tmp_path)
    part = next(n for n in _names(out) if n.startswith('word/charts/chart'))
    xml = _read(out, part)
    assert '<c:barChart>' in xml
    assert '<c:grouping val="stacked"/>' in xml
    assert xml.count('<c:ser>') == 2
    assert '22C55E' in xml and 'EF4444' in xml


# ── multiple charts in one document get distinct part names + rels ────────────
def test_multiple_charts_distinct_parts(tmp_path):
    doc = Document()
    dn.add_bar_chart(doc, ['a'], [1], 'one')
    dn.add_bar_chart(doc, ['b'], [2], 'two')
    dn.add_pie_chart(doc, ['c', 'd'], [1, 2], 'three')
    out = _save_reopen(doc, tmp_path)
    charts = [n for n in _names(out) if n.startswith('word/charts/chart') and n.endswith('.xml')]
    assert len(charts) == 3
    assert len(set(charts)) == 3          # unique part names


# ── embedded workbook gives "Edit Data" (best-effort, must not break export) ──
def test_bar_chart_embeds_editable_workbook(tmp_path):
    doc = Document()
    dn.add_bar_chart(doc, ['Jun', 'Jul'], [3, 4], 'Cash flow')
    out = _save_reopen(doc, tmp_path)
    names = _names(out)
    # embedded workbook is optional but expected; if present it is a real xlsx
    embeds = [n for n in names if n.startswith('word/embeddings/') and n.endswith('.xlsx')]
    if embeds:
        # the chart references it via externalData, and the chart part has a rel
        part = next(n for n in names if n.startswith('word/charts/chart'))
        xml = _read(out, part)
        assert '<c:externalData' in xml
        chart_rels = 'word/charts/_rels/' + os.path.basename(part) + '.rels'
        assert chart_rels in names
        assert 'spreadsheetml' in _read(out, '[Content_Types].xml')


# ── None-safety: bad / empty inputs return None and never raise ────────────────
def test_none_safe_bar():
    doc = Document()
    assert dn.add_bar_chart(doc, None, None, 'x') is None
    assert dn.add_bar_chart(doc, [], [], 'x') is None
    assert dn.add_bar_chart(doc, ['a', 'b'], [1], 'mismatch len') is None
    assert dn.add_bar_chart(doc, ['a'], [None], 'nonnumeric') is None
    assert dn.add_bar_chart(None, ['a'], [1], 'no doc') is None


def test_none_safe_pie():
    doc = Document()
    assert dn.add_pie_chart(doc, None, None, 'x') is None
    assert dn.add_pie_chart(doc, [], [], 'x') is None
    assert dn.add_pie_chart(doc, ['a'], [], 'x') is None
    assert dn.add_pie_chart(None, ['a'], [1], 'no doc') is None


def test_none_safe_stacked():
    doc = Document()
    assert dn.add_bar_chart_stacked(doc, None, None, 'x') is None
    assert dn.add_bar_chart_stacked(doc, [], [], 'x') is None
    assert dn.add_bar_chart_stacked(doc, ['a'], [], 'x') is None
    assert dn.add_bar_chart_stacked(doc, ['a'], [{'name': 'S', 'values': None}], 'x') is None
    assert dn.add_bar_chart_stacked(None, ['a'], [{'name': 'S', 'values': [1]}], 'x') is None


# ── a chart still adds a drawing so it is visible inline in the doc body ───────
def test_chart_adds_inline_drawing(tmp_path):
    doc = Document()
    before = len(doc.paragraphs)
    dn.add_bar_chart(doc, ['a', 'b'], [1, 2], 'T')
    assert len(doc.paragraphs) == before + 1
    out = _save_reopen(doc, tmp_path)
    body = _read(out, 'word/document.xml')
    assert '<w:drawing' in body
    assert 'graphicData' in body
    assert '/drawingml/2006/chart' in body
