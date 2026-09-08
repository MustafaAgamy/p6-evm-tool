"""SLICE D — Word Section-5 calendar renderer tests.

Builds the narrative calendar payload with SLICE A's ``build_report`` from a real
calendars fixture, renders it with ``docx_calendar.render_calendar`` into a fresh
Document (``chrome=None`` → the monthly histogram falls back to a stats table),
saves + reopens the .docx, and asserts the five pieces landed WITHOUT the excluded
Baseline Start / Baseline Finish dates.
"""
import io
import textwrap

from docx import Document

from p6_evm.parser import parse_file
from p6_narrative import docx_calendar
from p6_narrative.report import build_report

# 5-day-week default calendar (Fri+Sat off) with a New-Year holiday (Wed 01 Jan 2025)
# and two activities spanning Jan–Mar 2025 — same shape as the SLICE A payload fixture.
_XML = textwrap.dedent('''\
<?xml version="1.0"?>
<APIBusinessObjects xmlns="http://xmlns.oracle.com/Primavera/P6/V19.12/API/BusinessObjects">
  <Calendar>
    <ObjectId>C1</ObjectId><Name>5 Days/Week</Name><Type>Global</Type><IsDefault>true</IsDefault>
    <HoursPerDay>8</HoursPerDay>
    <StandardWorkWeek>
      <StandardWorkHours><DayOfWeek>Friday</DayOfWeek></StandardWorkHours>
      <StandardWorkHours><DayOfWeek>Saturday</DayOfWeek></StandardWorkHours>
      <StandardWorkHours><DayOfWeek>Sunday</DayOfWeek><WorkTime><Start>08:00:00</Start><Finish>16:00:00</Finish></WorkTime></StandardWorkHours>
      <StandardWorkHours><DayOfWeek>Monday</DayOfWeek><WorkTime><Start>08:00:00</Start><Finish>16:00:00</Finish></WorkTime></StandardWorkHours>
      <StandardWorkHours><DayOfWeek>Tuesday</DayOfWeek><WorkTime><Start>08:00:00</Start><Finish>16:00:00</Finish></WorkTime></StandardWorkHours>
      <StandardWorkHours><DayOfWeek>Wednesday</DayOfWeek><WorkTime><Start>08:00:00</Start><Finish>16:00:00</Finish></WorkTime></StandardWorkHours>
      <StandardWorkHours><DayOfWeek>Thursday</DayOfWeek><WorkTime><Start>08:00:00</Start><Finish>16:00:00</Finish></WorkTime></StandardWorkHours>
    </StandardWorkWeek>
    <HolidayOrExceptions>
      <HolidayOrException><Date>2025-01-01T00:00:00</Date></HolidayOrException>
    </HolidayOrExceptions>
  </Calendar>
  <Project>
    <ObjectId>1</ObjectId><Id>P1</Id><Name>Calendar Project</Name>
    <DataDate>2025-02-01T00:00:00</DataDate>
    <PlannedStartDate>2025-01-01T00:00:00</PlannedStartDate>
    <ScheduledFinishDate>2025-03-31T17:00:00</ScheduledFinishDate>
    <WBS><ObjectId>10</ObjectId><Name>Construction</Name><ParentObjectId></ParentObjectId></WBS>
    <Activity><ObjectId>A1</ObjectId><Id>A1</Id><Name>a</Name><Status>Not Started</Status>
      <CalendarObjectId>C1</CalendarObjectId><WBSObjectId>10</WBSObjectId><PercentComplete>0</PercentComplete>
      <PlannedStartDate>2025-01-02T00:00:00</PlannedStartDate>
      <PlannedFinishDate>2025-02-15T00:00:00</PlannedFinishDate></Activity>
    <Activity><ObjectId>A2</ObjectId><Id>A2</Id><Name>b</Name><Status>Not Started</Status>
      <CalendarObjectId>C1</CalendarObjectId><WBSObjectId>10</WBSObjectId><PercentComplete>0</PercentComplete>
      <PlannedStartDate>2025-02-16T00:00:00</PlannedStartDate>
      <PlannedFinishDate>2025-03-31T00:00:00</PlannedFinishDate></Activity>
  </Project>
</APIBusinessObjects>
''')


def _calendar_payload(tmp_path):
    p = tmp_path / 'cal.xml'
    p.write_text(_XML, encoding='utf-8')
    data = parse_file(str(p))
    doc = build_report(data, path=str(p)).to_dict()
    for s in doc['sections']:
        if s.get('title', '').strip().lower() == 'project calendars & holidays':
            return s['payload']
    raise AssertionError('calendar section not found in report')


def _render(tmp_path, chrome=None):
    payload = _calendar_payload(tmp_path)
    document = Document()
    docx_calendar.render_calendar(document, payload, chrome)
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return Document(buf)


def _all_text(document):
    parts = [pa.text for pa in document.paragraphs]
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return '\n'.join(parts)


def _table_texts(table):
    return [cell.text for row in table.rows for cell in row.cells]


def test_renders_the_five_pieces(tmp_path):
    doc = _render(tmp_path)
    # dashboard tiles + monthly-fallback + holidays + hours + comparison + usage
    assert len(doc.tables) >= 4, f'expected >=4 tables, got {len(doc.tables)}'
    text = _all_text(doc)
    # the five sub-headings are present
    for label in ('5.1)', '5.2)', '5.3)', '5.4)', '5.5)'):
        assert label in text, f'missing sub-heading {label}'


def test_excludes_baseline_start_finish(tmp_path):
    doc = _render(tmp_path)
    text = _all_text(doc)
    assert 'Baseline Start' not in text
    assert 'Baseline Finish' not in text
    # but the live dashboard labels that ARE kept do appear
    assert 'Working Days' in text
    assert 'Forecast Finish' in text


def test_holidays_table_has_populated_weekday(tmp_path):
    doc = _render(tmp_path)
    # find the holidays table (header Date | Weekday | Reason)
    hol = next((t for t in doc.tables
                if [c.text for c in t.rows[0].cells][:3] == ['Date', 'Weekday', 'Reason']), None)
    assert hol is not None, 'holidays table not found'
    weekdays = [t for t in _table_texts(hol) if t]
    # 01 Jan 2025 is a Wednesday — the weekday column must be populated
    assert 'Wednesday' in weekdays
    assert hol.rows[1].cells[1].text.strip(), 'weekday cell should be populated'


def test_hours_table_shows_days_per_week(tmp_path):
    doc = _render(tmp_path)
    hrs = next((t for t in doc.tables
                if [c.text for c in t.rows[0].cells][:1] == ['Profile']), None)
    assert hrs is not None, 'hours table not found'
    # a Days/Week value is rendered (5-day week → '5' in that column)
    dpw_col = [row.cells[3].text.strip() for row in hrs.rows[1:]]
    assert any(v == '5' for v in dpw_col), f'days/week not shown: {dpw_col}'


def test_comparison_and_usage_rows_render(tmp_path):
    doc = _render(tmp_path)
    comp = next((t for t in doc.tables
                 if [c.text for c in t.rows[0].cells][:2] == ['Calendar', 'Hours/Day']), None)
    usage = next((t for t in doc.tables
                  if [c.text for c in t.rows[0].cells] == ['Calendar', 'Activities', '%', 'Role']), None)
    assert comp is not None, 'comparison table not found'
    assert usage is not None, 'usage table not found'
    assert len(comp.rows) >= 2, 'comparison should have at least one data row'
    assert len(usage.rows) >= 2, 'usage should have at least one data row'
    # the default calendar is flagged Yes in the comparison table
    assert 'Yes' in _table_texts(comp)
