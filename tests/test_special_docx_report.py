"""Tests for the Reporting Studio native Word (.docx) generator.

Unlike the Office-HTML ``.doc`` wrapper (:mod:`p6_special.word_export`), this exporter
writes a genuine ``.docx`` package via python-docx: a double navy page border on every
page (``w:pgBorders``), a 3-logo header part, a page-number footer, a navy cover, navy
tables, and — importantly — ``bars`` drawn as real proportional bars (not a value
table). These tests hand-build a ``rendered`` selection (as ``registry.render`` returns)
covering table, kpi_group, bars, findings and a reused-feature ``html`` section, write a
``.docx``, then open it back with python-docx / zipfile and assert the structure.
"""
import zipfile

import docx

from p6_special import payloads as P
from p6_special.docx_report import build_docx


def _meta():
    return {'project_name': 'Grain Bulk Terminal', 'data_date': '2026-07-19'}


def _rendered():
    """A small report: a data table, a KPI group, a bars result, a findings register,
    and a REUSED feature-report ``html`` section (its inner ``<td>`` must be
    reconstructed into a real docx table)."""
    return [
        {'id': 'evm:cat', 'title': 'By category', 'feature': 'evm',
         'feature_title': 'EVM Report', 'ctype': 'table',
         'payload': P.table(['Category', 'Planned %', 'Actual %'],
                            [['Piling', 100, ('95', 'good')],
                             ['Concrete', ('80', 'warn'), 72]])},
        {'id': 'evm:kpi', 'title': 'Key metrics', 'feature': 'evm',
         'feature_title': 'EVM Report', 'ctype': 'kpi',
         'payload': P.kpi_group([P.kpi('SPI', '0.87', sub='behind plan', tone='warn'),
                                 P.kpi('CPI', '1.02', tone='good')])},
        {'id': 'evm:bars', 'title': 'Planned vs Actual', 'feature': 'evm',
         'feature_title': 'EVM Report', 'ctype': 'chart',
         'payload': P.bars(rows=[{'label': 'Piling', 'values': [100, 95],
                                  'display': ['100%', '95%']},
                                 {'label': 'Concrete', 'values': [80, 72],
                                  'display': ['80%', '72%']}],
                           series=[{'label': 'Planned', 'tone': 'neutral'},
                                   {'label': 'Actual', 'tone': 'good'}],
                           axis_max=100, note='Weighted by BAC')},
        {'id': 'audit:find', 'title': 'Open findings', 'feature': 'audit',
         'feature_title': 'Schedule Audit', 'ctype': 'findings',
         'payload': P.findings([{'severity': 'high', 'title': 'Negative float',
                                 'detail': '12 activities below zero'},
                                {'severity': 'low', 'title': 'Long duration',
                                 'detail': '3 activities over 60d'}])},
        {'id': 'audit:reuse', 'title': 'Detailed section', 'feature': 'audit',
         'feature_title': 'Schedule Audit', 'ctype': 'html',
         'payload': {'kind': 'html',
                     'css': '.x{color:red}',
                     'html': '<div><h3>Cat</h3>'
                             '<table><tr><th>A</th></tr><tr><td>1</td></tr></table></div>'}},
    ]


def _letterhead():
    return {'kicker': 'Project Progress Report', 'company': 'Acme Marine JV',
            'prepared_by': 'Planning Dept',
            'logos': {'owner': None, 'consultant': None, 'contractor': None}}


def _all_table_cell_texts(document):
    out = []
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                out.append(cell.text)
                # nested tables inside a cell, if any
                for nt in cell.tables:
                    for r in nt.rows:
                        for c in r.cells:
                            out.append(c.text)
    return out


def _all_paragraph_texts(document):
    return [p.text for p in document.paragraphs]


def test_build_docx_writes_valid_docx(tmp_path):
    out = tmp_path / 'out.docx'
    ret = build_docx(out, 'Weekly', _meta(), _rendered(), letterhead=_letterhead())

    # written + a real zip (docx is a zip package)
    assert out.exists()
    assert str(ret) == str(out)
    assert zipfile.is_zipfile(out)

    document = docx.Document(str(out))
    paras = _all_paragraph_texts(document)
    cell_texts = _all_table_cell_texts(document)

    # cover carries the report name
    assert any('Weekly' in p for p in paras)
    # a Contents heading is present
    assert any(p.strip() == 'Contents' for p in paras)
    # the numbered section headings (pick order)
    assert any(p.strip() == '1 By category' for p in paras)
    assert any(p.strip() == '3 Planned vs Actual' for p in paras)
    assert any(p.strip() == '5 Detailed section' for p in paras)

    # >= 1 real table with the expected header cell(s)
    assert 'Category' in cell_texts and 'Planned %' in cell_texts
    assert 'Metric' in cell_texts and 'Value' in cell_texts     # kpi_group table
    assert 'Severity' in cell_texts and 'Finding' in cell_texts  # findings table
    # findings severity values are title-cased
    assert 'High' in cell_texts and 'Low' in cell_texts

    # the html section's inner <td> '1' appears in a reconstructed table (chrome=None
    # here, so the html section falls back to native text/table extraction)
    assert '1' in cell_texts
    # the section's OWN leading heading is stripped (FIX 3) — 'Cat' is NOT re-emitted
    # under the Studio's '5 Detailed section' heading
    assert not any(p.strip() == 'Cat' for p in paras)


def test_docx_has_navy_page_border_and_header_part(tmp_path):
    out = tmp_path / 'framed.docx'
    build_docx(out, 'Weekly', _meta(), _rendered(), letterhead=_letterhead())

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        # a page border is declared on the section
        document_xml = z.read('word/document.xml').decode('utf-8')
        assert 'w:pgBorders' in document_xml
        # navy frame colour
        assert '1F4E79' in document_xml.upper()
        # a header part exists (the 3-logo running header)
        assert 'word/header1.xml' in names
        # a footer part exists (the page-number footer)
        assert any(n.startswith('word/footer') for n in names)
        # a live PAGE field in the footer
        footer = next(n for n in names if n.startswith('word/footer'))
        assert 'PAGE' in z.read(footer).decode('utf-8')


def test_bars_drawn_as_shaded_bars_not_a_table(tmp_path):
    """A bars payload draws proportional shaded cells (a series colour + a grey track),
    NOT a numeric value table — so the document XML carries the fill colours."""
    out = tmp_path / 'bars.docx'
    build_docx(out, 'Weekly', _meta(),
               [{'id': 'b', 'title': 'Bars only', 'feature': 'f', 'feature_title': 'Feat',
                 'ctype': 'chart',
                 'payload': P.bars(rows=[{'label': 'A', 'values': [100, 50]}],
                                   series=[{'label': 'Planned', 'tone': 'neutral'},
                                           {'label': 'Actual', 'tone': 'good'}],
                                   axis_max=100)}],
               letterhead=_letterhead())

    with zipfile.ZipFile(out) as z:
        document_xml = z.read('word/document.xml').decode('utf-8').upper()
    # the navy first-series fill, the green second-series fill and the grey track
    assert '1F4E79' in document_xml            # navy (series 0 + chrome)
    assert '2E7D32' in document_xml            # green (series 1)
    assert 'E9EDF2' in document_xml            # grey bar track


def test_build_docx_empty_selection_still_valid(tmp_path):
    out = tmp_path / 'empty.docx'
    build_docx(out, 'Weekly', _meta(), [], letterhead=_letterhead())
    assert out.exists()
    assert zipfile.is_zipfile(out)
    document = docx.Document(str(out))
    paras = _all_paragraph_texts(document)
    assert any('Weekly' in p for p in paras)                 # cover still written
    assert any(p.strip() == 'Contents' for p in paras)


def test_build_docx_never_raises_on_malformed_items(tmp_path):
    """A None payload, an unknown kind and a no_data item are skipped, never fatal."""
    rendered = [
        {'id': 'ok', 'title': 'Good table', 'feature': 'f', 'feature_title': 'Feat',
         'ctype': 'table', 'payload': P.table(['A', 'B'], [['x', 'y']])},
        {'id': 'none', 'title': 'Broken', 'feature': 'f', 'feature_title': 'Feat',
         'ctype': 'text', 'payload': None},
        {'id': 'weird', 'title': 'Mystery', 'feature': 'f', 'feature_title': 'Feat',
         'ctype': 'text', 'payload': {'kind': 'zzz'}},
        {'id': 'nd', 'title': 'Nothing', 'feature': 'f', 'feature_title': 'Feat',
         'ctype': 'text', 'payload': P.NO_DATA},
    ]
    out = tmp_path / 'malformed.docx'
    build_docx(out, 'Weekly', _meta(), rendered, letterhead=_letterhead())
    assert zipfile.is_zipfile(out)
    document = docx.Document(str(out))
    paras = _all_paragraph_texts(document)
    # every picked item still gets its numbered heading (even the empty ones)
    for n, title in ((1, 'Good table'), (2, 'Broken'), (3, 'Mystery'), (4, 'Nothing')):
        assert any(p.strip() == ('%d %s' % (n, title)) for p in paras)


def test_group_and_segbar_and_line_and_status(tmp_path):
    """A group recurses its blocks; segbar/line/status_header each render without error."""
    rendered = [
        {'id': 'g', 'title': 'Grouped', 'feature': 'f', 'feature_title': 'Feat',
         'ctype': 'group',
         'payload': P.group([P.keyvals([('Data date', '2026-07-19')]),
                             P.segbar([{'label': 'Steel', 'value': 60, 'tone': 'accent'},
                                       {'label': 'Concrete', 'value': 40, 'tone': 'good'}])])},
        {'id': 'l', 'title': 'SPI trend', 'feature': 'f', 'feature_title': 'Feat',
         'ctype': 'chart',
         'payload': P.line([{'label': 'SPI', 'tone': 'accent', 'points': [0.9, 0.92, 0.87]}],
                           x=['W1', 'W2', 'W3'], ref={'value': 1.0, 'label': 'target'})},
        {'id': 's', 'title': 'Status', 'feature': 'f', 'feature_title': 'Feat',
         'ctype': 'summary',
         'payload': P.status_header([{'domain': 'Schedule', 'tone': 'warn',
                                      'headline': 'Slipping'}])},
    ]
    out = tmp_path / 'mixed.docx'
    build_docx(out, 'Weekly', _meta(), rendered, letterhead=_letterhead())
    assert zipfile.is_zipfile(out)
    document = docx.Document(str(out))
    cell_texts = _all_table_cell_texts(document)
    # keyvals from inside the group
    assert 'Data date' in cell_texts
    # line trend point table + status header table
    assert 'Point' in cell_texts and 'SPI' in cell_texts
    assert 'Area' in cell_texts and 'Schedule' in cell_texts


# ── FIX 1 — contents page carries live page numbers (bookmarks + PAGEREF) ──────
def test_contents_has_pageref_page_numbers_and_bookmarks(tmp_path):
    """Each contents entry ends with a live PAGEREF field pointing at a section-heading
    bookmark, and the settings part asks Word to refresh fields on open."""
    out = tmp_path / 'toc.docx'
    build_docx(out, 'Weekly', _meta(), _rendered(), letterhead=_letterhead())

    with zipfile.ZipFile(out) as z:
        document_xml = z.read('word/document.xml').decode('utf-8')
        settings_xml = z.read('word/settings.xml').decode('utf-8')

    # PAGEREF fields exist in the contents list (one per selected item)
    assert 'PAGEREF' in document_xml
    assert document_xml.count('PAGEREF') >= len(_rendered())
    # bookmarks wrap the numbered section headings (sec1 … secN)
    assert 'w:bookmarkStart' in document_xml
    assert 'w:bookmarkEnd' in document_xml
    for i in range(1, len(_rendered()) + 1):
        assert ('sec%d' % i) in document_xml
    # Word refreshes fields (the page numbers) on open
    assert 'w:updateFields' in settings_xml


# ── FIX 3 — every numbered section names its source feature ────────────────────
def test_each_section_has_feature_caption(tmp_path):
    """Directly under each numbered section heading sits a caption naming the feature
    the result came from (matching the contents-page feature tag)."""
    out = tmp_path / 'cap.docx'
    build_docx(out, 'Weekly', _meta(), _rendered(), letterhead=_letterhead())
    paras = [p.text.strip() for p in docx.Document(str(out)).paragraphs]

    def _caption_after(head, cap):
        for i, t in enumerate(paras):
            if t == head:
                for t2 in paras[i + 1:]:
                    if t2:                      # first non-empty paragraph after heading
                        return t2 == cap
        return False

    assert _caption_after('1 By category', 'EVM Report')
    assert _caption_after('4 Open findings', 'Schedule Audit')
    assert _caption_after('5 Detailed section', 'Schedule Audit')


# ── FIX 3 — a reused section's own leading heading is stripped ─────────────────
def test_html_section_leading_heading_is_stripped(tmp_path):
    """A reused 'html' section whose fragment opens with its own heading has that
    heading removed, so the Studio's numbered heading is not followed by a duplicate
    title — while the section's body table still comes through (chrome=None fallback)."""
    rendered = [
        {'id': 'r', 'title': 'Executive dashboard', 'feature': 'calendar',
         'feature_title': 'Calendar & Weather', 'ctype': 'html',
         'payload': {'kind': 'html', 'css': '.srf-calendar h2{color:navy}',
                     'html': '<div class="srf-calendar">'
                             '<h2>1 - Execution Dashboard</h2>'
                             '<table><tr><th>Metric</th></tr>'
                             '<tr><td>innercell</td></tr></table></div>'}},
    ]
    out = tmp_path / 'strip.docx'
    build_docx(out, 'Weekly', _meta(), rendered, letterhead=_letterhead())
    document = docx.Document(str(out))
    paras = [p.text.strip() for p in document.paragraphs]
    cell_texts = _all_table_cell_texts(document)

    # Studio heading present, section's own duplicate heading gone
    assert any(p == '1 Executive dashboard' for p in paras)
    assert not any('Execution Dashboard' in p for p in paras)
    # body survives the strip
    assert 'innercell' in cell_texts


# ── FIX 2 — with no chrome, html sections fall back to text/table extraction ───
def test_html_falls_back_to_extraction_without_chrome(tmp_path):
    """With ``chrome=None`` there is no image; the reused html section is reconstructed
    as native docx tables so its cell text still appears (no crash)."""
    rendered = [
        {'id': 'r', 'title': 'Detailed', 'feature': 'audit',
         'feature_title': 'Schedule Audit', 'ctype': 'html',
         'payload': {'kind': 'html', 'css': '.srf-audit td{padding:2px}',
                     'html': '<div class="srf-audit"><table>'
                             '<tr><th>Head</th></tr><tr><td>innercell</td></tr>'
                             '</table></div>'}},
    ]
    out = tmp_path / 'fallback.docx'
    build_docx(out, 'Weekly', _meta(), rendered, letterhead=_letterhead(), chrome=None)
    assert zipfile.is_zipfile(out)
    cell_texts = _all_table_cell_texts(docx.Document(str(out)))
    assert 'innercell' in cell_texts


def test_rasterize_section_returns_none_without_chrome():
    """The rasteriser cleanly returns None when no chrome is given, or the path is
    missing — so the caller always has a safe fallback (never an exception)."""
    from p6_special.docx_report import _rasterize_section
    assert _rasterize_section('<p>x</p>', '.x{}', 'light', None) is None
    assert _rasterize_section('<p>x</p>', '.x{}', 'light', 'C:/nope/chrome-does-not-exist.exe') is None


def test_section_doc_does_not_nest_style_tags():
    """Regression: the theme block (a full ``<style>…</style>`` element) must NOT be
    nested inside another ``<style>``. Nesting let the theme's ``</style>`` close the
    block early, spilling the section's scoped CSS onto the page as visible text —
    which the rasteriser then screenshotted (Word showed a wall of CSS instead of the
    chart). The generated page must be well-formed: the scoped selector lives inside a
    style block in the head, never in the visible body, and no <style> opens while
    another is still open."""
    from p6_special.docx_report import _section_doc
    css = ".srf-weather h2{color:var(--rpt-accent)} .srf-weather td{padding:2px}"
    frag = '<div class="srf-weather"><table><tr><td>cell</td></tr></table></div>'
    doc = _section_doc(frag, css, 'light')
    head, _, body = doc.partition('</head>')
    assert '.srf-weather' in head          # the scoped CSS made it into a style block
    assert '.srf-weather h2' not in body   # ... and did NOT leak into the visible body
    # no <style> opens while another is still open (no nesting)
    depth = 0
    i = 0
    low = doc.lower()
    while i < len(low):
        if low.startswith('<style', i):
            depth += 1
            assert depth == 1, 'nested <style> — theme tag was wrapped again'
            i += 6
        elif low.startswith('</style>', i):
            depth -= 1
            i += 8
        else:
            i += 1
    assert depth == 0
