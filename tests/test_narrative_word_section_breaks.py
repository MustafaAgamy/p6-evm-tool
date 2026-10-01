"""Baseline Narrative Word — every section opens a new page through its heading's
'page break before', never through a separate page-break paragraph (NARRFIX).

The writer used to add a page-break paragraph after every section. When a section filled
its last page to the bottom margin, that paragraph did not fit, spilled onto the next page
and broke THERE — a page with nothing on it but its page number (SG_BASELINE_XER Word p12
after §7 Scope of Work). The page checker counted the text-less break paragraph as content,
so the empty page was reported only as an informational 'section_break_blank'.

XML / checker tests need python-docx only; the Word proof is skipped without Microsoft Word.
"""
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from p6_export import pagination_check as pc
from p6_narrative import docx_writer as W

_P = qn('w:p')
_FILLS = list(range(50, 63))          # 12 pt lines under a section heading: one of them fills
                                      # the page to within less than a break paragraph's line


def _pbb(p):
    ppr = p.find(qn('w:pPr'))
    el = ppr.find(qn('w:pageBreakBefore')) if ppr is not None else None
    return el is not None and el.get(qn('w:val'), 'true') not in ('0', 'false', 'off')


def _break_only(p):
    txt = ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip()
    return not txt and any(b.get(qn('w:type')) == 'page' for b in p.iter(qn('w:br')))


def _write(tmp_path, monkeypatch, name='new.docx'):
    """The narrative writer on sections whose body is ``fill`` lines of exactly 12 pt."""
    def render(document, section, number):
        from docx.enum.text import WD_LINE_SPACING
        from docx.shared import Pt
        p = document.add_paragraph()                  # ONE paragraph of ``fill`` lines
        for i in range(section['fill']):
            r = p.add_run('Section %s filler line %d.' % (number, i + 1))
            if i + 1 < section['fill']:
                r.add_break()
        f = p.paragraph_format
        f.line_spacing_rule, f.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(12)
        f.space_before = f.space_after = Pt(0)
        f.widow_control = False
    monkeypatch.setattr(W, '_render', render)
    doc = {'meta': {'project_name': 'Break test'},
           'sections': [{'number': k + 1, 'title': 'Fill %d' % n, 'kind': 'text', 'fill': n}
                        for k, n in enumerate(_FILLS)]}
    out = str(tmp_path / name)
    W.write_docx(doc, out)
    return out


def _headings(d):
    return [p for p in d.element.body.iter(_P)
            if ''.join(t.text or '' for t in p.iter(qn('w:t'))).startswith(tuple(
                '%d) Fill' % (k + 1) for k in range(len(_FILLS))))]


def test_sections_open_a_page_by_page_break_before_not_a_break_paragraph(tmp_path, monkeypatch):
    d = Document(_write(tmp_path, monkeypatch))
    heads = _headings(d)
    assert len(heads) == len(_FILLS)
    assert all(_pbb(h) for h in heads)                    # every section starts a new page …
    body = [el for el in d.element.body if el.tag == _P]
    first = body.index(heads[0])
    assert not any(_break_only(p) for p in body[first - 1:])   # … with no break paragraph
    # … and still keeps with its content (the shared pass treats 'page break before' as
    # starting the heading's page, not ending one)
    assert all(h.find(qn('w:pPr')).find(qn('w:keepNext')) is not None for h in heads)


def test_checker_flags_a_page_holding_only_a_spilled_break_paragraph():
    lay = {'page': {'h': 842.0, 'w': 595.0, 'top': 61.0, 'bottom': 48.0}, 'pages': 3, 'items': [
        {'k': 'p', 'page': 1, 'page_end': 1, 'y': 95.0, 'y_last': 780.0, 't': 'Section 7 text', 'kwn': False,
         'brk': False, 'len': 900, 'size': 11, 'bold': False, 'style': 'Normal', 'ol': 10},
        {'k': 'p', 'page': 2, 'page_end': 2, 'y': 61.0, 'y_last': 61.0, 't': '', 'brk': True, 'len': 0,
         'size': 11, 'bold': False, 'style': 'Normal', 'ol': 10},
        {'k': 'p', 'page': 3, 'page_end': 3, 'y': 95.0, 'y_last': 95.0, 't': '8) Project WBS', 'kwn': True,
         'brk': False, 'len': 14, 'size': 16, 'bold': True, 'style': 'Heading 1', 'ol': 1}]}
    flags, _, _ = pc.analyze_word_layout(lay)
    assert [(f['type'], f['page']) for f in flags] == [('empty_page', 2)]


def _old_breaks(src, dst):
    """The pre-fix writer's layout: a page-break paragraph after each section instead of
    'page break before' on the next section's heading."""
    d = Document(src)
    heads = _headings(d)
    for h in heads[1:]:
        h.find(qn('w:pPr')).remove(h.find(qn('w:pPr')).find(qn('w:pageBreakBefore')))
        p = OxmlElement('w:p')
        r = OxmlElement('w:r')
        br = OxmlElement('w:br')
        br.set(qn('w:type'), 'page')
        r.append(br)
        p.append(r)
        h.addprevious(p)
    d.save(dst)


def test_word_a_section_filling_its_page_leaves_no_empty_page(tmp_path, monkeypatch):
    """Word proof: sections of 50…62 lines — one of them fills its page to within less than
    a line of the bottom margin. With a break paragraph after each section that one leaves
    an EMPTY page (flagged); with 'page break before' on the headings there is none."""
    if not pc.word_available():
        pytest.skip('Microsoft Word (COM) is not available')
    new = _write(tmp_path, monkeypatch)
    old = str(tmp_path / 'old.docx')
    _old_breaks(new, old)
    old_flags, _, _ = pc.analyze_word_layout(pc.word_layout(old))
    assert 'empty_page' in {f['type'] for f in old_flags}, old_flags
    lay = pc.word_layout(new)
    new_flags, _, _ = pc.analyze_word_layout(lay)
    assert 'empty_page' not in {f['type'] for f in new_flags}, new_flags
    # every page holds text: no page is left with only its page number
    pages = set()
    for it in lay['items']:
        if (it.get('t') or '').strip() or it['k'] == 'table':
            pages.update(range(it['page'], it['page_end'] + 1))
    assert pages >= set(range(3, lay['pages'] + 1))
