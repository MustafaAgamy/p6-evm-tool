"""Owner point 14 — the ONE shared page-composition layer (PDF/HTML: report_theme; Word:
p6_export.docx_pagination). Unit tests for both, plus a Chrome proof on a synthetic long
report: the same document breaks the rules WITHOUT the layer and keeps them WITH it."""
import io
import os
import re
import tempfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

import report_theme as rt
from p6_export import docx_pagination as DP


# ── PDF / HTML layer ────────────────────────────────────────────────────────────
def test_every_themed_report_carries_the_pagination_layer_once():
    tag = rt.theme_style_tag('dark')
    assert tag.startswith('<style id="rpt-theme"')
    assert tag.count('id="rpt-pagination"') == 1 and tag.count('id="rpt-pagination-js"') == 1
    doc = f'<html><head>{tag}</head><body></body></html>'
    light = rt.force_light(doc)                 # Word / Excel re-theme keeps ONE layer
    assert light.count('id="rpt-pagination"') == 1
    assert 'data-rpt-theme="light"' in light and 'data-rpt-theme="dark"' not in light


def test_rules_cover_headings_parts_tables_and_scroll_boxes():
    css = rt.pagination_css()
    assert css.lstrip().startswith('@media print')          # the screen is never touched
    for needle in ('h2,', 'div.sub', '.ct', '.sr-sec-h', 'break-after: avoid',
                   'thead { display: table-header-group; }', 'tr { break-inside: avoid',
                   'tbody > tr:nth-child(-n+3)', 'tbody > tr:nth-last-child(-n+2)',
                   '.rpt-flow { break-inside: auto !important', 'orphans: 3',
                   '[style*="overflow-x"]', 'svg,', '.mgrid-wrap'):
        assert needle in css, needle
    assert ':is(' not in css                    # safe for the Studio scoper / Word CSS engine
    js = rt.pagination_script()
    assert '</script' not in js.lower()
    assert "addEventListener('beforeprint'" in js and "addEventListener('afterprint'" in js


def test_with_pagination_is_idempotent_and_only_defaults_the_page_size():
    base = '<html><head><style>@page { margin: 20mm 14mm; }</style></head><body>x</body></html>'
    once = rt.with_pagination(base)
    assert once.count('id="rpt-pagination"') == 1
    assert '@page { size: A4 portrait; }' in once
    assert once.index('rpt-page-size') < once.index('margin: 20mm 14mm')   # report margins win
    assert rt.with_pagination(once) == once
    land = '<html><head><style>@page { size: A4 landscape; margin: 11mm; }</style></head><body></body></html>'
    assert 'rpt-page-size' not in rt.with_pagination(land)                  # never overrides
    frag = rt.with_pagination('<div>fragment</div>')
    assert 'id="rpt-pagination"' in frag and frag.endswith('<div>fragment</div>')


def test_rules_survive_the_studio_css_scoper_and_the_word_css_engine():
    from p6_special.reuse import scope_css
    scoped = scope_css(rt.pagination_css(), '.x')
    assert '.x tbody > tr:nth-child(-n+3)' in scoped and '.x h2' in scoped
    from p6_export import css as C
    C.StyleSheet([rt.pagination_css()], page_width_px=700)   # parses without raising


def test_keep_together_wraps_a_group():
    assert rt.keep_together('<h3>A</h3>', '', '<table></table>') == \
        '<div class="rpt-keep"><h3>A</h3><table></table></div>'


def test_the_exported_html_keeps_the_rules_but_no_script():
    from p6_export.to_html import standalone_html
    out = standalone_html('<html><head></head><body><h2>A</h2></body></html>')
    assert 'id="rpt-pagination"' in out and '<script' not in out
    assert out.index('<meta charset="utf-8">') < out.index('rpt-pagination')


# ── Word layer ──────────────────────────────────────────────────────────────────
def _kwn(p_el):
    ppr = p_el.find(qn('w:pPr'))
    k = ppr.find(qn('w:keepNext')) if ppr is not None else None
    return k is not None and k.get(qn('w:val'), 'true') not in ('0', 'false')


def _row_kwn(tr):
    return all(_kwn(p) for tc in tr.findall(qn('w:tc')) for p in tc.findall(qn('w:p')))


def _flag(tr, name):
    trPr = tr.find(qn('w:trPr'))
    return trPr is not None and trPr.find(qn('w:' + name)) is not None


def _table(doc, n, header=True, text='cell'):
    t = doc.add_table(rows=n + (1 if header else 0), cols=3)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            r = c.paragraphs[0].add_run('Head' if header and i == 0 else f'{text} {i}')
            r.bold = bool(header and i == 0)
    return t


def test_word_small_table_is_kept_whole_with_its_heading_and_intro():
    doc = Document()
    h = doc.add_heading('Holidays', level=2)
    doc.add_paragraph('')                                   # a spacer between heading and intro
    intro = doc.add_paragraph('The calendar non-working days:')
    t = _table(doc, 11)
    DP.paginate_docx(doc)
    rows = t._tbl.findall(qn('w:tr'))
    assert all(_flag(tr, 'cantSplit') for tr in rows)
    assert all(_row_kwn(tr) for tr in rows[:-1]) and not _row_kwn(rows[-1])
    assert _kwn(intro._p) and _kwn(h._p)
    assert _flag(rows[0], 'tblHeader')                      # bold first row over a plain body


def test_word_long_table_repeats_its_header_and_never_strands_rows():
    doc = Document()
    doc.add_heading('Major Milestones', level=2)
    t = _table(doc, 60)
    DP.paginate_docx(doc)
    rows = t._tbl.findall(qn('w:tr'))
    assert _flag(rows[0], 'tblHeader') and not _flag(rows[1], 'tblHeader')
    # header + first 3 body rows chained; last 3 rows chained; the middle may break
    assert [_row_kwn(tr) for tr in rows[:4]] == [True, True, True, False]
    assert [_row_kwn(tr) for tr in rows[-4:]] == [False, True, True, False]
    assert not any(_row_kwn(tr) for tr in rows[5:-4])


def test_word_row_taller_than_a_page_is_allowed_to_break():
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].paragraphs[0].add_run('long text ' * 3000)
    DP.paginate_docx(doc)
    assert not _flag(t._tbl.findall(qn('w:tr'))[0], 'cantSplit')   # Word would clip it


def test_word_picture_keeps_with_its_label_and_caption():
    png = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00'
           b'\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfe\xa7\x35\x81'
           b'\x84\x00\x00\x00\x00IEND\xaeB`\x82')
    doc = Document()
    doc.add_paragraph('A long body paragraph. ' * 40)
    label = doc.add_paragraph('Roots Silos - 24 Hrs - Phase I Scope — 17 activities')
    doc.add_picture(io.BytesIO(png), width=Pt(300))
    pic = doc.paragraphs[-1]
    cap = doc.add_paragraph('Figure 3 — working-day timeline')
    DP.paginate_docx(doc)
    assert _kwn(label._p) and _kwn(pic._p) and not _kwn(cap._p)
    assert not _kwn(doc.paragraphs[0]._p)                  # a long body paragraph is not chained


def test_word_export_applies_the_rules_end_to_end(tmp_path):
    from p6_export.to_docx import html_to_docx
    rows = ''.join(f'<tr><td>d{i}</td><td>v{i}</td></tr>' for i in range(11))
    html = ('<html><head>' + rt.theme_style_tag('light') + '</head><body><section data-sec="a">'
            '<h2>Holidays</h2><p>Intro line</p><table><thead><tr><th>Date</th><th>Name</th></tr>'
            f'</thead><tbody>{rows}</tbody></table></section></body></html>')
    out = tmp_path / 'r.docx'
    html_to_docx(html, str(out), use_chrome=False)
    d = Document(str(out))
    t = d.tables[-1]._tbl.findall(qn('w:tr'))
    assert _flag(t[0], 'tblHeader') and all(_row_kwn(tr) for tr in t[:-1])


# ── Chrome proof: the same document, without and with the layer ─────────────────
_A4 = '@page { size: A4 portrait; margin: 14mm; } body { margin: 0; font: 9px Arial; }'
_ROW = 'td, th { height: 6mm; padding: 0; border: 0; font: 9px Arial; } table { border-collapse: collapse; width: 100%; }'


def _rows(tag, n):
    return ''.join(f'<tr><td>{tag}-ROW {i + 1}</td><td>v</td></tr>' for i in range(n))


def _synthetic_report():
    page = '<section style="break-before: page">{}</section>'
    cases = [
        # 1 · a heading at the page bottom followed by a small table
        '<div style="height:246mm">C1-SPACER</div><h2 style="margin:0;font-size:14px">C1 HEAD</h2>'
        '<table><thead><tr><th>C1-H</th><th>v</th></tr></thead><tbody>' + _rows('C1', 8) + '</tbody></table>',
        # 2 · a long table whose last row would be stranded alone on the next page
        '<table><thead><tr><th>C2-H</th><th>v</th></tr></thead><tbody>' + _rows('C2', 45) + '</tbody></table>',
        # 3 · a tall block the renderer marked "avoid" — must not leave a blank page behind
        '<div style="height:100mm">C3-SPACER</div><div style="break-inside:avoid">'
        '<table><thead><tr><th>C3-H</th><th>v</th></tr></thead><tbody>' + _rows('C3', 60) + '</tbody></table></div>',
        # 4 · a thead-less long table (header row in the body) — the header must repeat
        '<table class="dt"><tr><th>C4-HEADER</th><th>v</th></tr>' + _rows('C4', 70) + '</table>',
        # 5 · KPI tiles under a heading at the page bottom
        '<div style="height:250mm">C5-SPACER</div><h3 style="margin:0">C5 HEAD</h3>'
        '<div class="tiles" style="display:flex;gap:4px"><div style="flex:1;height:20mm">C5 TILE A</div>'
        '<div style="flex:1;height:20mm">C5 TILE B</div></div>',
    ]
    return _doc(''.join(page.format(c) for c in cases))


def _wide_report():
    # 6 · a table wider than the page inside a screen scroll box (its own document: Chrome
    #     shrinks a WHOLE document that overflows the page width, which would mask cases 1-5)
    return _doc('<div style="overflow-x:auto"><table style="width:1500px"><tr><td>C6-FIRST</td>'
                + ''.join(f'<td>c{i}</td>' for i in range(18)) + '<td>C6-LASTCOL</td></tr></table></div>')


def _doc(body):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{_A4} {_ROW}</style></head>'
            f'<body>{body}</body></html>')


def _print(html, chrome, folder, name):
    from p6_export.pdf import run_chrome
    src = os.path.join(folder, name + '.html')
    with open(src, 'w', encoding='utf-8') as fh:
        fh.write(html)
    out = os.path.join(folder, name + '.pdf')
    run_chrome(chrome, [f'--print-to-pdf={out}', '--no-pdf-header-footer',
                        'file:///' + src.replace(os.sep, '/')], timeout=120)
    import pymupdf
    with pymupdf.open(out) as d:
        return [p.get_text() for p in d]


def _violations(pages):
    """The rule breaches this synthetic report is built to provoke (page texts in order).
    A case whose markers are not in the document is not checked."""
    v = []

    def page_of(marker):
        return next((i for i, t in enumerate(pages) if marker in t), None)

    p = page_of('C1 HEAD')
    if p is not None:
        if 'C1-ROW 1\n' not in pages[p]:
            v.append('C1: heading orphaned from its table')
        elif len(re.findall(r'\bC1-ROW \d+\n', pages[p])) != 8:
            v.append('C1: small table split across pages')
    for tag in ('C2', 'C3', 'C4'):
        for i, t in enumerate(pages):
            n = len(re.findall(rf'\b{tag}-ROW \d+\n', t))
            if 0 < n < 3:
                v.append(f'{tag}: {n} row(s) stranded on page {i + 1}')
    p = page_of('C3-SPACER')
    if p is not None and 'C3-ROW 1\n' not in pages[p]:
        v.append('C3: tall block pushed to the next page (blank area left)')
    for i, t in enumerate(pages):
        if 'C4-ROW' in t and 'C4-HEADER' not in t:
            v.append(f'C4: header not repeated on page {i + 1}')
    p = page_of('C5 HEAD')
    if p is not None and 'C5 TILE A' not in pages[p]:
        v.append('C5: KPI tiles separated from their heading')
    if any('C6-FIRST' in t for t in pages) and not any('C6-LASTCOL' in t for t in pages):
        v.append('C6: wide table cut at the page edge')
    return v


def test_chrome_prints_the_synthetic_report_breaking_rules_without_and_keeping_them_with():
    try:
        import pymupdf  # noqa: F401
    except ImportError:
        pytest.skip('PyMuPDF not installed')
    from p6_export.pdf import chrome_candidates
    found = chrome_candidates(None)
    if not found:
        pytest.skip('no Chromium installed')
    html, wide = _synthetic_report(), _wide_report()
    with tempfile.TemporaryDirectory() as folder:
        before = (_violations(_print(html, found[0], folder, 'before'))
                  + _violations(_print(wide, found[0], folder, 'wide_before')))
        after = (_violations(_print(rt.with_pagination(html), found[0], folder, 'after'))
                 + _violations(_print(rt.with_pagination(wide), found[0], folder, 'wide_after')))
    # without the layer the document breaks the rules (the test is meaningful) …
    assert any(s.startswith('C1') for s in before), before
    assert any(s.startswith('C2') for s in before), before
    assert any(s.startswith('C3') for s in before), before
    assert any(s.startswith('C4') for s in before), before
    assert any(s.startswith('C6') for s in before), before
    # … and with the ONE shared layer it keeps every one of them
    assert after == [], after
