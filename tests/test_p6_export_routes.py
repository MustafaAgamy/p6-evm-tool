"""POST /api/export/{pdf,html,docx,xlsx} — every output built from the ONE final report HTML.

The HTML is what the picker holds after the user unticked a part (the real EVM renderer's
output with one [data-part] removed, exactly like report_parts.pruneHtml does)."""
import json
import re
import urllib.request

import docx
import openpyxl
import pytest

from p6_evm.evm_report import render_evm_report
from p6_export.pdf import chrome_candidates

RESULT = {'pv': 5.0e6, 'ev': 4.2e6, 'ac': 4.0e6, 'spi': 0.84, 'cpi': 1.05, 'delay_days': 9,
          'categories': {'Civil Works': {'weight': 0.7, 'planned_pct': 0.5, 'actual_pct': 0.42},
                         'MEP Works': {'weight': 0.3, 'planned_pct': 0.2, 'actual_pct': 0.1}}}
META = {'project_name': 'Route Test', 'data_date': '01-Mar-2026', 'report_date': '27-Sep-2026',
        'source_file': 'route.xml'}


def _post(port, path, body):
    req = urllib.request.Request(f'http://127.0.0.1:{port}/{path}', data=json.dumps(body).encode(),
                                 headers={'Content-Type': 'application/json'})
    return json.loads(urllib.request.urlopen(req, timeout=300).read())


def _final_html():
    html = render_evm_report(RESULT, META, theme='light')
    # untick the "Planned vs Actual — bar chart" part (drop the element, as the picker does)
    return re.sub(r'<div data-part="progress\.chart".*?</div></div></div>(?=</div>)', '', html, flags=re.S)


def test_final_html_fixture_really_dropped_the_part():
    h = _final_html()
    assert 'data-part="progress.chart"' not in h and 'data-part="progress.kpis"' in h


def _body(tmp_path, ext):
    return {'html': _final_html(), 'output_path': str(tmp_path / f'r.{ext}'),
            'title': 'Earned Value (EVM)', 'meta': {'feature': 'Earned Value (EVM)', 'project': 'Route Test'}}


def test_export_html_route(test_server, tmp_path):
    r = _post(test_server, 'api/export/html', _body(tmp_path, 'html'))
    assert r['ok'], r
    txt = (tmp_path / 'r.html').read_text(encoding='utf-8')
    assert '<meta charset="utf-8">' in txt and 'Civil Works' in txt
    assert 'Planned vs Actual — bar chart' not in txt


def test_export_docx_route(test_server, tmp_path):
    r = _post(test_server, 'api/export/docx', _body(tmp_path, 'docx'))
    assert r['ok'], r
    d = docx.Document(str(tmp_path / 'r.docx'))
    cells = ' '.join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert 'Civil Works' in cells and 'MEP Works' in cells and '84%' in cells
    text = ' '.join(p.text for p in d.paragraphs)
    assert 'Category Weights & Overall Progress' in text
    head = d.sections[0].header.paragraphs[0].text
    assert 'Earned Value (EVM)' in head and 'Route Test' in head


def test_export_xlsx_route(test_server, tmp_path):
    r = _post(test_server, 'api/export/xlsx', _body(tmp_path, 'xlsx'))
    assert r['ok'], r
    wb = openpyxl.load_workbook(str(tmp_path / 'r.xlsx'))
    first = [c.value for c in wb.worksheets[0]['A'][:2]]
    assert first[0].endswith('— Earned Value (EVM)')
    flat = [v for ws in wb for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    assert 'Civil Works' in flat and 'Planned vs Actual — bar chart' not in flat


@pytest.mark.skipif(not chrome_candidates(None), reason='no Chrome/Chromium installed')
def test_export_pdf_route(test_server, tmp_path):
    import pymupdf
    r = _post(test_server, 'api/export/pdf', _body(tmp_path, 'pdf'))
    assert r['ok'], r
    doc = pymupdf.open(str(tmp_path / 'r.pdf'))
    text = ' '.join(p.get_text() for p in doc)
    assert 'Civil Works' in text and 'MEP Works' in text


@pytest.mark.parametrize('kind', ['pdf', 'html', 'docx', 'xlsx'])
def test_export_routes_validate_input(test_server, tmp_path, kind):
    r = _post(test_server, f'api/export/{kind}', {'html': '<p>x</p>'})
    assert r == {'ok': False, 'error': 'No output path provided'}
    r = _post(test_server, f'api/export/{kind}', {'html': '  ', 'output_path': str(tmp_path / 'x')})
    assert not r['ok'] and 'empty' in r['error']
