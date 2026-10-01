"""Word page composition on the REAL writers (owner point 14 — findings NARR-WORD-2, CAL-WORD-1).

NARR-WORD-2: the Baseline Narrative Word report left a calendar's title ("Roots Silos - 24 Hrs -
Phase I Scope — 17 activities") at the bottom of one page and its 255 pt timeline chart on the
next. CAL-WORD-1: the Calendar Audit Word export split its 11-row non-working-days table 5 + 6
across two pages. Both are now kept together: the calendar title keeps with its chart (in the
narrative writer itself, and by the shared Word pass), a table that fits in about a third of
a page keeps every row with the next (the shared pass that p6_export.to_docx runs).

The XML tests read the keep flags the writers emit; the Word proofs (skipped when Microsoft
Word is not installed) push the block to the bottom of a page and read Word's OWN pagination:
without the keep flags Word splits them, with them the block moves whole.
"""
import datetime as _dt
import os

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from p6_export import docx_pagination as DP
from p6_export import pagination_check as pc

HOLIDAYS = [('2026-01-07', 'Coptic Christmas'), ('2026-01-25', 'Revolution Day'),
            ('2026-03-20', 'Eid al-Fitr'), ('2026-03-21', 'Eid al-Fitr'),
            ('2026-03-22', 'Eid al-Fitr'), ('2026-04-12', 'Easter Sunday'),
            ('2026-04-13', 'Sham El-Nessim'), ('2026-04-25', 'Sinai Liberation Day'),
            ('2026-05-01', 'Labour Day'), ('2026-05-27', 'Eid al-Adha'),
            ('2026-06-30', 'June 30 Revolution')]


# ── helpers ──────────────────────────────────────────────────────────────────────
def _kwn(p):
    ppr = p.find(qn('w:pPr'))
    k = ppr.find(qn('w:keepNext')) if ppr is not None else None
    return k is not None and k.get(qn('w:val'), 'true') not in ('0', 'false')


def _row_kwn(tr):
    return all(_kwn(p) for tc in tr.findall(qn('w:tc')) for p in tc.findall(qn('w:p')))


def _trflag(tr, name):
    trPr = tr.find(qn('w:trPr'))
    return trPr is not None and trPr.find(qn('w:' + name)) is not None


def _drop_keep(p):
    ppr = p.find(qn('w:pPr'))
    if ppr is not None:
        for k in ppr.findall(qn('w:keepNext')):
            ppr.remove(k)


def _body_items(document):
    return [el for el in document.element.body if el.tag in (qn('w:p'), qn('w:tbl'))]


def _text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def _has_drawing(p):
    return p.find('.//' + qn('w:drawing')) is not None


# ── CAL-WORD-1 · the Calendar Audit Word export ─────────────────────────────────
def _calendar_html():
    from p6_calendar.report import render_calendar_report
    hd = [{'date': d, 'weekday': _dt.date.fromisoformat(d).strftime('%A'), 'reason': r}
          for d, r in HOLIDAYS]
    res = {'dashboard': {'data_date': '2025-12-11'}, 'project': {}, 'primary_calendar_id': 1,
           'by_calendar': {1: {'monthly_stats': [], 'hours_profiles': [],
                               'exceptions': {'holidays': [], 'special': [], 'shutdowns': [],
                                              'holiday_dates': hd}}},
           'assigned_calendars': [{'object_id': 1, 'name': 'Roots Silos Phase I Scope'}]}
    return render_calendar_report(res, {'project_name': 'Synthetic'}, sections=['exceptions'])


def _calendar_docx(path):
    from p6_export.to_docx import html_to_docx
    html_to_docx(_calendar_html(), path, feature='Calendar Audit', use_chrome=False)
    return path


def _holiday_table(document):
    for t in document.tables:
        if [c.text.strip() for c in t.rows[0].cells][:3] == ['Date', 'Day', 'Description']:
            return t
    raise AssertionError('no Date | Day | Description table')


def test_calendar_word_non_working_days_table_is_kept_whole(tmp_path):
    d = Document(_calendar_docx(str(tmp_path / 'cal.docx')))
    rows = _holiday_table(d)._tbl.findall(qn('w:tr'))
    assert len(rows) == 12
    assert all(_trflag(tr, 'cantSplit') for tr in rows)               # a row never splits
    assert _trflag(rows[0], 'tblHeader')                               # header repeats if ever needed
    assert all(_row_kwn(tr) for tr in rows[:-1]) and not _row_kwn(rows[-1])   # kept whole


# ── NARR-WORD-2 · the Baseline Narrative calendar timeline ──────────────────────
def _calendar_payload(n=3):
    months = ['Jan 2026', 'Feb 2026', 'Mar 2026', 'Apr 2026', 'May 2026', 'Jun 2026']
    cals = [{'name': name, 'activity_count': cnt, 'months': months,
             'net_working_days': [22, 20, 21, 22, 21, 22], 'nonworking_days': [9, 8, 10, 8, 10, 8]}
            for name, cnt in (('Roots Silos Phase I Scope', 1474),
                              ('Roots Silos - 24 Hrs - Phase I Scope', 17),
                              ('Roots Silos Phase I Scope - For Loading Test', 12))[:n]]
    return {'calendars': cals}


def _narrative_calendar_doc(filler=0, paginate=True):
    from p6_narrative import docx_calendar
    d = Document()
    for i in range(filler):
        d.add_paragraph(f'Filler line {i + 1} of the page above the calendar timeline.')
    docx_calendar.render_calendar(d, _calendar_payload(), None, 8)
    if paginate:
        DP.paginate_docx(d)
    return d


def _calendar_titles(d):
    items = _body_items(d)
    return [(el, items[i + 1]) for i, el in enumerate(items[:-1])
            if el.tag == qn('w:p') and _text(el).endswith(' activities')]


def test_narrative_calendar_title_keeps_with_its_chart_straight_from_the_writer():
    d = _narrative_calendar_doc(paginate=False)            # the writer alone, before the shared pass
    pairs = _calendar_titles(d)
    assert len(pairs) == 3
    for title, nxt in pairs:
        assert nxt.tag == qn('w:tbl') or _has_drawing(nxt), _text(nxt)[:40]
        assert _kwn(title), _text(title)
    # the shared pass keeps it, and never chains a chart into the next calendar's title
    d = _narrative_calendar_doc()
    for title, nxt in _calendar_titles(d):
        assert _kwn(title)
        if nxt.tag == qn('w:p'):
            assert not _kwn(nxt)


# ── Word proofs ──────────────────────────────────────────────────────────────────
def _need_word():
    if not pc.word_available():
        pytest.skip('Microsoft Word (COM) is not available')


def _insert_filler(document, n):
    first = document.element.body[0]
    for i in range(n):
        p = OxmlElement('w:p')
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = f'Filler line {i + 1} that moves the block down the page.'
        r.append(t)
        p.append(r)
        first.addprevious(p)


def _layout(document, path):
    document.save(path)
    return pc.word_layout(path)


def _pitch_and_bottom(lay):
    ys = [it['y'] for it in lay['items'] if it['k'] == 'p' and (it.get('t') or '').startswith('Filler line')
          and it['page'] == 1]
    return (ys[-1] - ys[0]) / (len(ys) - 1), lay['page']['h'] - lay['page']['bottom']


def test_word_keeps_the_calendar_holidays_table_whole_at_a_page_bottom(tmp_path):
    """CAL-WORD-1 with Word's own pagination: the holidays table pushed to the bottom of a
    page, room for about 4 of its 12 rows. Without the keep flags Word splits it; with them
    the whole table (and its heading) moves to the next page."""
    _need_word()
    src = _calendar_docx(str(tmp_path / 'cal.docx'))
    probe = Document(src)
    _insert_filler(probe, 10)
    lay = _layout(probe, str(tmp_path / 'probe.docx'))
    pitch, bottom = _pitch_and_bottom(lay)
    tbl = next(it for it in lay['items'] if it['k'] == 'table' and it['t'].startswith('Date'))
    add = int((bottom - 4 * 15.5 - tbl['rows'][0]['y']) / pitch)       # ~4 rows of room left
    assert add > 0, (bottom, tbl['rows'][0]['y'])

    def run(strip):
        d = Document(src)
        _insert_filler(d, 10 + add)
        if strip:                                                     # the audit-era output
            for tr in _holiday_table(d)._tbl.findall(qn('w:tr')):
                for tc in tr.findall(qn('w:tc')):
                    for p in tc.findall(qn('w:p')):
                        _drop_keep(p)
        lay = _layout(d, str(tmp_path / ('before.docx' if strip else 'after.docx')))
        t = next(it for it in lay['items'] if it['k'] == 'table' and it['t'].startswith('Date'))
        return {r['page'] for r in t['rows']}
    assert len(run(strip=True)) == 2                                  # the defect, reproduced
    assert len(run(strip=False)) == 1                                 # kept whole


def test_word_keeps_the_narrative_calendar_title_with_its_chart(tmp_path):
    """NARR-WORD-2 with Word's own pagination: a calendar title at the very bottom of a page.
    Without keep-with-next Word leaves it there and starts the chart on the next page; with it
    the title moves down with its chart."""
    _need_word()
    lay = _layout(_narrative_calendar_doc(filler=10), str(tmp_path / 'probe.docx'))
    pitch, bottom = _pitch_and_bottom(lay)
    first = next(it for it in lay['items'] if it['k'] == 'p' and (it.get('t') or '').endswith(' activities'))
    add = int((bottom - 1.5 * pitch - first['y']) / pitch)            # room for the title only
    assert add > 0

    def run(strip):
        d = _narrative_calendar_doc(filler=10 + add)
        if strip:                                                     # the audit-era output
            for title, _ in _calendar_titles(d):
                _drop_keep(title)
        lay = _layout(d, str(tmp_path / ('before.docx' if strip else 'after.docx')))
        items = lay['items']
        i = next(j for j, it in enumerate(items) if it['k'] == 'p' and (it.get('t') or '').endswith(' activities'))
        chart = next(it for it in items[i + 1:] if it['k'] != 'p' or it.get('shape_h'))
        return items[i]['page'], chart['page']
    t, c = run(strip=True)
    assert c == t + 1, (t, c)                                         # the defect, reproduced
    t, c = run(strip=False)
    assert c == t, (t, c)                                             # title travels with its chart
