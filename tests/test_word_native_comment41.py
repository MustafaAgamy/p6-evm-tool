"""Owner comment 41 — "the Word file must match the PDF exactly, NOT as a picture: Word that can
be edited, with the same format and style."

Charts, gauges and tile rows used to reach Word as pictures of the PDF.  They are now rebuilt
from the PDF's own vector drawing as native Word shapes (every bar selectable) and real Word
text (every label retypable) at the same positions — ``p6_export/vector_shapes.py``.
"""
import re
import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml

from p6_export import docx_pagination, html_model as HM, vector_shapes as VS
from p6_export.pdf import chrome_candidates
from tests.test_server import _post_json


def _rect(x0, y0, x1, y1, fill='2563EB'):
    return {'k': 'path', 'segs': [], 'fill': fill, 'stroke': None, 'width': 0, 'fa': None, 'sa': None,
            'dash': False, 'rect': True, 'bbox': (x0, y0, x1, y1)}


def _text(text, x, base, size=8.0, font='Segoe UI'):
    return {'k': 'text', 'text': text, 'size': size, 'color': '111111', 'bold': False, 'italic': False,
            'font': font, 'vert': False, 'up': False, 'base': base,
            'bbox': (x, base - size, x + len(text) * size * 0.5, base + size * 0.3)}


def test_a_drawing_is_a_group_of_word_shapes_and_word_text_never_a_picture():
    prims = [_rect(10, 20, 200, 28), _text('Phase A', 10, 15), _text('95%', 220, 27)]
    xml = VS.group_drawing_xml(prims, 300, 60, [100])
    el = parse_xml(xml)                                              # well-formed, namespaced
    assert el.tag.endswith('}drawing')
    assert '<wpg:wgp>' in xml and 'pic:pic' not in xml and 'a:blip' not in xml
    assert xml.count('prstGeom prst="rect"') >= 2                    # the frame + the bar
    assert '>Phase A</w:t>' in xml and '>95%</w:t>' in xml           # labels are real, editable text
    assert 'w:lineRule="exact"' in xml


def test_labels_share_a_few_text_layers_not_one_text_box_each():
    # one text box per label made Word take minutes to open a report (0.1 s per box)
    prims = [_text('row %d' % i, 10, 20 + i * 14) for i in range(40)] + [_text('%d%%' % i, 200, 20 + i * 14) for i in range(40)]
    xml = VS.group_drawing_xml(prims, 300, 600, [100])
    assert xml.count('txBox="1"') <= 3
    assert xml.count('<w:t ') == 80                                  # …and no label is lost
    assert xml.count('<w:tab w:val="left"') == 40                    # the value sits at its own tab stop


def test_the_line_model_puts_each_baseline_where_the_pdf_has_it():
    # Word, exact line spacing: baseline = top of the line + 0.8 x line height (measured in Word)
    lines = VS._text_lines([_text('a', 0, 30.0, 8), _text('b', 0, 44.0, 8), _text('c', 0, 120.0, 10)])
    (layer,) = VS._layers(lines)
    y = layer['top']
    for ln, spacer, height in layer['rows']:
        y += spacer
        assert abs((y + VS.ASC * height) - ln['base']) < 0.05
        assert height >= VS.MIN_K * ln['size'] - 1e-6               # never so low that glyphs are cut
        y += height
    # two labels too close together to stack go to separate layers
    close = VS._layers(VS._text_lines([_text('a', 0, 30.0, 8), _text('b', 60, 33.0, 8)]))
    assert len(close) == 2


def test_a_scaled_drawing_scales_its_text_too_and_is_shown_one_to_one():
    # Word scales the SHAPES of a group but not the line heights of its text: labels drifted
    xml = VS.group_drawing_xml([_rect(0, 0, 590, 40), _text('x', 5, 20, 10.0)], 590, 40, [100], max_w_emu=300 * VS.EMU_PT)
    ext = re.search(r'<wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="(\d+)" cy="(\d+)"/><a:chOff x="0" y="0"/>'
                    r'<a:chExt cx="(\d+)" cy="(\d+)"/>', xml).groups()
    assert ext[0] == ext[2] and ext[1] == ext[3]                     # 1:1 — no group scaling
    assert int(ext[0]) <= 300 * VS.EMU_PT
    assert '<w:sz w:val="10"/>' in xml                               # 10 pt x 0.5 = 5 pt = 10 half-points


def test_the_first_child_is_an_invisible_frame_the_size_of_the_drawing():
    # without it Word moved the shapes to the top of the frame but not the text
    xml = VS.group_drawing_xml([_rect(10, 30, 50, 40), _text('x', 5, 20)], 100, 60, [100])
    first = re.search(r'<wps:wsp>.*?</wps:wsp>', xml, flags=re.S).group(0)
    assert 'name="Frame' in first and '<a:noFill/><a:ln><a:noFill/></a:ln>' in first
    assert f'<a:off x="0" y="0"/><a:ext cx="{VS._emu(100 + VS.RIGHT_PAD)}" cy="{VS._emu(60)}"/>' in first


def test_a_dotted_border_is_one_dashed_shape_not_hundreds_of_dashes():
    segs = []
    for i in range(40):                                              # 40 tiny dashes along a 200 pt line
        x = i * 5.0
        segs += [('M', (x, 0)), ('L', (x + 3, 0)), ('L', (x + 3, 1)), ('L', (x, 1)), ('Z',)]
    p = VS._dashes_to_outline(segs, 'D2D8E2', None, (0, 0, 200, 1))
    assert p and p['dash'] and p['stroke'] == 'D2D8E2' and p['fill'] is None and len(p['segs']) == 2
    box = [('M', (0, 0)), ('L', (100, 0)), ('L', (100, 80)), ('L', (0, 80)), ('Z',)]
    assert VS._dashes_to_outline(box, 'FFFFFF', None, (0, 0, 100, 80)) is None       # a real shape is kept


def test_word_uses_the_fonts_the_pdf_uses():
    assert VS._family('BCDEFG+SegoeUI-Semibold') == 'Segoe UI Semibold'
    assert VS._family('SegoeUIBlack') == 'Segoe UI Black' and VS._family('SegoeUI-Bold') == 'Segoe UI'
    assert VS._bold({'font': 'SegoeUI-Bold', 'flags': 16}) and not VS._bold({'font': 'SegoeUIBlack', 'flags': 16})
    # 'system-ui, -apple-system, Arial' is Segoe UI in the browser (and so in the PDF) — not Arial
    assert HM._first_family('system-ui, -apple-system, Arial, sans-serif') == 'Segoe UI'
    assert HM._first_family('"Inter", Arial') == 'Inter' and HM._first_family('sans-serif') is None


def test_the_page_fitter_never_rescales_a_native_drawing():
    d = Document()
    p = d.add_paragraph()
    p.add_run()._r.append(parse_xml(VS.group_drawing_xml([_rect(0, 0, 100, 400), _text('x', 5, 20)], 100, 400, [100])))
    before = p._p.xml
    assert docx_pagination.fit_picture(p._p, 120.0) is False
    assert p._p.xml == before


@pytest.mark.skipif(not chrome_candidates(None), reason='no Chrome/Chromium installed')
def test_the_word_file_of_a_report_has_no_picture_and_keeps_tables_as_tables(test_server, tmp_path):
    bars = ('<div class="bars">' + ''.join(
        f'<div class="wrow"><div class="wname">Phase {c}</div><div class="wtrack"><div class="wfill" '
        f'style="width:{pc}%"></div></div><div class="wpct">{pc}%</div></div>' for c, pc in (('A', 95), ('B', 68))) + '</div>')
    table = '<table><thead><tr><th>Activity ID</th><th>Float</th></tr></thead><tbody><tr><td>A1</td><td>3</td></tr></tbody></table>'
    html = ('<html><head><style>body{font-family:system-ui,Arial}.wrow{display:flex;gap:8px;align-items:center;margin:6px 0}'
            '.wname{width:90px}.wtrack{flex:1;height:8px;background:#e5e7eb;border-radius:4px}'
            '.wfill{background:#2563eb;height:8px;border-radius:4px}</style></head><body>'
            f'<div data-sec="scope"><h2>Scope weight</h2>{bars}{table}</div></body></html>')
    out = tmp_path / 'r.docx'
    _, d = _post_json(test_server, '/api/export/docx', {'html': html, 'output_path': str(out),
                                                        'meta': {'feature': 'Update Analysis', 'project': 'P'}})
    assert d['ok'], d
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        body = z.read('word/document.xml').decode('utf-8')
    assert not [n for n in names if n.startswith('word/media/')], 'no picture in the Word file'
    assert '<wpg:wgp>' in body and '<wps:wsp>' in body               # the chart: native Word shapes
    group = re.search(r'<wpg:wgp>.*?</wpg:wgp>', body, flags=re.S).group(0)
    assert '>Phase A</w:t>' in group and '95%' in group              # its labels: editable Word text
    assert body.count('<w:tbl>') >= 1 and 'Activity ID' in body      # the table: a real Word table
    Document(str(out))                                               # opens as a valid document
