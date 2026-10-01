"""Owner point 14 — the automatic page checker (p6_export.pagination_check).

Unit tests on small PDFs drawn with PyMuPDF and on a Word layout dump, the CLI contract, and
the PROOF on synthetic long reports built here: the same report printed by Chrome WITHOUT the
shared pagination rules is flagged for every defect type, and WITH them (report_theme.
with_pagination) the checker finds nothing; the same for a Word report without / with the
Word rules (p6_export.docx_pagination) — laid out by Word itself (COM) or by Spire.Doc.
"""
import io
import json
import os
import tempfile

import pytest

import report_theme as rt
from p6_export import pagination_check as pc

pymupdf = pytest.importorskip('pymupdf')


def _types(res):
    return {f['type'] for f in res['flags']}


# ── small PDFs drawn with PyMuPDF ───────────────────────────────────────────────
A4 = (595, 842)


def _pdf(path, pages):
    """pages: list of op lists — ('t', x, y, text, size, bold) text baseline at y;
    ('r', x0, y0, x1, y1) a filled box."""
    doc = pymupdf.open()
    for ops in pages:
        pg = doc.new_page(width=A4[0], height=A4[1])
        for op in ops:
            if op[0] == 't':
                _, x, y, text, size, bold = op
                pg.insert_text((x, y), text, fontsize=size, fontname='hebo' if bold else 'helv')
            else:
                _, x0, y0, x1, y1 = op
                pg.draw_rect(pymupdf.Rect(x0, y0, x1, y1), color=None, fill=(0.85, 0.9, 0.95))
    doc.save(path)
    doc.close()
    return path


def _running(no, n):
    return [('t', 40, 30, 'PROJECT REPORT  Grain Bulk Terminal', 8, False),
            ('t', 280, 825, f'Page {no} of {n}', 8, False)]


def _para(y0, y1, tag):
    return [('t', 40, y, f'{tag} line {y} of the narrative text of this section, plain words.', 10, False)
            for y in range(int(y0), int(y1), 14)]


def _rows(y0, n, tag, header=True, x=(40, 200, 400)):
    ops, y = [], y0
    if header:
        ops += [('t', x[0], y, 'Activity ID', 10, True), ('t', x[1], y, 'Activity name', 10, True),
                ('t', x[2], y, 'Duration', 10, True)]
        y += 18
    for i in range(n):
        ops += [('t', x[0], y, f'{tag}-{i + 1}', 10, False), ('t', x[1], y, f'{tag} activity {i + 1}', 10, False),
                ('t', x[2], y, f'{(i + 1) * 3} d', 10, False)]
        y += 18
    return ops


def test_clean_report_has_no_flags(tmp_path):
    n = 3
    pages = [
        _running(1, n) + [('t', 40, 70, 'Project overview', 16, True)] + _para(95, 790, 'A'),
        _running(2, n) + [('t', 40, 70, 'Activities', 16, True)] + _rows(95, 38, 'B'),
        _running(3, n) + [('t', 40, 70, 'Notes', 16, True)] + _para(95, 790, 'C'),
    ]
    res = pc.check_pdf(_pdf(str(tmp_path / 'clean.pdf'), pages))
    assert res['status'] == 'ok' and res['pages'] == 3
    assert res['flags'] == [], res['flags']


def test_orphaned_heading_is_flagged_with_its_page(tmp_path):
    n = 3
    pages = [
        _running(1, n) + [('t', 40, 70, 'Scope of work', 16, True)] + _para(95, 780, 'A')
        + [('t', 40, 800, 'Resource summary', 13, True)],                    # heading ends the page
        _running(2, n) + _rows(70, 30, 'B'),
        _running(3, n) + [('t', 40, 70, 'Notes', 16, True)] + _para(95, 790, 'C'),
    ]
    res = pc.check_pdf(_pdf(str(tmp_path / 'orphan.pdf'), pages))
    assert [(f['type'], f['page']) for f in res['flags']] == [('orphaned_heading', 1)], res['flags']
    assert 'Resource summary' in res['flags'][0]['detail']


def test_table_stranding_rows_and_losing_its_header_is_flagged(tmp_path):
    n = 3
    pages = [
        _running(1, n) + [('t', 40, 70, 'Activities', 16, True)] + _rows(95, 38, 'T'),     # ends ~797
        _running(2, n) + _rows(70, 2, 'U', header=False)                                 # 2 rows, no header
        + [('t', 40, 140, 'Next part', 16, True)] + _para(165, 790, 'D'),
        _running(3, n) + [('t', 40, 70, 'Notes', 16, True)] + _para(95, 790, 'C'),
    ]
    res = pc.check_pdf(_pdf(str(tmp_path / 'table.pdf'), pages))
    got = {(f['type'], f['page']) for f in res['flags']}
    assert ('table_split_few_rows', 1) in got, res['flags']
    assert ('table_header_not_repeated', 2) in got, res['flags']


def test_large_blank_before_a_pushed_block_but_not_before_a_new_section(tmp_path):
    n = 4
    pages = [
        _running(1, n) + [('t', 40, 70, 'Overview', 18, True)] + _para(95, 300, 'A'),   # 70 % blank …
        _running(2, n) + [('r', 40, 60, 555, 700)] + _para(720, 790, 'B'),               # … a pushed block
        _running(3, n) + [('t', 40, 70, 'Costs', 18, True)] + _para(95, 250, 'C'),       # blank before …
        _running(4, n) + [('t', 40, 70, 'Risks', 18, True)] + _para(95, 790, 'E'),       # … a new section
    ]
    res = pc.check_pdf(_pdf(str(tmp_path / 'blank.pdf'), pages))
    assert [(f['type'], f['page']) for f in res['flags']] == [('large_blank_then_continuation', 1)], res
    assert [(f['type'], f['page']) for f in res['info']] == [('section_break_blank', 3)], res['info']


def test_html_headings_reads_heading_tags_and_renderer_heading_classes():
    heads = pc.html_headings('<h2>Key Dates</h2><div class="sub">Phase I</div><p>body text</p>'
                             '<div class="ct">Cost curve</div>')
    assert heads == ['Key Dates', 'Phase I', 'Cost curve']


# ── Word: the layout analyser on a synthetic Word layout dump ────────────────────
def _p(t, page, y, **kw):
    it = {'k': 'p', 't': t, 'page': page, 'page_end': page, 'y': y, 'y_last': y, 'style': 'Normal',
          'ol': 10, 'bold': False, 'size': 11, 'len': len(t), 'pbb': False, 'brk': False}
    it.update(kw)
    return it


def _tbl(t, rows, ncols=3):
    return {'k': 'table', 't': t, 'rows': rows, 'ncols': ncols, 'page': rows[0]['page'],
            'page_end': rows[-1]['page'], 'y': rows[0]['y']}


def _r(page, y, hdr=False, bold=False):
    return {'page': page, 'y': y, 'hdr': hdr, 'bold': bold, 'n': 3, 't': ''}


def test_word_layout_analysis_flags_each_defect_with_its_page():
    body = [_r(1, 100 + 18 * i) for i in range(35)] + [_r(2, 72), _r(2, 90)]
    layout = {'page': {'h': 842, 'w': 595, 'top': 57, 'bottom': 57}, 'pages': 4, 'items': [
        _p('Scope of work', 1, 60, style='Heading 1'),
        _tbl('Activities', [_r(1, 82, bold=True)] + body),                    # 2 rows stranded, no header
        _p('Resource summary', 2, 760, style='Heading 2'),                    # heading ends page 2
        _tbl('Resources', [_r(3, 57, bold=True)] + [_r(3, 75 + 18 * i) for i in range(8)]),
        _p('Figure 1 - cost curve', 3, 780),                                  # label ends page 3 …
        {'k': 'pic', 't': '', 'page': 4, 'page_end': 4, 'y': 57, 'y_last': 57, 'shape_h': 300,
         'pbb': False, 'brk': False},                                         # … picture on page 4
    ]}
    flags, info, npages = pc.analyze_word_layout(layout)
    got = {(f['type'], f['page']) for f in flags}
    assert npages == 4
    assert ('table_split_few_rows', 2) in got, flags
    assert ('table_header_not_repeated', 2) in got, flags
    assert ('orphaned_heading', 2) in got, flags
    assert ('picture_separated_from_caption', 4) in got, flags


def test_word_table_running_over_a_page_fills_that_page():
    """A table that starts half-way down page 1, fills it and continues on page 2 is page 1's
    last content AND page 2's first — page 1 is full, not "ends early" (the old analyser
    counted a table only on its last page and flagged a large blank there)."""
    page = {'h': 842, 'w': 595, 'top': 57, 'bottom': 57}
    first = [_r(1, 100 + 18 * i) for i in range(10)]
    split = ([_r(1, 300, hdr=True)] + [_r(1, 318 + 18 * i) for i in range(25)]
             + [_r(2, 57, hdr=True)] + [_r(2, 75 + 18 * i) for i in range(10)])
    layout = {'page': page, 'pages': 2, 'items': [
        _tbl('First', [_r(1, 82, hdr=True)] + first),
        _tbl('Codes', split),
        _tbl('Next', [_r(2, 300, hdr=True)] + [_r(2, 318 + 18 * i) for i in range(5)]),
    ]}
    flags, info, _ = pc.analyze_word_layout(layout)
    assert flags == [], flags
    # a real early end is still caught: the long table pushed whole to page 2
    pushed = [_r(2, 57, hdr=True)] + [_r(2, 75 + 18 * i) for i in range(30)]
    layout['items'][1] = _tbl('Codes', pushed)
    layout['items'][2] = _tbl('Next', [_r(2, 640, hdr=True)] + [_r(2, 658 + 18 * i) for i in range(5)])
    flags, info, _ = pc.analyze_word_layout(layout)
    assert [(f['type'], f['page']) for f in flags] == [('large_blank_then_continuation', 1)], flags


def test_word_table_fragment_counts_its_wrapped_last_row():
    """NARR-WORD-1 — a Word table's height was the rows' first-line positions plus ONE usual row
    step per page fragment, so a wrapped (3-line) row ending a fragment was counted as one line:
    a 40 %-of-a-page table split 8+3 (legal: >= 3 rows each side, header repeated) was read as
    a 35 % "small" table and falsely flagged small_table_split."""
    page = {'h': 842, 'w': 595, 'top': 57, 'bottom': 57}
    rows = ([_r(1, 500, hdr=True)] + [_r(1, 518 + 18 * i) for i in range(8)]
            + [_r(2, 57, hdr=True)] + [_r(2, 75 + 18 * i) for i in range(3)])
    layout = {'page': page, 'pages': 2, 'items': [_tbl('Rates', rows)]}
    flags, _, _ = pc.analyze_word_layout(layout)
    assert [f['type'] for f in flags] == ['small_table_split'], flags      # a short-row table IS small
    rows[8]['y_last'] = rows[8]['y'] + 36          # page 1's last row wraps onto 3 lines
    flags, _, _ = pc.analyze_word_layout(layout)
    assert flags == [], flags


def test_word_page_break_by_design_is_not_a_defect():
    layout = {'page': {'h': 842, 'w': 595, 'top': 57, 'bottom': 57}, 'pages': 2, 'items': [
        _p('Overview', 1, 57, style='Heading 1'), _p('Short text.', 1, 90),
        _p('Costs', 2, 57, style='Heading 1', pbb=True), _p('More text.', 2, 90)]}
    flags, info, _ = pc.analyze_word_layout(layout)
    assert flags == [] and [i['type'] for i in info] == ['section_break_blank']
    # the same with a paragraph that holds only a manual page break
    layout['items'][2]['pbb'] = False
    layout['items'].insert(2, _p('', 2, 57, brk=True))
    flags, info, _ = pc.analyze_word_layout(layout)
    assert flags == [] and [i['type'] for i in info] == ['section_break_blank']


# ── CLI ─────────────────────────────────────────────────────────────────────────
def test_cli_prints_json_and_exits_1_on_flags_0_when_clean(tmp_path, capsys):
    n = 3
    bad = _pdf(str(tmp_path / 'bad.pdf'), [
        _running(1, n) + [('t', 40, 70, 'Scope', 16, True)] + _para(95, 780, 'A')
        + [('t', 40, 800, 'Resource summary', 13, True)],
        _running(2, n) + _rows(70, 30, 'B'),
        _running(3, n) + [('t', 40, 70, 'Notes', 16, True)] + _para(95, 790, 'C')])
    assert pc.main([bad]) == pc.EXIT_FLAGS
    out = json.loads(capsys.readouterr().out)
    assert out['flag_count'] == 1 and out['flags'][0]['page'] == 1 and out['counts'] == {'orphaned_heading': 1}
    good = _pdf(str(tmp_path / 'good.pdf'), [_running(1, 1) + [('t', 40, 70, 'Scope', 16, True)]
                                             + _para(95, 790, 'A')])
    assert pc.main([good]) == pc.EXIT_CLEAN
    assert json.loads(capsys.readouterr().out)['flags'] == []
    assert pc.main([str(tmp_path / 'missing.pdf')]) == pc.EXIT_ERROR


def test_word_file_without_a_renderer_is_skipped_with_a_clear_message(tmp_path, monkeypatch, capsys):
    from docx import Document
    path = str(tmp_path / 'r.docx')
    Document().save(path)
    monkeypatch.setattr(pc, 'word_available', lambda: False)
    monkeypatch.setattr(pc, 'spire_available', lambda: False)
    assert pc.main([path]) == pc.EXIT_SKIPPED
    out = json.loads(capsys.readouterr().out)
    assert out['status'] == 'skipped' and 'NOT checked' in out['message'] and out['flags'] == []


# ── PROOF 1 · Chrome: a synthetic long report without / with the shared rules ─────
_CSS = ('@page { size: A4 portrait; margin: 14mm; } body { margin: 0; font: 10px Arial; color: #222; }'
        'h1 { font-size: 20px; margin: 0 0 3mm; } h2 { font-size: 15px; margin: 2mm 0 1mm; }'
        'h3 { font-size: 13px; margin: 2mm 0 1mm; } p { margin: 0 0 1.5mm; }'
        'table { border-collapse: collapse; width: 100%; }'
        'th, td { border: 1px solid #889; padding: 0 4px; height: 6mm; text-align: left; font-size: 10px; }'
        'th { background: #dde4ee; }'
        '.tiles { display: flex; gap: 8px; } .tile { flex: 1; height: 22mm; border: 1px solid #99a;'
        ' background: #eef3fa; padding: 2mm; }'
        '.chart { border: 1px solid #556; height: 80mm; display: flex; align-items: flex-end; gap: 6px;'
        ' padding: 4px; }')


def _html_rows(tag, n):
    return ''.join(f'<tr><td>{tag}-ROW {i}</td><td>{tag} activity {i}</td><td>{i * 3} d</td></tr>'
                   for i in range(1, n + 1))


def _html_table(tag, n, thead=True):
    head = f'<tr><th>{tag} Activity ID</th><th>Activity name</th><th>Duration</th></tr>'
    if thead:
        return f'<table><thead>{head}</thead><tbody>{_html_rows(tag, n)}</tbody></table>'
    return f'<table><tbody>{head}{_html_rows(tag, n)}</tbody></table>'   # header row inside the body


def _block(mm, word):
    """Earlier content of the section, a known height (a filled block, so it is on the page)."""
    return f'<div style="height:{mm}mm;background:#e9edf3">{word} section text (earlier content).</div>'


def _case(title, body):
    return f'<section style="break-before: page"><h1>{title}</h1>{body}</section>'


def _synthetic_long_report():
    """Eight sections, each laid out so that a plain print breaks one composition rule."""
    bars = ''.join(f'<div style="flex:1;background:#4a7fb5;height:{15 + i * 8}%"></div>' for i in range(10))
    cases = [
        _case('Case 1 - Orphaned heading', _block(250, 'Alpha') + '<h2>C1 Resource summary</h2>'
              + _html_table('C1', 8)),
        _case('Case 2 - Heading and intro, chart pushed', _block(215, 'Bravo') + '<h2>C2 Cost curve</h2>'
              '<p>C2 intro: the planned cost curve of the project.</p>'
              '<svg width="600" height="300" viewBox="0 0 600 300"><rect x="1" y="1" width="598" height="298"'
              ' fill="#f4f6fa" stroke="#556"/><path d="M10 290 C 200 280, 300 100, 590 20" fill="none"'
              ' stroke="#c33" stroke-width="3"/></svg>'),
        _case('Case 3 - Chart cut by the page break', _block(200, 'Charlie')
              + f'<div class="chart">{bars}</div><p>C3 after chart</p>'),
        _case('Case 4 - Long table strands rows', _html_table('C4', 42)),
        _case('Case 5 - Header not repeated', _html_table('C5', 70, thead=False)),
        _case('Case 6 - KPI cards separated', _block(252, 'Foxtrot') + '<h3>C6 Key indicators</h3>'
              '<div class="tiles"><div class="tile">SPI<br><b>0.94</b></div><div class="tile">CPI<br>'
              '<b>1.02</b></div><div class="tile">Delay<br><b>12 d</b></div></div>'),
        _case('Case 7 - Small table split', _block(223, 'Golf') + _html_table('C7', 10)),
        _case('Case 8 - Tall block pushed', _block(100, 'Hotel') + '<div style="break-inside: avoid">'
              + _html_table('C8', 60) + '</div>'),
    ]
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{_CSS}</style></head>'
            f'<body>{"".join(cases)}</body></html>')


def _chrome_pdf(html, chrome, folder, name):
    from p6_export.pdf import run_chrome
    src = os.path.join(folder, name + '.html')
    with open(src, 'w', encoding='utf-8') as fh:
        fh.write(html)
    out = os.path.join(folder, name + '.pdf')
    run_chrome(chrome, [f'--print-to-pdf={out}', '--no-pdf-header-footer',
                        'file:///' + src.replace(os.sep, '/')], timeout=120)
    return out


def test_chrome_report_is_flagged_without_the_rules_and_clean_with_them():
    from p6_export.pdf import chrome_candidates
    found = chrome_candidates(None)
    if not found:
        pytest.skip('no Chromium installed')
    html = _synthetic_long_report()
    with tempfile.TemporaryDirectory() as folder:
        before = pc.check_pdf(_chrome_pdf(html, found[0], folder, 'before'))
        after = pc.check_pdf(_chrome_pdf(rt.with_pagination(html), found[0], folder, 'after'))
    # without the rules every kind of defect the report was built to provoke is found …
    assert _types(before) >= {
        'orphaned_heading', 'heading_separated_from_block', 'graphic_cut', 'table_split_few_rows',
        'table_header_not_repeated', 'kpi_separated_from_heading', 'small_table_split',
        'large_blank_then_continuation'}, before['flags']
    assert all(f['page'] >= 1 for f in before['flags'])
    # … on the page where it happens (case N starts on page 2N-1)
    got = {(f['type'], f['page']) for f in before['flags']}
    assert ('orphaned_heading', 1) in got and ('kpi_separated_from_heading', 11) in got, before['flags']
    # … and with the ONE shared pagination layer there is nothing left to flag
    assert after['flags'] == [], after['flags']


# ── PROOF 2 · Word: a synthetic long Word report without / with the Word rules ────
_LONG = ('This paragraph stands for the earlier narrative of the section: it describes the scope, the '
         'assumptions and the data sources in enough words to be a real body paragraph rather than a '
         'short label or an introduction line, so no keep rule treats it as a lead-in to what follows.')


def _synthetic_long_docx(fix):
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from p6_export import docx_pagination as DP

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)

    def block(pt, word):              # earlier content of a known height (exact line spacing)
        p = doc.add_paragraph(f'{word} section text (earlier content).')
        f = p.paragraph_format
        f.line_spacing_rule, f.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(pt)
        f.space_before = f.space_after = Pt(0)

    def heading(text, level):         # a renderer that does not keep headings with their content
        h = doc.add_heading(text, level)
        h.paragraph_format.keep_with_next = False
        if level == 1:
            h.paragraph_format.page_break_before = True

    def table(tag, n):
        t = doc.add_table(rows=n + 1, cols=3)
        t.style = 'Table Grid'
        for j, h in enumerate((f'{tag} Activity ID', 'Activity name', 'Duration')):
            t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
        for i in range(1, n + 1):
            for j, v in enumerate((f'{tag}-ROW {i}', f'{tag} activity {i}', f'{i * 3} d')):
                t.rows[i].cells[j].text = v
        for tr in t._tbl.findall(qn('w:tr')):
            h = OxmlElement('w:trHeight')
            h.set(qn('w:val'), '340')
            h.set(qn('w:hRule'), 'exact')
            tr.get_or_add_trPr().append(h)

    def picture():
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 600, 300), 0)
        pix.set_rect(pix.irect, (70, 120, 180))
        doc.add_picture(io.BytesIO(pix.tobytes('png')), height=Pt(220))

    doc.add_paragraph('Synthetic long report - Word pagination proof')
    heading('Case W1 - Orphaned heading', 1); block(648, 'Alpha')
    heading('W1 Resource summary', 2); table('W1', 8)
    heading('Case W2 - Long table strands rows', 1); table('W2', 40)
    heading('Case W3 - Picture and its label', 1); block(520, 'Charlie')
    doc.add_paragraph('Figure W3 - planned cost curve'); picture()
    heading('Case W4 - Small table split', 1); block(560, 'Delta'); doc.add_paragraph(_LONG); table('W4', 10)
    heading('Case W5 - KPI cards', 1); block(650, 'Echo'); heading('W5 Key indicators', 2)
    k = doc.add_table(rows=2, cols=4)
    k.style = 'Table Grid'
    for j, (a, b) in enumerate((('SPI', '0.94'), ('CPI', '1.02'), ('Delay', '12 d'), ('Float', '4 d'))):
        k.rows[0].cells[j].text, k.rows[1].cells[j].text = a, b
    heading('Case W6 - Heading and intro', 1); block(628, 'Foxtrot')
    heading('W6 Cost table', 2); doc.add_paragraph('W6 intro: the cost of each package.'); table('W6', 12)
    if fix:
        DP.paginate_docx(doc)
    return doc


def _check_word_pair(engine, tmp_path):
    before_path, after_path = str(tmp_path / 'before.docx'), str(tmp_path / 'after.docx')
    _synthetic_long_docx(False).save(before_path)
    _synthetic_long_docx(True).save(after_path)
    before = pc.check_docx(before_path, engine=engine, timeout=600)
    after = pc.check_docx(after_path, engine=engine, timeout=600)
    assert before['status'] == 'ok' and after['status'] == 'ok', (before['message'], after['message'])
    return before, after


def test_word_report_laid_out_by_word_is_flagged_without_the_rules_and_clean_with_them(tmp_path):
    if not pc.word_available():
        pytest.skip('Microsoft Word (COM) is not available')
    before, after = _check_word_pair('word', tmp_path)
    assert before['engine'] == 'word-com'
    assert _types(before) >= {
        'orphaned_heading', 'table_split_few_rows', 'table_header_not_repeated',
        'picture_separated_from_caption', 'small_table_split', 'kpi_separated_from_heading',
        'heading_separated_from_block'}, before['flags']
    assert after['flags'] == [], after['flags']


def test_word_report_rendered_by_spire_is_flagged_without_the_rules_and_clean_with_them(tmp_path):
    if not pc.spire_available():
        pytest.skip('Spire.Doc is not installed')
    before, after = _check_word_pair('spire', tmp_path)
    assert before['engine'] == 'spire' and 'first 10 pages' in before['message']
    assert _types(before) >= {'orphaned_heading', 'table_split_few_rows', 'small_table_split'}, before['flags']
    assert after['flags'] == [], after['flags']
