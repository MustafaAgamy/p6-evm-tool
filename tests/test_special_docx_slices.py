"""Reporting Studio Word (.docx): reused feature sections as page-sized pictures
(owner point 14 — findings STUDIO-WORD-1 / STUDIO-WORD-2).

A reused section used to be ONE screenshot of a fixed 920 x 1400 px viewport placed 6.3 in
wide (690 pt — taller than the room under its heading): Word left the numbered heading
alone on a near-blank page, a long section was CUT at the screenshot edge and a short one
was a mostly-white picture. The section is now printed by Chrome on Word-page-sized pages
and every printed page becomes one trimmed picture. These tests prove it:

* the page checker's picture rules (``docx_picture_flags``) catch a cut and a mostly-blank
  picture and pass a clean one (no browser needed);
* with Chrome: the OLD screenshot of a long / a short section is flagged, the NEW slices of
  the same sections are clean, complete (several slices) and each fits a Word page — the
  first one under the heading + caption;
* with Word installed: Word's own pagination of the new file has zero flags.
"""
import io
import os

import pytest

from p6_export import pagination_check as pc
from p6_special import docx_report as DR


# ── helpers ────────────────────────────────────────────────────────────────────────
def _png(lines, h_px, w_px=900, rule_at_bottom=False, top=10):
    """A white PNG ``w_px`` x ``h_px`` with ``lines`` text lines from ``top`` down (each
    line ~22 px), optionally a full-width rule on its last pixel row."""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=w_px, height=h_px)
    y = top + 16
    for i in range(lines):
        page.insert_text((12, y), 'Row %03d   CONS.PL.S%02d.1000   Concrete Pouring   14 wd' % (i, i % 40),
                         fontsize=15)
        y += 22
    if rule_at_bottom:
        page.draw_rect(pymupdf.Rect(0, h_px - 1.5, w_px, h_px), color=(0.6, 0.6, 0.6),
                       fill=(0.6, 0.6, 0.6))
    return page.get_pixmap(alpha=False).tobytes('png')


def _docx_with(pictures, path):
    """A .docx holding each (png, width_in) picture under its own Heading 1."""
    from docx import Document
    from docx.shared import Inches
    d = Document()
    for i, (png, w_in) in enumerate(pictures, 1):
        d.add_heading('%d Section' % i, level=1)
        d.add_picture(io.BytesIO(png), width=Inches(w_in))
    d.save(path)
    return path


def _chrome():
    from p6_export.pdf import chrome_candidates
    found = chrome_candidates(None)
    return found[0] if found else None


_CSS = ('.sx{font-family:Calibri,sans-serif;font-size:13px}'
        '.sx table{border-collapse:collapse;width:100%}'
        '.sx td,.sx th{border:1px solid #cfd8e3;padding:6px 8px;text-align:left}'
        '.sx th{background:#dbe5f1}')


def _long_section(rows=140):
    body = ''.join('<tr><td>CONS.PL.S%02d.%04d</td><td>Concrete pouring for pile cap %d</td>'
                   '<td>2026-07-%02d</td><td>%d wd</td></tr>' % (i % 40, 1000 + i, i, i % 28 + 1, i % 17 + 1)
                   for i in range(1, rows + 1))
    return ('<div class="sx"><p>Lag &amp; lead register - every relationship with a lag.</p>'
            '<table><thead><tr><th>Activity ID</th><th>Activity name</th><th>Start</th>'
            '<th>Duration</th></tr></thead><tbody>%s</tbody></table>'
            '<p>End of register (row %d).</p></div>' % (body, rows))


def _short_section():
    return ('<div class="sx"><p>Each distinct working-time period in the calendar.</p>'
            '<table><thead><tr><th>Period</th><th>Hours</th><th>Days / week</th></tr></thead>'
            '<tbody><tr><td>Normal</td><td>08:00-15:59</td><td>6</td></tr></tbody></table></div>')


def _item(n, html):
    return {'id': 'x:%d' % n, 'title': 'Reused section %d' % n, 'feature': 'audit',
            'feature_title': 'Schedule Audit', 'ctype': 'html',
            'payload': {'kind': 'html', 'css': _CSS, 'html': html}}


def _pictures(path):
    """(height_pt, heading) of every body picture, in reading order."""
    from docx import Document
    from docx.oxml.ns import qn
    d, out, head = Document(path), [], ''
    for p in d.element.body.iter(qn('w:p')):
        st = p.find(qn('w:pPr') + '/' + qn('w:pStyle'))
        if st is not None and str(st.get(qn('w:val'))).startswith('Heading1'):
            head = ''.join(t.text or '' for t in p.iter(qn('w:t')))
        for ext in p.iter(qn('wp:extent')):
            out.append((int(ext.get('cy')) / 12700.0, head))
    return out


# ── the checker's picture rules (no browser) ──────────────────────────────────────
def test_picture_check_flags_a_cut_and_a_mostly_blank_picture_and_passes_clean_ones(tmp_path):
    cut = _png(40, 40 * 22 + 6)                    # the last text line runs off the bottom
    blank = _png(3, 900)                           # 3 lines on a 900 px canvas
    clean = _png(12, 12 * 22 + 30)                 # trimmed, white below the last line
    rule = _png(12, 12 * 22 + 20, rule_at_bottom=True)   # ends on a table's bottom border
    path = _docx_with([(cut, 6.3), (blank, 6.3), (clean, 6.3), (rule, 6.3)],
                      str(tmp_path / 'pics.docx'))
    flags = pc.docx_picture_flags(path)
    got = sorted((f['type'], f['detail'].split(':')[0]) for f in flags)
    assert got == [('picture_mostly_blank', "picture 2 under '2 Section'"),
                   ('picture_truncated', "picture 1 under '1 Section'")], flags
    # pages are stamped from Word's layout when it lists exactly the pictures found
    stamped = pc.docx_picture_flags(path, pic_pages=[3, 5, 6, 7])
    assert sorted(f['page'] for f in stamped) == [3, 5]
    assert all(f['page'] == 0 for f in pc.docx_picture_flags(path, pic_pages=[3]))
    assert {'picture_truncated', 'picture_mostly_blank'} <= set(pc.DEFECT_TYPES)


# ── Chrome: the old screenshot vs the new page-sized slices ───────────────────────
@pytest.mark.skipif(_chrome() is None, reason='no Chrome / Chromium on this machine')
def test_old_screenshot_is_cut_or_blank_new_slices_are_complete_and_fit_the_page(tmp_path):
    chrome = _chrome()
    # BEFORE — the fixed 920 x 1400 screenshot: a long section is cut, a short one blank
    old = [DR._rasterize_section(h, _CSS, 'light', chrome) for h in (_long_section(), _short_section())]
    assert all(old), 'Chrome did not produce the screenshots'
    before = pc.docx_picture_flags(_docx_with([(p, 6.3) for p in old], str(tmp_path / 'old.docx')))
    assert sorted(f['type'] for f in before) == ['picture_mostly_blank', 'picture_truncated'], before

    # AFTER — the Studio's own Word export of the same two sections
    out = str(tmp_path / 'studio.docx')
    DR.build_docx(out, 'Monthly Progress Report', {'project_name': 'Synthetic Terminal',
                                                   'data_date': '2026-07-19'},
                  [_item(1, _long_section()), _item(2, _short_section())], chrome=chrome)
    assert pc.docx_picture_flags(out) == []
    pics = _pictures(out)
    long_pics = [h for h, head in pics if 'Reused section 1' in head]
    short_pics = [h for h, head in pics if 'Reused section 2' in head]
    assert len(long_pics) >= 3, pics            # 140 rows = several pages, nothing cut away
    assert len(short_pics) == 1 and short_pics[0] < 120, pics   # trimmed to its content
    from docx import Document
    room = DR._page_room_pt(Document(out))
    assert all(h <= room + 0.5 for h in long_pics + short_pics), (room, pics)
    # the first slice sits under the numbered heading + feature caption on one page
    assert long_pics[0] <= room - DR._LEAD_PT + 0.5 or long_pics[0] <= room * 0.5 + 0.5


@pytest.mark.skipif(_chrome() is None or not pc.word_available(),
                    reason='needs Chrome and Microsoft Word (COM)')
def test_word_lays_out_the_sliced_sections_with_zero_flags(tmp_path):
    out = str(tmp_path / 'studio_word.docx')
    DR.build_docx(out, 'Monthly Progress Report', {'project_name': 'Synthetic Terminal',
                                                   'data_date': '2026-07-19'},
                  [_item(1, _short_section()), _item(2, _long_section()),
                   _item(3, _short_section())], chrome=_chrome())
    res = pc.check_docx(out, engine='word', timeout=600)
    assert res['status'] == 'ok', res['message']
    assert res['flag_count'] == 0, res['flags']
