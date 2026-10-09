"""The final report HTML → a REAL Word document (.docx) that matches the PDF.

Built from the SAME HTML string the preview shows and Chrome prints (via the neutral
block model of :mod:`p6_export.html_model`), so Word carries exactly the ticked parts,
in the chosen order — ALWAYS in the standard LIGHT style (owner decision: the appearance
mode is a screen + PDF choice; the Word page is never dark — see
:func:`report_theme.force_light`), with the same sections, tables, values and formats:

* page size / orientation / margins from the report's ``@page``;
* a white page (a report that paints a non-white page WITHOUT the theme tokens keeps it);
* a running header (app · feature · project) and a "Page X of Y" footer;
* headings kept with what follows them (no orphaned titles), paragraphs of styled runs,
  bullet / numbered lists;
* tables with the header row repeated on every page, concrete cell colours, column widths
  from the report, merged cells;
* the shared Word pagination rules (:mod:`p6_export.docx_pagination`): small tables kept
  whole, long tables never strand 1-2 rows, headings / intros / captions kept with the
  first block;
* KPI / stat tiles → a small grid table (label · value · note);
* charts → PNG pictures (:mod:`p6_export.svg_raster`), or — when no picture can be made —
  the numbers behind the chart as a table;
* ``<img>`` data URIs embedded.

Every colour written is a concrete hex — Word ignores ``var(--x)``.
"""
import io
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor

from . import docx_pagination
from . import html_model as HM

_ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
          'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
_HEX_RE = re.compile(r'^[0-9A-Fa-f]{6}$')
_BAD_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _hex(v):
    if not v:
        return None
    v = str(v).lstrip('#')
    return v.upper() if _HEX_RE.match(v) else None


def _clean(t):
    return _BAD_XML.sub('', t or '')


def _lum(hexv):
    h = _hex(hexv) or 'FFFFFF'
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _solid_png(hexv, size=8):
    """A tiny opaque PNG of one colour (stdlib only) — stretched to the page it is the page colour."""
    import struct
    import zlib
    rgb = bytes.fromhex(_hex(hexv) or 'FFFFFF')
    raw = b''.join(b'\x00' + rgb * size for _ in range(size))

    def chunk(tag, data):
        return (struct.pack('>I', len(data)) + tag + data
                + struct.pack('>I', zlib.crc32(tag + data) & 0xFFFFFFFF))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', size, size, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def _page_fill(paragraph, hexv, width_emu, height_emu):
    """Paint the page colour as a full-page picture anchored BEHIND the text, from the header, so it
    repeats on every page. Word does not print ``w:background`` (nor keep it in Save as PDF) by
    default — on a dark theme that would leave near-white ink on white paper. A picture prints."""
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(_solid_png(hexv)), width=Emu(width_emu), height=Emu(height_emu))
    inline = run._r.find('.//' + qn('wp:inline'))
    anchor = OxmlElement('wp:anchor')
    for k, v in (('distT', '0'), ('distB', '0'), ('distL', '0'), ('distR', '0'), ('simplePos', '0'),
                 ('relativeHeight', '0'), ('behindDoc', '1'), ('locked', '1'),
                 ('layoutInCell', '1'), ('allowOverlap', '1')):
        anchor.set(k, v)
    sp = OxmlElement('wp:simplePos')
    sp.set('x', '0')
    sp.set('y', '0')
    anchor.append(sp)
    for tag, rel in (('wp:positionH', 'page'), ('wp:positionV', 'page')):
        pos = OxmlElement(tag)
        pos.set('relativeFrom', rel)
        off = OxmlElement('wp:posOffset')
        off.text = '0'
        pos.append(off)
        anchor.append(pos)
    kids = {c.tag: c for c in inline}
    anchor.append(kids[qn('wp:extent')])
    eff = kids.get(qn('wp:effectExtent'))
    if eff is None:
        eff = OxmlElement('wp:effectExtent')
        for k in ('l', 't', 'r', 'b'):
            eff.set(k, '0')
    anchor.append(eff)
    anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ('wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
        if kids.get(qn(tag)) is not None:
            anchor.append(kids[qn(tag)])
    doc_pr = anchor.find(qn('wp:docPr'))
    if doc_pr is not None:              # header ids restart at 1 → keep clear of body picture ids
        doc_pr.set('id', '9000')
        doc_pr.set('name', 'Page colour')
    inline.getparent().replace(inline, anchor)
    return anchor


# ── low-level XML helpers ──────────────────────────────────────────────────────
def _shade(pr, fill):
    fill = _hex(fill)
    if not fill:
        return
    for old in pr.findall(qn('w:shd')):
        pr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def _cell_shade(cell, fill):
    _shade(cell._tc.get_or_add_tcPr(), fill)


def _borders(tag, spec):
    """spec: {side: (size_eighths_pt, 'RRGGBB') | None}"""
    el = OxmlElement(tag)
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if side not in spec:
            continue
        b = OxmlElement(f'w:{side}')
        v = spec[side]
        if not v or not _hex(v[1]):
            b.set(qn('w:val'), 'nil')
        else:
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), str(max(2, int(v[0]))))
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), _hex(v[1]))
        el.append(b)
    return el


def _cell_borders(cell, spec):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(old)
    tcPr.append(_borders('w:tcBorders', spec))


def _cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for side, v in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tcPr.append(mar)


def _cell_width(cell, emu):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcW')):
        tcPr.remove(old)
    w = OxmlElement('w:tcW')
    w.set(qn('w:w'), str(int(emu / 635)))          # EMU → twentieths of a point
    w.set(qn('w:type'), 'dxa')
    tcPr.insert(0, w)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader')
    h.set(qn('w:val'), 'true')
    trPr.append(h)


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))


def _fixed_layout(table):
    tblPr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)


def _para_border(p, side, width_pt, color):
    color = _hex(color)
    if not color:
        return
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        pPr.append(bdr)
    b = OxmlElement(f'w:{side}')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(max(2, int(round((width_pt or 0.75) * 8)))))
    b.set(qn('w:space'), '4')
    b.set(qn('w:color'), color)
    bdr.append(b)


def _field(run, instr):
    for kind, text in (('begin', None), (None, instr), ('separate', None), (None, '1'), ('end', None)):
        if kind:
            f = OxmlElement('w:fldChar')
            f.set(qn('w:fldCharType'), kind)
            run._r.append(f)
        elif text == instr:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = f' {instr} '
            run._r.append(it)
        else:
            t = OxmlElement('w:t')
            t.text = text
            run._r.append(t)


# OOXML is order-sensitive inside *Pr elements — Word rejects a file whose children are
# out of schema order. Everything above appends freely; _normalize_order() fixes it once.
_SETTINGS_BEFORE_BG = ('writeProtection', 'view', 'zoom', 'removePersonalInformation',
                       'removeDateAndTime', 'doNotDisplayPageBoundaries')
_ORDER = {
    'pPr': ('pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
            'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
            'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN',
            'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind', 'contextualSpacing',
            'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
            'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'),
    'rPr': ('rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike',
            'outline', 'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish',
            'webHidden', 'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs', 'highlight',
            'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang',
            'eastAsianLayout', 'specVanish', 'oMath'),
    'tcPr': ('cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap',
             'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'),
    'tblPr': ('tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
              'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd',
              'tblLayout', 'tblCellMar', 'tblLook'),
    'trPr': ('cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit',
             'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden'),
}


def _normalize_order(root):
    rank = {qn('w:' + k): {qn('w:' + n): i for i, n in enumerate(v)} for k, v in _ORDER.items()}
    for el in root.iter(*rank.keys()):
        order = rank[el.tag]
        kids = list(el)
        if len(kids) < 2:
            continue
        # keep only the last of duplicated singletons (e.g. two tcMar)
        seen, uniq = set(), []
        for ch in reversed(kids):
            if ch.tag in seen and ch.tag in order:
                continue
            seen.add(ch.tag)
            uniq.append(ch)
        uniq.reverse()
        srt = sorted(uniq, key=lambda ch: order.get(ch.tag, len(order)))
        if srt != kids:
            for ch in kids:
                el.remove(ch)
            for ch in srt:
                el.append(ch)


# ── the writer ─────────────────────────────────────────────────────────────────
def _fit_widths(widths, mins):
    """Column widths with every column at least ``mins`` wide, the total unchanged.  The extra is
    taken from the columns that are wider than they need, in proportion to their spare room.  When
    the needs do not fit the page, every column gives up the same share of its need."""
    total = sum(widths)
    if not widths:
        return widths
    if sum(mins) > total:                       # not enough room for all: share the shortage
        k = total / float(sum(mins))
        mins = [int(m * k) for m in mins]
    need = sum(max(0, m - w) for w, m in zip(widths, mins))
    spare = sum(max(0, w - m) for w, m in zip(widths, mins))
    if need <= 0 or spare <= 0:
        return widths
    out = []
    for w, m in zip(widths, mins):
        out.append(m if w < m else w - int((w - m) * need / spare))
    spare_at = max(range(len(out)), key=lambda i: out[i] - mins[i])     # rounding goes where there is room
    out[spare_at] += total - sum(out)
    return out


class _Writer:
    def __init__(self, rep, app_name='', feature='', project=''):
        self.rep = rep
        self.doc = Document()
        self.app_name = app_name
        self.feature = feature
        self.project = project
        self.ink = _hex(rep.ink) or '000000'
        self.bg = _hex(rep.bg) or 'FFFFFF'
        self.dark = _lum(self.bg) < 110
        theme = rep.theme or {}
        self.hair = _hex(theme.get('rpt-hair')) or ('3A4656' if self.dark else 'D5DCE6')
        self.muted = _hex(theme.get('rpt-muted')) or ('9AA7B8' if self.dark else '6B7686')
        self.accent = _hex(theme.get('rpt-accent')) or '1F4E79'
        self.base_pt = max(7.0, min(float(rep.base_pt or 8.25), 12.0))
        self._last_break = True
        self._setup()

    # page, fonts, background, header / footer
    def _setup(self):
        doc, rep = self.doc, self.rep
        page = rep.page or {}
        sec = doc.sections[0]
        w, h = page.get('width_mm') or 210.0, page.get('height_mm') or 297.0
        sec.orientation = WD_ORIENT.LANDSCAPE if w > h else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Mm(w), Mm(h)
        t, r, b, l = page.get('margins_mm') or (15, 12, 15, 12)
        # the running header/footer sit inside the top/bottom margins → keep room for them
        sec.top_margin, sec.bottom_margin = Mm(max(t, 16)), Mm(max(b, 14))
        sec.left_margin, sec.right_margin = Mm(l), Mm(r)
        sec.header_distance, sec.footer_distance = Mm(6), Mm(6)
        self.content_emu = sec.page_width - sec.left_margin - sec.right_margin
        self.max_h_emu = int(sec.page_height - sec.top_margin - sec.bottom_margin - Mm(12))
        normal = doc.styles['Normal']
        normal.font.name = rep.font_family or 'Segoe UI'
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.append(fonts)
        for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            fonts.set(qn(k), rep.font_family or 'Segoe UI')
        normal.font.size = Pt(self.base_pt)
        normal.font.color.rgb = RGBColor.from_string(self.ink)
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.line_spacing = 1.1
        for i in range(1, 7):
            try:
                hs = doc.styles[f'Heading {i}']
            except KeyError:
                continue
            hs.font.name = rep.font_family or 'Segoe UI'
            hs.font.color.rgb = RGBColor.from_string(self.accent)
            hs.paragraph_format.keep_with_next = True
            hs.paragraph_format.space_before = Pt(10 if i <= 2 else 6)
            hs.paragraph_format.space_after = Pt(4)
            hrpr = hs.element.get_or_add_rPr()
            hf = hrpr.find(qn('w:rFonts'))
            if hf is None:
                hf = OxmlElement('w:rFonts')
                hrpr.append(hf)
            for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
                hf.set(qn(k), rep.font_family or 'Segoe UI')
            for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
                if hf.get(qn(k)) is not None:
                    del hf.attrib[qn(k)]
        if self.bg != 'FFFFFF':
            bgel = OxmlElement('w:background')
            bgel.set(qn('w:color'), self.bg)
            doc.element.insert(0, bgel)
            settings = doc.settings.element
            if settings.find(qn('w:displayBackgroundShape')) is None:
                idx = 0
                for i, ch in enumerate(list(settings)):
                    if ch.tag in {qn('w:' + n) for n in _SETTINGS_BEFORE_BG}:
                        idx = i + 1
                settings.insert(idx, OxmlElement('w:displayBackgroundShape'))
        # running header: app — feature · project ; footer: Page X of Y
        bits = [x for x in (self.app_name, self.feature) if x]
        head = ' — '.join(bits)
        if self.project:
            head = f'{head} · {self.project}' if head else self.project
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(_clean(head))
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(self.muted)
        _para_border(hp, 'bottom', 0.5, self.hair)
        if self.bg != 'FFFFFF':
            # w:background shows on screen only; this prints (and survives Word's Save as PDF)
            _page_fill(hp, self.bg, sec.page_width, sec.page_height)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for txt, instr in (('Page ', None), (None, 'PAGE'), (' of ', None), (None, 'NUMPAGES')):
            r = fp.add_run(txt or '')
            r.font.size = Pt(7.5)
            r.font.color.rgb = RGBColor.from_string(self.muted)
            if instr:
                _field(r, instr)

    # runs
    def _run(self, p, r, base_pt=None, force_bold=None, force_color=None):
        if r.br:
            p.add_run().add_break()
            return
        text = _clean(r.text)
        if not text:
            return
        run = p.add_run(text)
        f = run.font
        if r.bold or force_bold:
            f.bold = True
        if r.italic:
            f.italic = True
        if r.underline:
            f.underline = True
        if r.strike:
            f.strike = True
        if r.caps:
            f.all_caps = True
        if r.sup:
            f.superscript = True
        if r.sub:
            f.subscript = True
        size = r.size_pt or base_pt
        if size:
            f.size = Pt(max(5.0, min(float(size), 40.0)))
        col = _hex(force_color) or _hex(r.color)
        if col:
            f.color.rgb = RGBColor.from_string(col)
        if r.bg and _hex(r.bg):
            _shade(run._r.get_or_add_rPr(), r.bg)
        if r.font and r.font.lower() not in ('segoe ui', 'arial', 'sans-serif', 'system-ui'):
            f.name = r.font

    def _runs(self, p, runs, base_pt=None, **kw):
        for r in runs:
            self._run(p, r, base_pt, **kw)

    def _para_fmt(self, p, blk):
        pf = p.paragraph_format
        p.alignment = _ALIGN.get(getattr(blk, 'align', 'left'), WD_ALIGN_PARAGRAPH.LEFT)
        if getattr(blk, 'keep_with_next', False):
            pf.keep_with_next = True
        bg = getattr(blk, 'bg', None)
        if bg and _hex(bg) and _hex(bg) != self.bg:
            _shade(p._p.get_or_add_pPr(), bg)
        bl = getattr(blk, 'border_left', None)
        if bl:
            _para_border(p, 'left', bl[0], bl[1])
        bb = getattr(blk, 'border_bottom', None)
        if bb:
            _para_border(p, 'bottom', bb[0], bb[1])

    # blocks
    def emit(self, blk, container=None):
        c = container or self.doc
        k = blk.kind
        if k == 'pagebreak':
            if container is None and not self._last_break:
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                self._last_break = True
            return
        if k == 'section':
            for b in blk.blocks:
                self.emit(b, container)
            return
        self._last_break = False
        if k == 'heading':
            lvl = max(1, min(int(blk.level or 2), 6))
            p = c.add_paragraph(style=f'Heading {lvl}') if container is None else c.add_paragraph()
            self._runs(p, blk.runs, blk.size_pt)
            self._para_fmt(p, blk)
            p.paragraph_format.keep_with_next = True
        elif k == 'paragraph':
            p = c.add_paragraph()
            self._runs(p, blk.runs, blk.size_pt)
            self._para_fmt(p, blk)
            if blk.role == 'pre':
                for r in p.runs:
                    r.font.name = 'Consolas'
        elif k == 'list':
            for i, item in enumerate(blk.items):
                style = 'List Number' if blk.ordered else 'List Bullet'
                try:
                    p = c.add_paragraph(style=style)
                except KeyError:
                    p = c.add_paragraph()
                    p.add_run(f'{i + 1}. ' if blk.ordered else '• ')
                self._runs(p, item, blk.size_pt)
        elif k == 'table':
            self.table(blk, c)
        elif k == 'kpis':
            self.kpis(blk, c)
        elif k == 'image':
            self.picture(blk.data, blk.width_px, c, blk.height_px)
        elif k == 'visual':
            self.visual(blk, c)

    def _space(self, c, pts=3):
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run('')
        r.font.size = Pt(pts)
        return p

    def picture(self, data, width_px, c, height_px=None):
        if not data:
            return
        max_emu = int(self.content_emu)
        width = Emu(min(max_emu, int((width_px or 600) * 9525))) if width_px else Emu(max_emu)
        if width_px and height_px and self.max_h_emu:
            h = width * height_px / width_px            # keep the aspect; never taller than a page
            if h > self.max_h_emu:
                width = Emu(int(width * self.max_h_emu / h))
        p = c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p.add_run().add_picture(io.BytesIO(data), width=width)
        except Exception:
            p.add_run('[picture]')

    def visual(self, v, c):
        if v.vectors and self._native(v, c):                 # real Word shapes + text boxes (comment 41)
            return
        if v.slices:
            for png, w, h in v.slices:
                self.picture(png, w, c, h)
            return
        if v.png:
            self.picture(v.png, v.width_px, c, v.height_px)
            return
        if v.data_headers and v.data_rows:
            rows = [[HM.Cell(runs=[HM.Run(text=str(h), bold=True)], header=True) for h in v.data_headers]]
            for r in v.data_rows:
                rows.append([HM.Cell(runs=[HM.Run(text='' if x is None else str(x))]) for x in r])
            self.table(HM.Table(rows=rows, header_rows=1), c)
            return
        for line in v.text_lines or []:
            p = c.add_paragraph()
            p.add_run(_clean(line))

    def _native(self, v, c):
        """The chart as groups of native Word shapes read from the PDF's own drawing — every
        bar and every label can be selected and edited.  False → nothing was written."""
        from . import vector_shapes as VS
        try:
            ids = self.__dict__.setdefault('_shape_ids', [50000])
            max_w = int(self.content_emu if c is self.doc else self.content_emu * 0.96)
            xmls = []
            for prims, w_pt, h_pt in v.vectors:
                x = VS.group_drawing_xml(prims, w_pt, h_pt, ids, max_w_emu=max_w,
                                         max_h_emu=self.max_h_emu or None, name=v.title or 'Chart')
                if not x:
                    return False
                xmls.append(x)
            for x in xmls:
                p = c.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                p.add_run()._r.append(parse_xml(x))
            return bool(xmls)
        except Exception:
            if os.environ.get('CX_EXPORT_DEBUG'):
                raise
            return False

    def table(self, t, c):
        rows = [r for r in t.rows if r]
        if not rows:
            return
        ncols = max(1, t.ncols)
        # occupancy grid for row/col spans
        grid = []
        placed = []                       # (r, c, rowspan, colspan, cell)
        for ri, row in enumerate(rows):
            while len(grid) <= ri:
                grid.append([None] * ncols)
            ci = 0
            for cell in row:
                while ci < ncols and grid[ri][ci] is not None:
                    ci += 1
                if ci >= ncols:
                    break
                cs = max(1, min(cell.colspan, ncols - ci))
                rs = max(1, min(cell.rowspan, len(rows) - ri))
                for dr in range(rs):
                    while len(grid) <= ri + dr:
                        grid.append([None] * ncols)
                    for dc in range(cs):
                        grid[ri + dr][ci + dc] = cell
                placed.append((ri, ci, rs, cs, cell))
                ci += cs
        nrows = len(grid)
        table = c.add_table(rows=nrows, cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _fixed_layout(table)
        weights = list(t.col_weights or [])
        if len(weights) != ncols or not all(isinstance(w, (int, float)) and w > 0 for w in weights):
            weights = [1.0] * ncols
        total = sum(weights)
        widths = [int(self.content_emu * w / total) for w in weights]
        base = t.size_pt or self.base_pt
        # a column is never narrower than its longest single word: "10" must not wrap to "1 / 0",
        # nor an Activity ID onto two lines (the room is taken from the columns that have spare)
        longest = [0.0] * ncols
        for ri, ci, rs, cs, cell in placed:
            if cs != 1:
                continue
            for r in cell.runs:
                if r.br:
                    continue
                pt = r.size_pt or cell.size_pt or base or 9
                mono = any(k in (r.font or '').lower() for k in ('mono', 'consol', 'courier'))
                per = 0.61 if mono else (0.56 if (cell.header or cell.bold or r.bold) else 0.5)
                for word in _clean(r.text or '').split():
                    longest[ci] = max(longest[ci], min(len(word), 40) * pt * per)
        widths = _fit_widths(widths, [int(w * 12700) + 101600 + 25400 for w in longest])
        hair = _hex(t.border_color) or self.hair
        side = (4, hair) if getattr(t, 'grid', False) else None     # full grid when the report boxes its cells
        tblPr = table._tbl.tblPr
        tblPr.append(_borders('w:tblBorders', {
            'top': (4, hair), 'bottom': (4, hair), 'left': side, 'right': side,
            'insideH': (4, hair), 'insideV': side}))
        trs = table.rows
        for ri in range(nrows):
            if ri < t.header_rows:
                _repeat_header(trs[ri])
            if nrows < 60 or ri < t.header_rows:
                _cant_split(trs[ri])
        cells_by_row = [trs[ri].cells for ri in range(nrows)]
        for ri in range(nrows):
            for ci in range(ncols):
                _cell_width(cells_by_row[ri][ci], widths[ci])
        for ri, ci, rs, cs, cell in placed:
            dcell = cells_by_row[ri][ci]
            if rs > 1 or cs > 1:
                dcell = dcell.merge(cells_by_row[ri + rs - 1][ci + cs - 1])
            _cell_margins(dcell)
            if cell.bg:
                _cell_shade(dcell, cell.bg)
            p = dcell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = _ALIGN.get(cell.align, WD_ALIGN_PARAGRAPH.LEFT)
            first = True
            chunks = [[]]
            for r in cell.runs:
                if r.br:
                    chunks.append([])
                else:
                    chunks[-1].append(r)
            for chunk in chunks:
                if not first:
                    p = dcell.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.alignment = _ALIGN.get(cell.align, WD_ALIGN_PARAGRAPH.LEFT)
                first = False
                for r in chunk:
                    self._run(p, r, cell.size_pt or base,
                              force_bold=cell.bold or cell.header or None,
                              force_color=None if r.color else cell.color)
            if cell.bar:
                width = sum(widths[ci:ci + cs]) - 2 * 80 * 635          # less the cell margins
                self._bar(p, cell.bar, width / 12700.0)
        self._space(c, 4)

    def _bar(self, p, bar, w_pt):
        """A Gantt time-line cell as native Word shapes the width of its column: the track,
        the bar with its % complete fill (or the milestone diamond) and the data-date line —
        or, in the header, the month labels with their tick marks.  As the PDF draws them."""
        from . import vector_shapes as VS
        w_pt -= VS.RIGHT_PAD                       # group_drawing_xml adds this room back
        if w_pt < 12:
            return
        x = lambda pct: w_pt * max(0.0, min(100.0, pct)) / 100.0

        def rect(x0, y0, x1, y1, fill):
            return {'k': 'path', 'segs': [], 'fill': fill, 'stroke': None, 'width': 0, 'rect': True,
                    'bbox': (x0, y0, max(x1, x0 + 0.75), y1)}
        prims = []
        if 'scale' in bar:
            rows = bar.get('scale_rows') or [0] * len(bar['scale'])
            row_h = 8.4                                  # one line of the scale (months alternate lines, years have their own)
            h = row_h * (max(rows) + 1) + 1.0
            for (left, label, col), r in zip(bar['scale'], rows):
                x0, y0 = x(left), 0.5 + r * row_h
                prims.append(rect(x0, y0, x0 + 0.75, h, self.hair))
                size = 6.6
                is_year = label.isdigit() and len(label) == 4
                prims.append({'k': 'text', 'text': label, 'size': size, 'color': _hex(col) or self.muted,
                              'bold': is_year, 'italic': False, 'font': 'Consolas', 'vert': False,
                              'base': 7.2, 'bbox': (x0 + 1.5, y0, x0 + 1.5 + len(label) * size * 0.6, y0 + row_h)})
        else:
            h = 9.0
            prims.append(rect(0, 0, w_pt, h, _hex(bar.get('track')) or 'F1F4F8'))
            if bar.get('bar'):
                left, width, col = bar['bar']
                x0, x1 = x(left), x(min(100.0, left + width))
                prims.append(rect(x0, 1.5, max(x1, x0 + 1.5), 7.5, _hex(col) or 'D6E4F5'))
                f = bar.get('fill')
                if f and f[0] > 0:
                    prims.append(rect(x0, 1.5, x0 + (max(x1, x0 + 1.5) - x0) * f[0] / 100.0, 7.5,
                                      _hex(f[1]) or self.accent))
            if bar.get('ms'):
                cx, col = x(bar['ms'][0]), _hex(bar['ms'][1]) or self.ink
                r = 4.0
                prims.append({'k': 'path', 'fill': col, 'stroke': None, 'width': 0, 'rect': False,
                              'segs': [('M', (cx, 4.5 - r)), ('L', (cx + r, 4.5)), ('L', (cx, 4.5 + r)),
                                       ('L', (cx - r, 4.5)), ('Z',)],
                              'bbox': (cx - r, 4.5 - r, cx + r, 4.5 + r)})
            if bar.get('dd'):
                dx = x(bar['dd'][0])
                prims.append(rect(dx - 0.375, 0, dx + 0.375, h, _hex(bar['dd'][1]) or self.accent))
        try:
            ids = self.__dict__.setdefault('_shape_ids', [50000])
            xml = VS.group_drawing_xml(prims, w_pt, h, ids, name='Gantt bar')
            if xml:
                p.add_run()._r.append(parse_xml(xml))
        except Exception:
            if os.environ.get('CX_EXPORT_DEBUG'):
                raise

    def kpis(self, g, c):
        tiles = [t for t in g.tiles if (t.label or t.value or t.note)]
        if not tiles:
            return
        cols = max(1, min(int(g.columns or 4), len(tiles), 6))
        nrows = (len(tiles) + cols - 1) // cols
        table = c.add_table(rows=nrows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _fixed_layout(table)
        gap_col = self.bg
        tblPr = table._tbl.tblPr
        tblPr.append(_borders('w:tblBorders', {
            'top': None, 'bottom': None, 'left': None, 'right': None,
            'insideH': (24, gap_col), 'insideV': (24, gap_col)}))
        width = int(self.content_emu / cols)
        for i in range(nrows * cols):
            cell = table.cell(i // cols, i % cols)
            _cell_width(cell, width)
            if i >= len(tiles):
                _cell_borders(cell, {'top': None, 'bottom': None, 'left': None, 'right': None})
                continue
            tile = tiles[i]
            edge = _hex(tile.border) or self.hair
            left = (18, _hex(tile.accent)) if _hex(tile.accent) else (6, edge)
            _cell_borders(cell, {'top': (6, edge), 'bottom': (6, edge), 'right': (6, edge), 'left': left})
            _cell_margins(cell, 70, 70, 110, 90)
            if tile.bg and _hex(tile.bg) != self.bg:
                _cell_shade(cell, tile.bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            if tile.label:
                r = p.add_run(_clean(tile.label).upper())
                r.font.size = Pt(tile.label_pt or max(6.0, self.base_pt - 1.5))
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(_hex(tile.label_color) or self.muted)
            if tile.value:
                p2 = cell.add_paragraph()
                p2.paragraph_format.space_after = Pt(1)
                r = p2.add_run(_clean(tile.value))
                r.font.size = Pt(min(tile.value_pt or 14.0, 22.0))
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(_hex(tile.value_color) or self.ink)
            if tile.note:
                p3 = cell.add_paragraph()
                p3.paragraph_format.space_after = Pt(0)
                r = p3.add_run(_clean(tile.note))
                r.font.size = Pt(max(6.0, self.base_pt - 1.5))
                r.font.color.rgb = RGBColor.from_string(_hex(tile.note_color) or self.muted)
        self._space(c, 4)

    def write(self, path):
        rep = self.rep
        for b in rep.front:
            self.emit(b)
        for s in rep.sections:
            self.emit(s)
        for b in rep.tail:
            self.emit(b)
        self.doc.core_properties.title = _clean(rep.title or self.feature or '')
        if self.app_name:
            self.doc.core_properties.author = self.app_name
        docx_pagination.paginate_docx(self.doc)      # the shared Word page-composition rules
        _normalize_order(self.doc.element)
        for part in (self.doc.sections[0].header, self.doc.sections[0].footer):
            _normalize_order(part._element)
        _normalize_order(self.doc.styles.element)
        self.doc.save(path)
        return path


def build_docx(rep, path, app_name='', feature='', project=''):
    """Write a parsed :class:`html_model.ReportDoc` (visuals already rasterised) to ``path``."""
    return _Writer(rep, app_name=app_name, feature=feature, project=project).write(path)


def html_to_docx(html, path, app_name='', feature='', project='', chrome=None,
                 use_chrome=True, sections=None):
    """The one-call path used by ``POST /api/export/docx``. Whatever appearance mode the
    preview is in, Word is written in the standard light style (report_theme.force_light)."""
    import report_theme
    from . import svg_raster
    rep = HM.parse_report(report_theme.force_light(html), sections=sections)
    svg_raster.rasterize(rep, chrome=chrome, use_chrome=use_chrome)
    project = project or (rep.meta or {}).get('project', '')
    return build_docx(rep, path, app_name=app_name, feature=feature or rep.title, project=project)
