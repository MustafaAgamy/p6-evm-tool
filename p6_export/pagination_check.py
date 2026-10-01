"""Automatic page checker (owner point 14) — finds page-composition defects in a FINISHED
report, PDF or Word, so every feature team can prove its report is laid out properly.

    python -m p6_export.pagination_check report.pdf
    python -m p6_export.pagination_check report.docx [--engine auto|word|spire]
    python -m p6_export.pagination_check report.pdf --html report.html   (heading hints)

prints ONE JSON object: ``{file, format, engine, status, pages, flag_count, counts, flags,
info, message}``; each flag is ``{type, page, detail}`` (``page`` is 1-based, the page the
defect is on — for a break, the page BEFORE it). Exit code: 0 clean · 1 flags found ·
2 error · 3 skipped (no renderer for a Word file on this machine).

Defects (``flags``):

  orphaned_heading              a heading is the last thing on a page; its content starts
                                on the next page
  kpi_separated_from_heading    a heading ends a page and its KPI / summary cards start the
                                next page
  heading_separated_from_block  a heading (+ a short intro) ends a page and its table /
                                chart / picture / cards start the next page
  picture_separated_from_caption (Word) a picture starts a page, its label / caption is on
                                the page before
  table_split_few_rows          a table breaks across pages leaving only 1-2 body rows on
                                one side (the rules keep >= 3)
  small_table_split             a table small enough to be kept whole (<= FIT of a page) is
                                split across two pages
  table_header_not_repeated     a table continues on the next page without its header row
  empty_repeated_header         (Word) a side-by-side part of a table (two tables built as
                                one, a blank column between) repeats its header on a page
                                where it has no rows left
  graphic_cut                   a chart / diagram / picture is cut by the page break (part
                                on each page, or clipped by the sheet edge)
  text_cut                      text runs off the sheet (content overflowed the page), or
                                (Word) a table cell's text is clipped by an EXACT row height
  content_in_margin             a page's content starts inside the top margin (the page lost
                                its margin / running header — an overflow page)
  large_blank_then_continuation a page ends more than 40 % blank and the next page goes on
                                with a pushed block (not a new section)
  stranded_fragment             a page holds only a small tail of a chart / cards block
  empty_page                    a page with nothing on it besides the running header/footer
  picture_truncated             (Word) a picture's content runs into its bottom edge - a
                                report section embedded as a picture was cut off
  picture_mostly_blank          (Word) a picture is mostly white (< 50 % painted, >= 2 in
                                of white on the page)

Informational (``info``, never counted): ``section_break_blank`` — a page ends early because
the next page starts a new top-level section, or it is the cover page (page 1, its content
set well down the page), or the contents list ends on it and the body starts on the next
page (page breaks by design).

How a PDF is read (PyMuPDF): text lines with their font size / bold, vector drawings and
pictures per page (only those that intersect the page: Chrome emits thousands of off-page
paths per sheet, y down to -30 000; a page frame / page background box is not content).
The running header / footer / page frame (the same thing at the same place on most pages)
and page numbers are set aside. Headings are found WITHOUT the source:
a short line alone on its row that is bold (or clearly larger than the body text) — the
same "heading-like" definition the print-time composer uses; the renderers' heading texts
can be passed as hints (``--html`` / ``headings=``). Tables are found as runs of rows of
side-by-side cells sharing column edges (PyMuPDF's table finder mistakes page frames for
tables); a table that continues across a break is matched by its column edges.

How a Word file is read: Word's OWN pagination through COM (Word installed + pywin32):
the document is opened read-only, repaginated, and the page + vertical position of every
paragraph, table row and picture is read (an inline picture's height, or the height of a
floating picture anchored to an empty / "/" paragraph - the Narrative WBS trees) — no PDF
export (that hangs when Word runs headless). Fallback: Spire.Doc renders the file to PDF
(its free edition converts only the first 10 pages - the result says so) and the PDF rules
run on that, with the document's heading texts as hints.
When neither is available the check is skipped with a clear message (status "skipped").

Thresholds are the SAME numbers the pagination rules use (report_theme.PAGINATION_*).
"""
import argparse
import collections
import html as _html
import json
import os
import re
import subprocess
import sys
import tempfile

try:                                   # the same thresholds as the rules the renderers follow
    import report_theme as _rt
    FIT = float(getattr(_rt, 'PAGINATION_FIT', 0.35))
    MIN_ROWS = int(getattr(_rt, 'PAGINATION_MIN_ROWS', 3))
    HEAD_MAX_CHARS = int(getattr(_rt, 'PAGINATION_HEAD_MAX_CHARS', 160))
except Exception:                      # pragma: no cover — report_theme always ships
    _rt = None
    FIT, MIN_ROWS, HEAD_MAX_CHARS = 0.35, 3, 160

# a page ending more than 40 % blank before a pushed block is a defect: the rules keep a
# block of up to FIT (35 %) of a page whole WITH its heading (~5 %), so pushing one leaves up
# to ~40 % blank by design (GBT Studio: a 33 % row of lag charts under its section title)
BLANK = FIT + 0.05
FRAGMENT = 0.12           # a page holding < 12 % of content (a figure / cards tail) is stranded
INTRO_MAX_PT = 60.0       # a heading's short intro (about 3-4 lines)
TOP_ZONE_PT = 26.0        # "what the page starts with" looks this far below the first item
COVER_DROP = 0.2          # page 1 whose content starts > 20 % of the page down is a cover
CONTENTS_ROWS = 3         # a page ending with >= 3 "title ... page-number" rows ends the contents
DEFECT_TYPES = (
    'orphaned_heading', 'kpi_separated_from_heading', 'heading_separated_from_block',
    'picture_separated_from_caption', 'table_split_few_rows', 'small_table_split',
    'table_header_not_repeated', 'empty_repeated_header', 'graphic_cut', 'text_cut',
    'content_in_margin', 'large_blank_then_continuation',
    'stranded_fragment', 'empty_page', 'picture_truncated', 'picture_mostly_blank',
    'block_split',
)
INFO_TYPES = ('section_break_blank',)

EXIT_CLEAN, EXIT_FLAGS, EXIT_ERROR, EXIT_SKIPPED = 0, 1, 2, 3


def norm(s):
    """Lower-case words only — for matching heading texts across renderers."""
    s = _html.unescape(s or '').replace('\xa0', ' ')
    return re.sub(r'[^0-9a-z]+', ' ', s.lower()).strip()


def _result(path, fmt, engine, pages, flags, info=(), status='ok', message=''):
    flags = sorted(flags, key=lambda f: (f['page'], f['type']))
    counts = collections.Counter(f['type'] for f in flags)
    return {'file': path, 'format': fmt, 'engine': engine, 'status': status,
            'pages': pages, 'flag_count': len(flags), 'counts': dict(sorted(counts.items())),
            'flags': flags, 'info': sorted(info, key=lambda f: f['page']), 'message': message}


def _flag(kind, page, detail):
    return {'type': kind, 'page': int(page), 'detail': detail}


# ══ PDF ══════════════════════════════════════════════════════════════════════════
class _Line:
    __slots__ = ('t', 'n', 'x0', 'y0', 'x1', 'y1', 'bold', 'size')

    def __init__(self, t, x0, y0, x1, y1, bold, size):
        self.t, self.n = t, norm(t)
        self.x0, self.y0, self.x1, self.y1 = x0, y0, x1, y1
        self.bold, self.size = bold, size


class _Band:
    """One row of text: the lines that sit side by side at the same height."""
    __slots__ = ('lines', 'x0', 'y0', 'x1', 'y1', 'heading', 'size')

    def __init__(self, lines):
        self.lines = sorted(lines, key=lambda l: l.x0)
        self.x0 = min(l.x0 for l in lines)
        self.x1 = max(l.x1 for l in lines)
        self.y0 = min(l.y0 for l in lines)
        self.y1 = max(l.y1 for l in lines)
        self.size = max(l.size for l in lines)
        self.heading = False

    @property
    def text(self):
        return ' '.join(l.t for l in self.lines)

    @property
    def grid(self):          # side-by-side cells (a table row / a row of cards)
        return len(self.lines) >= 2 and any(b.x0 - a.x1 > 3 for a, b in zip(self.lines, self.lines[1:]))

    @property
    def bold(self):
        return all(l.bold for l in self.lines)


class _Draw:
    __slots__ = ('x0', 'y0', 'x1', 'y1', 'ry0', 'ry1', 'w', 'h', 'curve', 'fill', 'thin', 'img')

    def __init__(self, x0, y0, x1, y1, ry0, ry1, curve=False, fill=None, img=False):
        self.x0, self.y0, self.x1, self.y1 = x0, y0, x1, y1
        self.ry0, self.ry1 = ry0, ry1
        self.w, self.h = x1 - x0, y1 - y0
        self.curve, self.fill, self.img = curve, fill, img
        self.thin = self.w < 2.5 or self.h < 2.5


class _Page:
    def __init__(self, no, W, H):
        self.no, self.W, self.H = no, W, H
        self.lines, self.draws, self.misc = [], [], []   # misc: rotated / tiny text extents
        self.bands, self.regions = [], []

    def items(self):
        return ([(b.y0, b.y1) for b in self.bands] + [(d.y0, d.y1) for d in self.draws]
                + list(self.misc))

    @property
    def empty(self):
        return not self.items()

    @property
    def top(self):
        return min(y0 for y0, _ in self.items())

    @property
    def bottom(self):
        return max(y1 for _, y1 in self.items())


def _is_bold(span):
    f = span.get('font') or ''
    return bool(span.get('flags', 0) & 16) or any(k in f for k in ('Bold', 'bold', 'Semibold', 'Black', 'Heavy'))


def _read_pdf(path):
    import pymupdf
    pages = []
    with pymupdf.open(path) as doc:
        for i, pg in enumerate(doc):
            W, H = pg.rect.width, pg.rect.height
            P = _Page(i + 1, W, H)
            for b in pg.get_text('dict')['blocks']:
                for ln in b.get('lines', []):
                    spans = [s for s in ln['spans'] if s['text'].strip()]
                    if not spans:
                        continue
                    x0, y0, x1, y1 = ln['bbox']
                    if y1 <= 0 or y0 >= H + 400:
                        continue
                    t = ' '.join(''.join(s['text'] for s in ln['spans']).split())
                    size = max(s['size'] for s in spans)
                    if size < 3 or t.startswith("Evaluation Warning"):
                        continue                  # invisible markers (a 1px section marker) / a renderer watermark
                    d = ln.get('dir', (1, 0))
                    if abs(d[0] - 1) > 0.01:
                        P.misc.append((max(y0, 0), min(y1, H)))   # rotated text (an axis label)
                        continue
                    P.lines.append(_Line(t, x0, y0, x1, y1, all(_is_bold(s) for s in spans), size))
            for d in pg.get_drawings():
                r = d['rect']
                if r.y1 <= 0.5 or r.y0 >= H - 0.5 or r.x1 <= 0 or r.x0 >= W:
                    continue                      # off-page paths (Chrome emits thousands)
                w, h = r.width, r.height
                fill, stroke = d.get('fill'), d.get('color')
                if w > 0.8 * W and h > 0.4 * H and (fill is None or (w > 0.95 * W and h > 0.9 * H)):
                    continue                      # page frame box / page background
                if (w < 2.5 and h > 0.25 * H) or (h < 2.5 and w > 0.85 * W):
                    continue                      # page frame lines
                if fill is not None and min(fill) > 0.985 and not stroke:
                    continue                      # white fills
                if fill is None and not stroke:
                    continue
                items = d.get('items') or []
                P.draws.append(_Draw(max(r.x0, 0), max(r.y0, 0), min(r.x1, W), min(r.y1, H), r.y0, r.y1,
                                     curve=any(it[0] == 'c' for it in items), fill=fill))
            for im in pg.get_image_info():
                x0, y0, x1, y1 = im['bbox']
                if y1 <= 0 or y0 >= H or x1 - x0 < 2 or y1 - y0 < 2:
                    continue
                P.draws.append(_Draw(max(x0, 0), max(y0, 0), min(x1, W), min(y1, H), y0, y1, img=True))
            pages.append(P)
    return pages


def _bands_of(lines):
    lines = sorted(lines, key=lambda l: (l.y0, l.x0))
    bands, cur = [], []
    for l in lines:
        if cur and (abs(l.y0 - cur[0].y0) <= 2.5 and l.y0 < max(c.y1 for c in cur)):
            cur.append(l)
        else:
            if cur:
                bands.append(_Band(cur))
            cur = [l]
    if cur:
        bands.append(_Band(cur))
    return bands


def _strip_running(pages):
    """Set aside the running header / footer / frame and page numbers: the same row of text
    (or the same drawing) at the same place on many pages. A table's REPEATED header row is
    kept — it also sits elsewhere on a page (where the table starts), right above table rows."""
    n = len(pages)
    pnum = re.compile(r'^(page )?\d+( (of|/) \d+)?$')
    # page numbers change from page to page (digits masked); a row of 3+ cells must match
    # exactly — table rows of the same shape at the same place are content, not a header
    sig = lambda b: re.sub(r'\d+', '#', norm(b.text)) if len(b.lines) <= 2 else norm(b.text)
    per_page = [_row_clusters(_bands_of(P.lines)) for P in pages]
    occ = collections.defaultdict(list)
    for bands in per_page:
        seen = set()
        for j, b in enumerate(bands):
            k = (round(b.y0 / 3), sig(b))
            if k not in seen:
                seen.add(k)
                # the next row under it - past the 1-2 wrapped lines of its own header cells
                # ('Start (before -> after)' wraps under a one-line 'Activity ID')
                nx = None
                for c in bands[j + 1:j + 4]:
                    if len(c.lines) >= 3 or c.y0 - b.y1 > 14:
                        nx = c
                        break
                occ[k].append((b, nx))
    need = max(3, 0.3 * n)

    def table_header(k):
        # a table's header row repeated on every page of the table: 3+ cells, and wherever
        # it sits, table rows with the same columns follow right under it
        if not all(len(b.lines) >= 3 and nx is not None and nx.grid and nx.y0 - b.y1 <= 30
                   and _col_match(b.lines, nx.lines) >= 3 for b, nx in occ[k]):
            return False
        # ... but when the row under it is the SAME on every page too, the two are a running
        # header TABLE (the narrative letterhead 'OWNER CONSULTANT CONTRACTOR' over 'logo logo
        # logo' - matched as one column grid since centred columns match by their centre),
        # not a table's header over its body rows
        under = {(round(nx.y0 / 3), sig(nx)) for _, nx in occ[k]}
        return not (len(occ[k]) >= 3 and len(under) == 1)

    H = pages[0].H if pages else 842.0

    def running(b):
        k = (round(b.y0 / 3), sig(b))
        if b.y1 < 0.1 * H or b.y0 > 0.9 * H:      # the header / footer band of the sheet:
            # on every page of a short report — but a table's header row repeated at the top
            # of each page (a report without a running header: the weekday row of the month
            # calendars that open every page) is content
            return n >= 2 and len(occ[k]) >= min(n, need) and not table_header(k)
        return n >= 3 and len(occ[k]) >= need and not table_header(k)
    dkey = lambda d: (round(d.ry0 / 3), round(d.x0 / 3), round(d.w / 3), round((d.ry1 - d.ry0) / 3))
    dcnt = collections.Counter()
    for P in pages:
        dcnt.update({dkey(d) for d in P.draws if not d.img})
    drun = {k for k, c in dcnt.items() if n >= 3 and c >= need}
    # the text of every running header / footer row (in the sheet's header / footer band) —
    # a report shell's repeated table footer is drawn right under the content on the LAST
    # page of its sheet (under the contents list, at the end of the report), away from its
    # usual place: as the last row of a page it is furniture too
    run_sigs = {k[1] for k, lst in occ.items()
                if k[1] and (lst[0][0].y1 < 0.1 * H or lst[0][0].y0 > 0.9 * H) and running(lst[0][0])}
    for P, bands in zip(pages, per_page):
        edge = 0.07 * P.H
        folio = lambda b: (len(b.lines) == 1 and pnum.match(b.lines[0].n or '#')
                           and (b.y1 < edge or b.y0 > P.H - edge))
        content = [b for b in bands if not running(b) and not folio(b)]
        last = max(content, key=lambda b: b.y1) if content else None
        keep = []
        for b in bands:
            if running(b) or (b is last and len(b.lines) <= 2 and sig(b) in run_sigs):
                continue
            keep.extend(l for l in b.lines
                        if not (pnum.match(l.n or '#') and (l.y1 < edge or l.y0 > P.H - edge)))
        P.lines = keep
        P.draws = [d for d in P.draws if d.img or dkey(d) not in drun]


def _make_bands(P):
    P.bands = _bands_of(P.lines)


def _body_size(pages):
    c = collections.Counter()
    for P in pages:
        for l in P.lines:
            c[round(l.size * 2) / 2] += len(l.t)
    return c.most_common(1)[0][0] if c else 10.0


def _mark_headings(pages, body, hints):
    for P in pages:
        for b in P.bands:
            txt = b.text
            n = norm(txt)
            if len(n) < 3 or len(txt) > HEAD_MAX_CHARS or not re.search(r'[a-z]', n):
                continue
            alnum = re.sub(r'[^0-9a-z]', '', n)
            if sum(ch.isdigit() for ch in alnum) > 0.4 * len(alnum):
                continue                      # a value ("247 WD", "4.2%"), not a title
            if len(b.lines) > 2 or (len(b.lines) == 2 and b.lines[1].x0 - b.lines[0].x1 > 30):
                continue                      # cells side by side, not a title
            styled = all((l.bold and l.size >= 0.8 * body) or l.size >= 1.25 * body for l in b.lines)
            hinted = bool(hints) and any(
                n == h or (len(n) >= 8 and (h.startswith(n) or n.startswith(h[:max(12, len(h) // 2)])))
                for h in hints)
            b.heading = ((styled and not _card_value(P, b, body)) or hinted) and not _list_band(P, b)


def _list_band(P, b):
    """A label inside one of a STACK of rounded boxes - a WBS tree band ('Approval', 'Remaining
    Towers', a wrapped 'Traffic Study Approval From / Authorities' in the Baseline Revision WBS
    comparison), a list of chips - is a list item, not a heading: its rounded box (1-2 lines
    tall) has another rounded box with the same right edge right above or below it (<= 8 pt
    apart). A group heading styled as a box stands apart from other boxes; table cells are
    square (STUDIO-RICH-8)."""
    lh = max(b.y1 - b.y0, 1.0)
    for d in P.draws:
        if (d.thin or d.img or not d.curve or d.h > max(lh + 26, 3.2 * lh)
                or not (d.x0 - 1 <= b.x0 and b.x1 <= d.x1 + 1 and d.y0 - 1 <= b.y0 and b.y1 <= d.y1 + 1)):
            continue
        for o in P.draws:
            if (o is not d and o.curve and not o.thin and not o.img and abs(o.x1 - d.x1) <= 2
                    and 0.4 * d.h <= o.h <= 2.5 * d.h
                    and (0 <= d.y0 - o.y1 <= 8 or 0 <= o.y0 - d.y1 <= 8)):
                return True
    return False


def _boxed_rows(P, reg):
    """Most of a region's rows are list bands side by side (two WBS trees in grid columns read
    as a 'table'): its few-row / small-table rules do not apply - each column is a list that
    continues on the next page, its branches kept with their first child by the renderer."""
    def boxed(row):
        return all(_list_band(P, _Band([l])) for l in row.lines)
    rows = reg.rows
    return bool(rows) and sum(1 for r in rows if boxed(r)) >= 0.8 * len(rows)


def _card_value(P, b, body):
    """A KPI tile's value ("0 days" under the label "DELAY"): a big / numeric line inside a card
    box that also holds a smaller label above it — card content, not a heading (STUDIO-PDF-2)."""
    if not re.search(r'[0-9]', b.text) and b.size < 1.8 * body:
        return False

    def inside(o, d):
        return d.x0 - 1 <= o.x0 and o.x1 <= d.x1 + 1 and d.y0 - 1 <= o.y0 and o.y1 <= d.y1 + 1
    for d in P.draws:
        if d.thin or d.img or d.h > 170 or not inside(b, d):
            continue
        if any(o is not b and o.y1 <= b.y0 + 1 and o.size < 0.8 * b.size and inside(o, d) for o in P.bands):
            return True
    return False


def _col_match(a_lines, b_lines):
    """How many of ``b_lines`` sit in one of ``a_lines``' columns: the same left edge, right
    edge or CENTRE (a centred column - 'Rev.00' over '24h/day' over 'Non-working' - shares
    neither edge from row to row; STUDIO-RICH-6)."""
    hits = 0
    for l in b_lines:
        c = (l.x0 + l.x1) / 2
        if any(abs(l.x0 - m.x0) < 4 or abs(l.x1 - m.x1) < 4 or abs(c - (m.x0 + m.x1) / 2) < 3
               for m in a_lines):
            hits += 1
    return hits


class _Region:
    """A table (or a grid of cells): consecutive rows of side-by-side cells."""

    def __init__(self, band):
        self.rows = [band]
        self.members = [band]
        self.y0, self.y1 = band.y0, band.y1
        self.col0 = band.x0

    def add(self, band, row):
        self.members.append(band)
        if row:
            self.rows.append(band)
        self.y1 = max(self.y1, band.y1)
        self.col0 = min(self.col0, band.x0)

    def pop_row(self):
        band = self.rows.pop()
        self.members = [m for m in self.members if m is not band]
        self.y1 = max(m.y1 for m in self.members)
        return band

    @property
    def lines(self):
        return [l for b in self.rows for l in b.lines]

    def header_rows(self, strict=False):
        """Leading bold rows = the header. ``strict``: only when no later row is bold (bold
        group / total / label rows make a bold first row ambiguous)."""
        k = 0
        while k < len(self.rows) and k < 3 and self.rows[k].bold:
            k += 1
        if len(self.rows) > 1 and (k == len(self.rows) or all(r.bold for r in self.rows)):
            return 0                     # every row bold (a list of bold bars): no header row
        if strict and (any(r.bold for r in self.rows[k:])
                       or not all(re.search('[a-z]', _sig(r)) for r in self.rows[:k])):
            return 0
        return k


def _sig(band):
    """A row's structure (digits ignored) — a repeated header row has the same signature."""
    return re.sub(r'[0-9]+', '#', norm(band.text))


def _continues_row(reg, b):
    """A second line of the previous row (wrapped cells) rather than a new row."""
    prev = reg.rows[-1]
    return (b.y0 - prev.y0 < 1.5 * max(b.size, 1) and len(b.lines) < len(prev.lines)
            and not any(abs(l.x0 - reg.col0) < 4 for l in b.lines))


def _cell_line(reg, b):
    """A heading-like line right under a table row that is really one of that table's CELL
    lines: it starts on one of the table's column edges RIGHT of its first column and stays
    inside the table. A heading sits at the left edge of the content, never inside a column.
    E.g. the first line of a tall row's vertically centred cell ('Predecessor link …' - the
    same text as a 'Fix these first' sub-line, so it was hinted as a heading) or a wrapped
    WBS cell ending '… > Machine Tower': read as headings they broke every tall-row register
    into 1-2 row 'tables' (STUDIO-RICH: 27 false table_split_few_rows on one GBT register)."""
    if b.x0 <= reg.col0 + 8:
        return False
    cols = {round(l.x0) for r in reg.rows for l in r.lines if l.x0 > reg.col0 + 8}
    if not any(abs(b.x0 - c) < 4 for c in cols):
        return False
    nxt = [c for c in cols if c > b.x0 + 4]          # ... and stays inside its column
    return not nxt or b.x1 <= min(nxt) + 2


def _row_clusters(bands):
    """Bands that overlap vertically are ONE table row (cells aligned top / middle / bottom,
    wrapped cells, a two-line header cell)."""
    out = []
    for b in bands:
        if out and not b.heading and not out[-1].heading and b.y0 < out[-1].y1 - 1.0:
            out[-1] = _Band(out[-1].lines + b.lines)
        else:
            out.append(b)
    return out


def _make_regions(P):
    regs, cur, pend = [], None, []

    def close_trailing():
        for m in pend:
            if m.y0 - cur.y1 <= 14 and cur.col0 - 4 <= m.x0 and not m.heading:
                cur.add(m, False)       # the last row's wrapped cell lines
            else:
                break
    for b in _row_clusters(P.bands):
        if b.grid and not b.heading:
            last_y = max(cur.y1, pend[-1].y1) if (cur is not None and pend) else (cur.y1 if cur else 0)
            if cur is not None and b.y0 - last_y <= 24 and _col_match(cur.lines, b.lines) >= 2:
                # A header repeats one of the table's HEADER rows (its leading bold rows, over
                # body rows that are not bold) and does not open with a number / date: all-bold
                # rows sharing a signature - dated rows '07 Jan 2025 Non-working' / '25 Jan
                # 2025 Non-working', bars 'Erection of Each Level of TR34 ...' / '... TR35 ...'
                # - are body rows, not a header repeated (STUDIO-RICH-6)
                hdr = cur.header_rows() if b.bold and len(cur.rows) > 1 else 0
                if hdr and re.search('[a-z]', _sig(b)) and not re.match(r'\s*\d', b.text) \
                        and any(_sig(r) == _sig(b) for r in cur.rows[:hdr]):
                    # the same header row again: a new table starts here, with the bold title
                    # row(s) sitting right on top of it
                    lead = []
                    while len(cur.rows) > 1 and cur.rows[-1].bold and b.y0 - cur.rows[-1].y1 <= 8 and not pend:
                        lead.insert(0, cur.pop_row())
                    new = _Region(lead[0] if lead else b)
                    for r in lead[1:] + ([b] if lead else []):
                        new.add(r, True)
                    cur = new
                    regs.append(cur)
                else:
                    for m in pend:
                        cur.add(m, False)
                    cur.add(b, not _continues_row(cur, b))
            else:
                if cur is not None:
                    close_trailing()
                cur = _Region(b)
                regs.append(cur)
            pend = []
        elif (cur is not None and b.heading and _cell_line(cur, b)
              and b.y0 - (pend[-1].y1 if pend else cur.y1) <= 14):
            b.heading = False            # a cell's line that reads like a heading (see _cell_line)
            pend.append(b)
        elif (cur is not None and len(pend) < 12
              and b.y0 - (pend[-1].y1 if pend else cur.y1) <= (6 if b.heading else 14)):
            pend.append(b)               # a wrapped cell line / a group row inside the table
        else:
            if cur is not None:
                close_trailing()
            cur, pend = None, []
    if cur is not None:
        close_trailing()
    P.regions = regs


def _tiles(P, y_top, y_lim):
    """Boxes of a KPI / summary card row starting in [y_top, y_lim]: two or more boxes of about
    the same height side by side with a gap between them, each holding text."""
    boxes = [(d.x0, d.y0, d.x1, d.y1) for d in P.draws
             if not d.thin and not d.img and 18 <= d.h <= 170 and 28 <= d.w <= 0.62 * P.W
             and y_top - 2 <= d.y0 <= y_lim]
    # boxes drawn as four thin borders: pair the vertical sides sharing a span
    vs = collections.defaultdict(list)
    for d in P.draws:
        if d.w < 2.5 and 18 <= d.h <= 170 and y_top - 2 <= d.y0 <= y_lim:
            vs[(round(d.y0), round(d.y1))].append((d.x0 + d.x1) / 2)
    for (a, b), xs in vs.items():
        xs = sorted(set(round(x, 1) for x in xs))
        for x0, x1 in zip(xs, xs[1:]):
            if 28 <= x1 - x0 <= 0.62 * P.W:
                boxes.append((x0, a, x1, b))
    boxes = [bx for bx in boxes if any(bx[0] - 1 <= l.x0 and l.x1 <= bx[2] + 1 and bx[1] - 1 <= l.y0
                                        and l.y1 <= bx[3] + 1 for l in P.lines)]
    rows = []
    for bx in sorted(boxes, key=lambda b: (b[1], b[0])):
        for row in rows:
            a = row[0]
            if abs(bx[1] - a[1]) <= 3 and abs((bx[3] - bx[1]) - (a[3] - a[1])) <= 0.25 * (a[3] - a[1]):
                if not any(abs(bx[0] - c[0]) < 3 and abs(bx[2] - c[2]) < 3 for c in row):
                    row.append(bx)       # (a filled box and its border pair are the same tile)
                break
        else:
            rows.append([bx])
    for row in rows:
        row.sort(key=lambda b: b[0])
        gaps = [b[0] - a[2] for a, b in zip(row, row[1:])]
        # cards stand apart; table cells touch (a row with touching boxes is a table row)
        if len(row) >= 2 and all(g >= 2 for g in gaps):
            return row
    return None


def _section_size(pages, body, firsts):
    """The smallest heading size that marks a TOP-LEVEL section: clearly larger than the body
    text and starting a page at least half of the times it occurs (sections that start on a
    new page by design). A page-top heading at least this big is a section break."""
    seen = collections.defaultdict(lambda: [0, 0])
    for P in pages[1:]:
        for b in P.bands:
            if b.heading:
                seen[round(b.size * 2) / 2][0] += 1
    for P, (kind, _, size) in zip(pages[1:], firsts[1:]):
        if kind == 'heading':
            seen[round(size * 2) / 2][1] += 1
    cands = [k for k, (n, t) in seen.items() if k >= 1.2 * body and t >= max(1, 0.5 * n)]
    # The report's LARGEST titles are its top-level sections even when they flow on under the
    # previous one (the Reporting Studio's numbered items: ~40 % start a page). A page that
    # ends early because the next item's opening - its title, intro and an unsplittable chart
    # row - does not fit under it is a section break, not a pushed block (STUDIO-RICH-5).
    common = [k for k, (n, t) in seen.items() if n >= 3]
    if common:
        top = max(common)
        n, t = seen[top]
        if top >= 1.4 * body and t >= 0.25 * n:
            cands.append(top)
    return min(cands) if cands else float('inf')


def _first_block(P):
    """What page P starts with: ('heading'|'kpi'|'table'|'figure'|'text'|'empty', text, size)."""
    if P.empty:
        return 'empty', '', 0
    top = P.top
    fb = P.bands[0] if P.bands else None
    figs_above = [d for d in P.draws if not d.thin and (fb is None or d.y1 <= fb.y0 + 0.5)
                  and not (fb is not None and d.x0 <= fb.x0 and d.x1 >= fb.x1 and d.y1 >= fb.y1)]
    if fb is not None and fb.heading and fb.y0 <= top + TOP_ZONE_PT and not figs_above:
        return 'heading', fb.text, fb.size
    if _tiles(P, top, top + TOP_ZONE_PT):
        return 'kpi', '', 0
    if fb is not None and fb.grid and fb.y0 <= top + TOP_ZONE_PT and not figs_above:
        return 'table', fb.text, fb.size
    zone = [d for d in P.draws if d.y0 <= top + TOP_ZONE_PT and not d.thin]
    textfree = [d for d in zone if not any(d.x0 - 1 <= l.x0 and l.x1 <= d.x1 + 1 and d.y0 - 1 <= l.y0
                                           and l.y1 <= d.y1 + 1 for l in P.lines)]
    if any(d.img or d.curve for d in zone) or len(textfree) >= 2:
        return 'figure', '', 0
    if fb is not None and fb.grid:
        return 'table', fb.text, fb.size
    return 'text', (fb.text if fb is not None else ''), (fb.size if fb is not None else 0)


def _below(P, y, head=None):
    """Content items on P starting below y (ignoring a heading's own underline / bar)."""
    out = []
    for b in P.bands:
        if b.y0 >= y - 0.5:
            out.append((b.y0, b.y1, 'text', b))
    for d in P.draws:
        if d.y0 < y - 0.5:
            continue
        if head is not None and d.thin and d.y0 <= y + 8 and d.h < 3:
            continue                      # the heading's underline
        out.append((d.y0, d.y1, 'figure' if (d.img or d.curve or not d.thin) else 'rule', d))
    for y0, y1 in P.misc:
        if y0 >= y - 0.5:
            out.append((y0, y1, 'figure', None))
    return out


def _table_split(A, B, a_reg, b_reg, area, flags):
    """Compare the table fragment ending page A with the one starting page B."""
    ha = a_reg.header_rows()
    body_a = len(a_reg.rows) - ha
    hb = 0
    sigs = [_sig(r) for r in a_reg.rows[:ha]]
    # the header's words on A - its rows plus the wrapped lines of its cells sitting right above
    # it ('Start' / '(before' over 'Activity ID ... Shift'): on B the same header may be laid
    # out as ONE row of 8 cells, so a header repeated is also one whose words are all A's
    words = lambda t: set(re.findall(r'[a-z]+', norm(t)))
    hw = set()
    for r in a_reg.rows[:ha]:
        hw |= words(r.text)
    for b in A.bands:
        if b.bold and b.y1 <= a_reg.y0 + 1 and a_reg.y0 - b.y1 <= 20 and b.x0 >= a_reg.col0 - 4:
            hw |= words(b.text)
    while (hb < len(b_reg.rows) and hb < 3 and b_reg.rows[hb].bold
           and (_sig(b_reg.rows[hb]) in sigs or (ha and words(b_reg.rows[hb].text) <= hw))):
        hb += 1
    if hb == len(b_reg.rows) and hb > 1:
        hb = 0
    body_b = len(b_reg.rows) - hb
    height = (a_reg.y1 - a_reg.y0) + (b_reg.y1 - b_reg.y0)
    what = f'table {a_reg.rows[0].text[:40]!r}'
    if min(body_a, body_b) < MIN_ROWS:
        flags.append(_flag('table_split_few_rows', A.no,
                           f'{what} breaks p{A.no}->p{B.no} with {body_a} body row(s) on p{A.no} '
                           f'and {body_b} on p{B.no} (the rules keep >= {MIN_ROWS})'))
    elif height <= FIT * area:
        flags.append(_flag('small_table_split', A.no,
                           f'{what} ({body_a}+{body_b} rows, {height:.0f}pt = {height / area:.0%} of a page) '
                           f'fits on one page but is split p{A.no}->p{B.no}'))
    if a_reg.header_rows(strict=True) and not hb and body_b > 0:
        flags.append(_flag('table_header_not_repeated', B.no,
                           f'{what} continues on p{B.no} without its header row '
                           f'{a_reg.rows[0].text[:40]!r}'))


def _opens_new_table(a_reg, B):
    """Page B opens a NEW table / grid rather than continuing A's: its own bold title row(s)
    (side-by-side month titles "Feb 2026  Mar 2026" — not a heading, so they start the region)
    sit on top of a row repeating A's column header ("Mon Tue ... Sun"). The next row of whole
    month calendars is a new grid, not a table continuing with 1 row (CAL-PDF-1 / STUDIO-PDF-3);
    a real continuation starts with the repeated header itself, or with body rows."""
    sigs = {_sig(r) for r in a_reg.rows[:3] if r.bold and re.search('[a-z]', _sig(r))}
    head = _row_clusters(B.bands)[:3]
    if not sigs or len(head) < 2 or not head[0].bold or _sig(head[0]) in sigs:
        return False
    for r in head[1:]:
        if not r.bold:
            return False
        if _sig(r) in sigs:
            return True
    return False


def _figure_cut(A, B, area_top, area_bottom):
    """A chart / diagram cut at the break between page A and page B (Chrome clips each part to
    its page: the part on A ends exactly on A's clip line, the part on B starts on B's)."""
    fa = [d for d in A.draws if not d.thin]
    fb = [d for d in B.draws if not d.thin]
    if not fa or not fb:
        return None
    clip_a = max(d.y1 for d in A.draws)
    clip_b = min(d.y0 for d in B.draws)
    if clip_a < area_bottom - 0.6 or clip_b > area_top + 0.6:
        return None                       # ends above the page-area bottom: nothing was cut
    if A.bands and max(b.y1 for b in A.bands) > clip_a + 0.5:
        return None                       # text below the graphic: it ended on this page
    if B.bands and min(b.y0 for b in B.bands) < clip_b - 0.5:
        return None                       # text above the graphic: it starts fresh
    wide = 0.85 * max(A.W, 1)

    def panel(d, P):
        # a long card / panel the content flows through (a bars list's card, 40 pages of
        # bars): taller on this page than anything kept whole and holding lines of text - its
        # background is sliced at the break like a long table's, it is not a chart (STUDIO-RICH-7)
        return (d.y1 - d.y0 > FIT * (area_bottom - area_top)
                and sum(1 for b in P.bands if d.x0 - 1 <= b.x0 and b.x1 <= d.x1 + 1
                        and d.y0 - 1 <= b.y0 and b.y1 <= d.y1 + 1) >= 3)
    ea = [d for d in fa if abs(d.y1 - clip_a) < 0.35 and d.w < wide and not panel(d, A)]
    eb = [d for d in fb if abs(d.y0 - clip_b) < 0.35 and d.w < wide and not panel(d, B)]
    if not ea or not eb:
        return None

    def overlap(a, b):
        return min(a.x1, b.x1) - max(a.x0, b.x0) > 0.5 * min(a.w, b.w)
    if not any(overlap(a, b) for a in ea for b in eb):
        return None
    # a table row / text box that breaks has text right at the break — a graphic does not.
    # Measured per LINE, not per band: a column chart (div bars) whose columns were
    # fragmented one by one has the other columns' value labels at the top of B in the same
    # row as the continuing bars, but not over them (NARR-PDF-5: GBT narrative p34->p35)
    def over(b, ds):
        return any(l.x0 < d.x1 and l.x1 > d.x0 for l in b.lines for d in ds)
    near_a = [b for b in A.bands if b.y1 > clip_a - 4 and over(b, ea)]
    near_b = [b for b in B.bands if b.y0 < clip_b + 4 and over(b, eb)]
    if near_a or near_b:
        return None

    # … and a WHOLE table row that happens to end on the break (its cell shading / borders
    # end there, its text sits inside the cell a few points above) followed by the table's
    # repeated header row (shaded cells, text inside) is a table continuing, not a graphic
    def holds_text(d, bands, bottom):
        return any(b.x0 < d.x1 and b.x1 > d.x0 and b.y0 >= d.y0 - 0.5 and b.y1 <= d.y1 + 0.5
                   and ((d.y1 - b.y1) if bottom else (b.y0 - d.y0)) <= 10 for b in bands)
    if all(holds_text(d, A.bands, True) for d in ea) and all(holds_text(d, B.bands, False) for d in eb):
        return None
    return len(ea), len(eb)


def _in_closed_box(P, h, area_top, area_bottom):
    """Heading-like band ``h`` is the title of a card drawn around it that also holds content
    under it and CLOSES on page P (above the page-area bottom): the card's own label (the Float
    Health card's bold 'High Float > 44 WD - construction' driver over its bar and note), not a
    heading whose block can be parted from it by the break (STUDIO-RICH-10). A page-spanning
    shell / frame or a title bar holding only the title does not count."""
    for d in P.draws:
        if d.thin or d.img or d.w >= 0.9 * P.W:
            continue
        if not (d.x0 - 1 <= h.x0 and h.x1 <= d.x1 + 1 and d.y0 - 1 <= h.y0 and h.y1 <= d.y1 + 1):
            continue
        if d.y1 >= area_bottom - 2 or (d.y0 <= area_top + 2 and d.y1 - d.y0 >= 0.9 * (area_bottom - area_top)):
            continue
        if any(b is not h and b.y0 >= h.y1 - 1 and d.x0 - 1 <= b.x0 and b.x1 <= d.x1 + 1
               and b.y1 <= d.y1 + 1 for b in P.bands):
            return True
    return False


def _block_split(A, B, area_top, area_bottom, area):
    """A SMALL card / box split by the break between page A and page B (STUDIO-RICH-10).

    Chrome paints each fragment of a split box as its own box: the part on A starts on A and
    ends exactly on A's clip line, the part on B (same left / right edges) starts on B's clip
    line and ends on B. A box that small (<= FIT of a page in all) is one the rules keep whole -
    a driving-path milestone card with its title on A and its dates on B. ``_figure_cut``
    leaves such a TEXT box alone (text at the break is how a long panel or a table row reads);
    a long panel that flows across pages is taller than FIT and is not reported here.
    A whole table row ending right on the break (its text sits a few points above its bottom
    edge) followed by the next row (its text inside the cell padding) is not a split box."""
    fa = [d for d in A.draws if not d.thin and not d.img]
    fb = [d for d in B.draws if not d.thin and not d.img]
    if not fa or not fb:
        return None
    clip_a = max(d.y1 for d in A.draws)
    clip_b = min(d.y0 for d in B.draws)
    # the area edges are measured on text boxes (a line's ascent reaches ~1 pt above the box
    # its fragment is painted in)
    if clip_a < area_bottom - 2 or clip_b > area_top + 2:
        return None
    wide = 0.85 * max(A.W, 1)
    ea = [d for d in fa if abs(d.y1 - clip_a) < 0.35 and d.w < wide and d.y0 > area_top + 2]
    eb = [d for d in fb if abs(d.y0 - clip_b) < 0.35 and d.w < wide and d.y1 < area_bottom - 2]

    def edge_band(d, bands, bottom):
        inside = [b for b in bands if b.x0 < d.x1 and b.x1 > d.x0
                  and b.y0 >= d.y0 - 1.5 and b.y1 <= d.y1 + 1.5]
        if not inside:
            return None
        return max(inside, key=lambda b: b.y1) if bottom else min(inside, key=lambda b: b.y0)

    def text_in(d, bands, bottom):
        e = edge_band(d, bands, bottom)
        if e is None:
            return None
        return (d.y1 - e.y1) if bottom else (e.y0 - d.y0)
    for a in ea:
        for b in eb:
            if abs(a.x0 - b.x0) > 1.5 or abs(a.x1 - b.x1) > 1.5:
                continue
            h = (a.y1 - a.y0) + (b.y1 - b.y0)
            if h > FIT * area:
                continue
            ga, gb = text_in(a, A.bands, True), text_in(b, B.bands, False)
            if ga is None or gb is None:
                continue                  # a part holding no text: a table's / chart's frame
                                          # (graphic_cut reads charts)
            # two whole rows meeting at the break: the next row's text sits inside its cell
            # padding - a box's continuation starts with its text right on its top edge (the
            # padding is not repeated on a sliced fragment)
            if ga <= 10 and 1.5 < gb <= 10:
                continue
            # ... or a table's last row on A (a shaded row of 3+ cells) and its header repeated
            # on B: the same columns
            la, lb = edge_band(a, A.bands, True), edge_band(b, B.bands, False)
            if (la is not None and lb is not None and len(la.lines) >= 3 and len(lb.lines) >= 3
                    and _col_match(la.lines, lb.lines) >= 3):
                continue
            return a, b, h
    return None


def _analyze_pages(pages, hints=()):
    n = len(pages)
    _strip_running(pages)
    for P in pages:
        _make_bands(P)
    body = _body_size(pages)
    _mark_headings(pages, body, [norm(h) for h in hints if norm(h)])
    for P in pages:
        _make_regions(P)
    full = [P for P in pages if not P.empty]
    if not full:
        return [], [], (0, 0)
    tops = sorted(P.top for P in full)
    bots = sorted(min(P.bottom, P.H) for P in full)
    T = tops[len(tops) // 2]
    B = bots[min(len(bots) - 1, round((len(bots) - 1) * 0.9))]
    area = max(B - T, 1.0)
    # the page-area edges: nothing on a normal page reaches past them (spilled pages aside)
    area_bottom = max((P.bottom for P in full if P.bottom <= P.H - 8), default=B)
    area_top = min((P.top for P in full if P.top >= 8), default=T)
    firsts = [_first_block(P) for P in pages]
    section_size = _section_size(pages, body, firsts)
    flags, info = [], []

    for i, P in enumerate(pages):
        nxt = pages[i + 1] if i + 1 < n else None
        kind_n, text_n, size_n = firsts[i + 1] if nxt is not None else ('none', '', 0)
        if P.empty:
            if n > 1:
                flags.append(_flag('empty_page', P.no, 'the page has nothing on it besides the running '
                                                       'header / footer / frame'))
            continue
        # content running off the sheet
        cut_t = [l for l in P.lines if l.y1 > P.H - 0.5 or l.y0 < 0.5]
        if cut_t:
            flags.append(_flag('text_cut', P.no, 'text runs off the sheet: '
                               + '; '.join(repr(l.t[:50]) for l in cut_t[:3])))
        elif P.no > 1 and T - 24 > 8:
            high = [l for l in P.lines if l.y0 < T - 24]
            if high:
                flags.append(_flag('content_in_margin', P.no,
                                   f'content starts at {min(l.y0 for l in high):.0f}pt, inside the top '
                                   f'margin (the body starts at {T:.0f}pt on the other pages) - the page '
                                   f'lost its margin / running header: ' + repr(high[0].t[:50])))
        edge = [d for d in P.draws if (d.ry1 > P.H + 0.5 and d.y0 < P.H - 2) or (d.ry0 < -0.5 and d.y1 > 2)]
        edge = [d for d in edge if d.img or (not d.thin and d.w < 0.85 * P.W and (d.ry1 - d.ry0) < 0.5 * P.H)]
        if edge:
            flags.append(_flag('graphic_cut', P.no, f'{len(edge)} picture / drawing part(s) are clipped by '
                                                   f'the edge of the sheet'))
        if nxt is None or nxt.empty:
            continue
        # headings that end the page (a title inside a card that closes on this page keeps its
        # content inside that card - it cannot be parted from it by the break)
        hs = [b for b in P.bands if b.heading and not _in_closed_box(P, b, area_top, area_bottom)]
        if hs:
            h = max(hs, key=lambda b: b.y0)
            below = [it for it in _below(P, h.y1, head=h) if it[2] != 'rule']
            ext = (max(y1 for _, y1, _, _ in below) - h.y1) if below else 0.0
            closing = kind_n == 'heading' and size_n >= h.size - 0.3
            if (not below or ext < 3) and not closing:
                kind = 'kpi_separated_from_heading' if kind_n == 'kpi' else 'orphaned_heading'
                flags.append(_flag(kind, P.no, f'heading {h.text[:60]!r} is the last thing on page {P.no}; '
                                               f'its {kind_n} starts page {nxt.no}'))
            elif (ext <= INTRO_MAX_PT and all(k == 'text' and not it.grid and not it.heading
                                              for _, _, k, it in below)
                  and kind_n in ('table', 'figure', 'kpi')):
                kind = 'kpi_separated_from_heading' if kind_n == 'kpi' else 'heading_separated_from_block'
                flags.append(_flag(kind, P.no, f'heading {h.text[:60]!r} + {ext:.0f}pt of intro end page '
                                               f'{P.no}; its {kind_n} starts page {nxt.no}'))
        # a table that continues across the break
        if P.regions and nxt.regions and kind_n == 'table':
            a_reg, b_reg = P.regions[-1], nxt.regions[0]
            tail = [it for it in _below(P, a_reg.y1 + 1) if it[2] != 'rule']
            if (not tail and abs(b_reg.y0 - nxt.bands[0].y0) < 0.5
                    and _col_match(a_reg.lines, b_reg.rows[0].lines) >= 2
                    and not _opens_new_table(a_reg, nxt)
                    and not (_boxed_rows(P, a_reg) and _boxed_rows(nxt, b_reg))):
                _table_split(P, nxt, a_reg, b_reg, area, flags)
        # a chart / diagram cut by the break
        cut = _figure_cut(P, nxt, area_top, area_bottom)
        if cut:
            flags.append(_flag('graphic_cut', P.no, f'a chart / diagram is cut by the page break '
                                                   f'p{P.no}->p{nxt.no} ({cut[0]} part(s) end on the '
                                                   f'break, {cut[1]} continue at the top of p{nxt.no})'))
        else:
            split = _block_split(P, nxt, area_top, area_bottom, area)
            if split:
                a, b, h = split
                words = lambda Q, d: ' '.join(x.text for x in Q.bands if x.x0 < d.x1 and x.x1 > d.x0
                                              and x.y0 >= d.y0 - 1.5 and x.y1 <= d.y1 + 1.5)[:40]
                flags.append(_flag('block_split', P.no,
                                   f'a card / box ({h:.0f}pt = {h / area:.0%} of a page) is split '
                                   f'p{P.no}->p{nxt.no}: {words(P, a)!r} on p{P.no}, '
                                   f'{words(nxt, b)!r} on p{nxt.no}'))
        # a large blank area before the next page
        fill = (min(P.bottom, P.H) - T) / area
        if fill < 1 - BLANK:
            what = f'page {P.no} ends at {max(fill, 0):.0%} of the page ({1 - max(fill, 0):.0%} blank); ' \
                   f'page {nxt.no} starts with {kind_n}' + (f' {text_n[:40]!r}' if text_n else '')
            if kind_n == 'heading' and size_n >= section_size - 0.3:
                info.append(_flag('section_break_blank', P.no, what + ' (a new section)'))
            elif P.no == 1 and P.top - T > COVER_DROP * area:
                info.append(_flag('section_break_blank', P.no, what + ' (the cover page)'))
            elif kind_n == 'heading' and _contents_end(P, pages[max(0, i - 2):i + 1]):
                info.append(_flag('section_break_blank', P.no, what + ' (the end of the contents)'))
            else:
                flags.append(_flag('large_blank_then_continuation', P.no, what))
        # a small tail of cards / a figure alone on a (middle) page
        if i > 0:
            k0 = firsts[i][0]
            if k0 in ('figure', 'kpi') and (min(P.bottom, P.H) - P.top) < FRAGMENT * area:
                flags.append(_flag('stranded_fragment', P.no,
                                   f'only {min(P.bottom, P.H) - P.top:.0f}pt of a {k0} block '
                                   f'({(min(P.bottom, P.H) - P.top) / area:.0%} of the page) sits on page {P.no}'))
    return flags, info, (T, B)


_CONTENTS_TITLE = re.compile(r'^(table of )?contents$')


def _contents_end(P, recent):
    """The contents list ends page P: a "(Table of) contents" title on P or on one of the
    pages just before it, and P's last rows are contents entries — a title with the page
    number as its last cell, flush with the right edge of the page's text, the numbers never
    going down. The report body then starts on a fresh page by design (STUDIO-PDF-4: the
    Reporting Studio's contents page)."""
    if not any(_CONTENTS_TITLE.match(b.lines[0].n) for Q in recent for b in Q.bands
               if len(b.lines) == 1):
        return False
    rows = sorted(P.bands, key=lambda b: b.y0)[-CONTENTS_ROWS:]
    if len(rows) < CONTENTS_ROWS:
        return False
    right = max(b.x1 for b in P.bands)
    nums = []
    for b in rows:
        last = b.lines[-1]
        if not (b.grid and re.fullmatch(r'\d{1,4}', last.t.strip()) and right - last.x1 <= 12
                and re.search(r'[a-z]', b.text.lower())):
            return False
        nums.append(int(last.t))
    return nums == sorted(nums)


def check_pdf(path, headings=None):
    """Check a PDF. ``headings``: optional heading texts (hints) — e.g. from the source HTML."""
    try:
        pages = _read_pdf(path)
    except ImportError:
        return _result(path, 'pdf', 'pymupdf', 0, [], status='skipped',
                       message='PyMuPDF (pymupdf) is not installed - the PDF cannot be read.')
    flags, info, (T, B) = _analyze_pages(pages, headings or ())
    res = _result(path, 'pdf', 'pymupdf', len(pages), flags, info)
    res['body_area_pt'] = [round(T, 1), round(B, 1)]
    return res


# ── heading hints from a source HTML ─────────────────────────────────────────────
def html_headings(html):
    """Heading texts of a report HTML: h1-h6 + the renderers' heading selectors.

    A ``tag.class`` selector matches that TAG only, exactly as the print composer applies it:
    ``p.rescap`` (a lead-in above a chart group) is a heading, but ``div.rescap`` (the 'Peak …'
    caption closing a chart) is not - reading the bare class made every chart caption a
    'heading' and flagged the table after it as orphaned. ``[attr]`` / ``[attr="v"]`` too."""
    from html.parser import HTMLParser
    sels = []                                   # (tag or None, class or None, attr or None, value)
    for s in (getattr(_rt, 'HEADING_SELECTORS', ()) if _rt else ()):
        m = re.fullmatch(r'([a-z0-9]+)?\.([\w-]+)', s)
        if m:
            sels.append((m.group(1), m.group(2), None, None))
            continue
        m = re.fullmatch(r'\[([\w-]+)(?:="([^"]*)")?\]', s)
        if m:
            sels.append((None, None, m.group(1), m.group(2)))

    def _is_heading(tag, attrs):
        if tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            return True
        a = dict(attrs)
        cls = set((a.get('class') or '').split())
        for t, c, at, v in sels:
            if t and t != tag:
                continue
            if c and c in cls:
                return True
            if at and at in a and (v is None or (a.get(at) or '') == v):
                return True
        return False

    class P(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.stack, self.out = [], []

        def handle_starttag(self, tag, attrs):
            if tag in ('br', 'img', 'hr', 'meta', 'link', 'input', 'col', 'wbr', 'source'):
                return
            self.stack.append([tag, _is_heading(tag, attrs), []])

        def handle_endtag(self, tag):
            while self.stack:
                t, is_h, buf = self.stack.pop()
                if is_h:
                    txt = ' '.join(''.join(buf).split())
                    if 2 < len(txt) <= HEAD_MAX_CHARS:
                        self.out.append(txt)
                if self.stack:
                    self.stack[-1][2].extend(buf)
                if t == tag:
                    break

        def handle_data(self, data):
            if self.stack:
                self.stack[-1][2].append(data)
    p = P()
    p.feed(html)
    return p.out


# ══ Word ═════════════════════════════════════════════════════════════════════════
def docx_headings(path):
    """Heading texts of a .docx (heading / title styles, outline levels, short bold lines)."""
    try:
        from docx import Document
    except ImportError:
        return []
    out = []
    try:
        d = Document(path)
    except Exception:
        return []
    for p in d.paragraphs:
        t = ' '.join(p.text.split())
        if not t or len(t) > HEAD_MAX_CHARS:
            continue
        st = (p.style.name if p.style is not None else '').lower()
        runs = [r for r in p.runs if r.text.strip()]
        if st.startswith(('heading', 'title', 'subtitle')) or (runs and all(r.bold for r in runs)):
            out.append(t)
    return out


PIC_BLANK_EXTENT = 0.5    # a picture whose painted part spans < 50 % of its height …
PIC_BLANK_PT = 144.0      # … and wastes >= 2 in of white on the page is "mostly blank"
PIC_MIN_PX = 40           # icons / logos / bullets are not checked
PIC_RULE = 0.9            # a bottom row painted across >= 90 % of the width is a rule / border


def _frame_sides_only(ink, last, run=24, dark=None):
    """The ink at a picture's bottom edge only CLOSES or CONTINUES boxes - nothing is cut:

    * thin vertical lines that run down the picture: the side borders of a card / table the
      section flows through (one page-sized slice of a long Baseline Revision card - the next
      slice goes on inside the same card), and/or
    * a thin horizontal line under an EMPTY strip: the bottom border of a box that ends right
      at the slice's edge (a WBS band's outline).

    A text line or a chart cut at the edge leaves many short runs, or ink right above a wide
    run, and is still flagged (STUDIO-RICH-9: 206 of 206 GBT slices were these)."""
    import numpy as np
    W = ink.shape[1]
    cols = np.nonzero(ink[last])[0]
    if not len(cols):
        return False
    runs, start = [], cols[0]
    for a, b in zip(cols, cols[1:]):
        if b != a + 1:
            runs.append((start, a))
            start = b
    runs.append((start, cols[-1]))
    if len(runs) > 8:
        return False                      # glyph pieces: a text line is cut there
    top = max(0, last - run + 1)
    for a, b in runs:
        if b - a + 1 <= 5:                # a side border: it runs down the picture
            if ink[top:last + 1, a:b + 1].any(axis=1).mean() < 0.9:
                return False
            continue
        if b - a + 1 < 0.15 * W:
            return False                  # a short wide piece: a bar / glyph cut at the edge
        # a horizontal line: thin, drawn darker than a box's light fill, nothing dark right
        # above it (the box's padding)
        m = ink if dark is None else dark
        if m[last, a:b + 1].mean() < 0.6:
            return False
        k = last
        while k > 0 and m[k - 1, a:b + 1].mean() > 0.6:
            k -= 1
        if last - k > 12 or m[max(0, k - 4):k, a + 6:max(a + 7, b - 5)].mean() > 0.05:
            return False
    return True


def docx_picture_flags(path, pic_pages=None):
    """Pictures in a .docx body that are CUT or MOSTLY BLANK (STUDIO-WORD-2).

    A report section embedded as a picture must hold the whole section and nothing but it:

      picture_truncated     the painted content runs into the picture's bottom edge (the
                            last pixel rows cut through text / a diagram) — the rest of the
                            section never reached Word. A picture that ends on a full-width
                            rule (a table's bottom border) ends cleanly and is not flagged.
      picture_mostly_blank  the painted part spans < 50 % of the picture's height and the
                            white left over is >= 2 in on the page.

    ``pic_pages``: the page of each body picture in reading order (Word's layout), used to
    stamp the page when it lists exactly the pictures found; otherwise the page is 0 and the
    detail names the heading the picture sits under.
    Needs PyMuPDF + numpy (both ship in the bundle); returns [] without them."""
    try:
        import numpy as np
        import pymupdf
        from docx import Document
        from docx.oxml.ns import qn
    except Exception:
        return []
    try:
        d = Document(path)
    except Exception:
        return []
    rels = d.part.rels
    found, k, head = [], 0, ''
    heads, open_end = {}, set()          # picture -> its heading; pictures ending on open boxes
    for p in d.element.body.iter(qn('w:p')):
        st = p.find(qn('w:pPr') + '/' + qn('w:pStyle'))
        sv = str(st.get(qn('w:val')) or '') if st is not None else ''
        if sv.lower().startswith(('heading', 'title')):
            head = ' '.join(''.join(t.text or '' for t in p.iter(qn('w:t'))).split())[:60]
        for ext, blip in zip(p.iter(qn('wp:extent')), p.iter(qn('a:blip'))):
            k += 1
            heads[k] = head
            try:
                pix = pymupdf.Pixmap(rels[blip.get(qn('r:embed'))].target_part.blob)
                if pix.alpha or pix.n > 3:
                    pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                    if pix.alpha:
                        pix = pymupdf.Pixmap(pix, 0)
                a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            except Exception:
                continue
            H, W = a.shape[0], a.shape[1]
            if H < PIC_MIN_PX or W < PIC_MIN_PX:
                continue
            ink = (a[:, :, :3] < 244).any(axis=2) if a.shape[2] >= 3 else (a[:, :, 0] < 244)
            rows = np.nonzero(ink.any(axis=1))[0]
            if not len(rows):
                continue
            where = f'picture {k}' + (f" under {head!r}" if head else '')
            last = int(rows[-1])
            if H - 1 - last <= max(2, H * 0.001) and float(ink[last].mean()) < PIC_RULE:
                if _frame_sides_only(ink, last, dark=(a[:, :, :3] < 200).any(axis=2)
                                     if a.shape[2] >= 3 else None):
                    open_end.add(k)       # fine IF the next picture goes on with the section
                found.append((k, 'picture_truncated',
                                   f'{where}: its content runs into the bottom edge of the picture '
                                   '(cut off - the rest never reached the page)'))
            extent = (last + 1 - int(rows[0])) / H
            try:
                h_pt = int(ext.get('cy') or 0) / 12700.0
            except ValueError:
                h_pt = 0.0
            if extent < PIC_BLANK_EXTENT and h_pt * (1 - extent) >= PIC_BLANK_PT:
                found.append((k, 'picture_mostly_blank',
                                   f'{where}: only {extent:.0%} of the picture is painted - '
                                   f'~{h_pt * (1 - extent):.0f} pt of white on the page'))
    # one page-sized slice of a long section ending inside a card / table (only its side
    # borders reach the edge, or a box closes right there) whose NEXT picture continues the same
    # section is not cut - the section goes on in that picture (STUDIO-RICH-9)
    found = [f for f in found if not (f[1] == 'picture_truncated' and f[0] in open_end
                                      and f[0] + 1 in heads and heads[f[0] + 1] == heads[f[0]])]
    pages = list(pic_pages) if pic_pages and len(pic_pages) == k else None
    return [_flag(kind, (pages[n - 1] if pages else 0) or 0, detail) for n, kind, detail in found]


def word_available():
    try:
        import win32com.client  # noqa: F401
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r'Word.Application\CLSID'))
        return True
    except Exception:
        return False


def spire_available():
    try:
        import spire.doc  # noqa: F401
        return True
    except Exception:
        return False


def _word_pid(word):
    """Process id of the Word instance ``word`` (found by a unique caption on its hidden main
    window); None on any error."""
    try:
        import uuid
        import win32gui
        import win32process
        tag = 'cx-wl-' + uuid.uuid4().hex
        old = word.Caption
        word.Caption = tag
        try:
            hwnd = win32gui.FindWindow('OpusApp', tag)
        finally:
            word.Caption = old
        return int(win32process.GetWindowThreadProcessId(hwnd)[1]) if hwnd else None
    except Exception:
        return None


def word_layout(path, pid_file=None):
    """Word's own pagination of a .docx / .doc via COM: in body order, paragraphs outside
    tables (page, y, style, outline level, bold, size, keep flags, picture height), and tables
    with the page + y of every row. No PDF export (it hangs when Word runs headless).
    ``pid_file``: where to write the process id of the Word this call starts (so a caller that
    gives up waiting can stop exactly that Word)."""
    import pythoncom
    import win32com.client as wc
    path = os.path.abspath(path)
    pythoncom.CoInitialize()
    word = wc.DispatchEx('Word.Application')
    res = {'file': path}
    try:
        word.Visible = False
        word.DisplayAlerts = 0
        if pid_file:
            pid = _word_pid(word)
            if pid:
                with open(pid_file, 'w', encoding='ascii') as fh:
                    fh.write(str(pid))
        doc = word.Documents.Open(path, False, True, False)   # ConfirmConversions, ReadOnly, AddToRecent
        try:
            doc.Repaginate()
            ps = doc.PageSetup
            res['page'] = {'h': ps.PageHeight, 'w': ps.PageWidth, 'top': ps.TopMargin,
                           'bottom': ps.BottomMargin}
            items, spans = [], []
            for ti in range(1, doc.Tables.Count + 1):
                tb = doc.Tables(ti)
                r = tb.Range
                spans.append((r.Start, r.End))
                rows = []
                try:
                    objs = []
                    for row in tb.Rows:
                        rr = row.Range
                        at = doc.Range(rr.Start, rr.Start)
                        objs.append(row)
                        rows.append({'page': at.Information(3), 'y': at.Information(6),
                                     'hdr': bool(row.HeadingFormat), 'n': row.Cells.Count,
                                     'bold': rr.Font.Bold == -1,
                                     't': rr.Text.replace('\r\x07', ' | ').replace('\x07', '').strip()[:60]})
                        if rows[-1]['n'] >= 4:        # a wide row: which cells hold text (side-by-side parts)
                            rows[-1]['cells'] = [bool(c.Range.Text.replace('\r', '').replace('\x07', '').strip())
                                                 for c in row.Cells]
                        cut = _w_exact_row_cut(doc, row)
                        if cut:
                            rows[-1]['cut'] = cut
                    # the row that ENDS a page fragment: where its tallest cell's last line sits
                    # (a wrapped last row is taller than the table's usual row step, so the
                    # fragment height is not under-counted - a table of 40 % of a page read as
                    # "small" was a false small_table_split)
                    for i, rw in enumerate(rows):
                        if i + 1 < len(rows) and rows[i + 1]['page'] == rw['page']:
                            continue
                        try:
                            last = []
                            for c in objs[i].Cells:
                                e = max(c.Range.Start, c.Range.End - 1)
                                last.append(doc.Range(e, e).Information(6))
                            rw['y_last'] = max(last)
                        except Exception:
                            pass
                except Exception:              # merged cells: rows not addressable one by one
                    at = doc.Range(r.Start, r.Start)
                    rows = [{'page': at.Information(3), 'y': at.Information(6), 'hdr': False, 'n': 1,
                             'bold': False, 't': '(table)'}]
                items.append({'k': 'table', 'start': r.Start, 'end': r.End, 'rows': rows,
                              'page': rows[0]['page'], 'page_end': r.Information(3),
                              'y': rows[0]['y'], 't': rows[0]['t'],
                              'ncols': max((x['n'] for x in rows), default=1)})
            spans.sort()
            import bisect
            starts = [s for s, _ in spans]

            def in_table(pos):
                k = bisect.bisect_right(starts, pos) - 1
                return k >= 0 and spans[k][0] <= pos < spans[k][1]
            for p in doc.Paragraphs:
                r = p.Range
                if in_table(r.Start):
                    continue
                raw = r.Text
                txt = raw.strip().replace('\r', ' ').replace('\x07', '').replace('\x0c', '').strip()
                shape_h = 0.0
                try:
                    if r.InlineShapes.Count:
                        shape_h = max(r.InlineShapes(k).Height for k in range(1, r.InlineShapes.Count + 1))
                    elif txt in ('', '/') and r.ShapeRange.Count:     # an anchored (floating) picture
                        shape_h = max(r.ShapeRange(k).Height for k in range(1, r.ShapeRange.Count + 1))
                except Exception:
                    shape_h = 0.0
                brk = '\x0c' in raw
                if not txt and not shape_h and not brk:
                    continue
                if not shape_h and not brk and (r.Font.Hidden == -1 or 0 < (r.Font.Size or 10) <= 2):
                    continue                      # hidden / 1pt marker text is not content
                at = doc.Range(r.Start, r.Start)
                last = doc.Range(max(r.End - 1, r.Start), max(r.End - 1, r.Start))
                try:
                    style = p.Style.NameLocal
                except Exception:
                    style = ''
                items.append({'k': 'pic' if (shape_h and not txt.strip('/ ')) else 'p', 'start': r.Start,
                              'end': r.End, 't': txt[:90], 'page': at.Information(3),
                              'page_end': r.Information(3), 'y': at.Information(6),
                              'y_last': last.Information(6), 'style': style, 'ol': p.OutlineLevel,
                              'bold': r.Font.Bold == -1, 'size': r.Font.Size, 'kwn': bool(p.KeepWithNext),
                              'pbb': bool(p.PageBreakBefore), 'brk': brk, 'ls': p.LineSpacing,
                              'lsr': p.LineSpacingRule, 'sa': p.SpaceAfter, 'shape_h': shape_h,
                              'len': len(txt)})
            items.sort(key=lambda x: x['start'])
            res['items'] = items
            res['pages'] = doc.Content.Information(3)
        finally:
            doc.Close(False)
    finally:
        try:
            word.Quit(False)
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return res


EXACT_ROW_LINE = 1.15      # a text line is ~1.15 x its font size (Times / Calibri single spacing)
EXACT_ROW_TOL = 1.0        # pt of slack before an EXACT-height row counts as clipping its text


def _w_exact_row_cut(doc, row):
    """A row whose height is EXACT (``wdRowHeightExactly``) clips whatever does not fit: Word
    still lays the lines out, so the first and last character of each cell give the height its
    text needs. Returns ``{'cell', 'need', 'have', 't'}`` for the worst clipped cell, else None
    (SG Word §15.2: a 4-line crew cell in a 40 pt row lost '+15 more', NARRFIX)."""
    try:
        if row.HeightRule != 2:              # wdRowHeightExactly
            return None
        have = float(row.Height)
        worst = None
        for ci, c in enumerate(row.Cells):
            r = c.Range
            txt = r.Text.replace('\r\x07', '').replace('\x07', '').strip()
            if not txt:
                continue
            s, e = r.Start, max(r.Start, r.End - 2)
            y0 = doc.Range(s, s).Information(6)
            y1 = doc.Range(e, e).Information(6)
            fs = doc.Range(e, e + 1).Font.Size or 10
            if fs > 200:                     # wdUndefined (mixed sizes): take the cell's first
                fs = doc.Range(s, s + 1).Font.Size or 10
            need = (y1 - y0) + fs * EXACT_ROW_LINE
            if need > have + EXACT_ROW_TOL and (worst is None or need - have > worst['need'] - worst['have']):
                worst = {'cell': ci + 1, 'need': round(need, 1), 'have': round(have, 1),
                         't': ' '.join(txt.split())[:60]}
        return worst
    except Exception:
        return None


def _w_heading_table(it):
    """A heading set as a one-row table (the Reporting Studio's numbered item title: a number
    badge + the title, both bold) — a heading, not a data table (STUDIO-DOC-1)."""
    rows = it.get('rows') or []
    t = ' '.join((it.get('t') or '').replace('|', ' ').split())
    return (it['k'] == 'table' and len(rows) == 1 and (it.get('ncols') or 1) <= 2
            and bool(rows[0].get('bold')) and 0 < len(t) <= HEAD_MAX_CHARS
            and bool(re.search(r'[A-Za-z]', t)))


def _w_is_heading(it):
    if it['k'] == 'table':
        return _w_heading_table(it)
    if it['k'] != 'p':
        return False
    t = it.get('t') or ''
    if not t or len(t) > HEAD_MAX_CHARS or not re.search(r'[A-Za-z]', t):
        return False
    st = (it.get('style') or '').lower()
    if st.startswith(('heading', 'title', 'subtitle')) or (it.get('ol') or 10) < 10:
        return True
    return bool(it.get('bold')) and (it.get('len') or len(t)) <= 120


def _w_kind(it):
    if it['k'] == 'table':
        if _w_heading_table(it):
            return 'heading'
        return 'kpi' if (len(it['rows']) <= 2 and (it.get('ncols') or 1) >= 3) else 'table'
    if it['k'] == 'pic':
        return 'picture'
    return 'heading' if _w_is_heading(it) else 'text'


def _w_end_y(it):
    """Where an item ends on its last page (points from the top of the page)."""
    if it['k'] == 'table':
        rows = [r for r in it['rows'] if r['page'] == it['page_end'] and r['y'] is not None and r['y'] >= 0]
        if not rows:
            return None
        ys = sorted(r['y'] for r in rows)
        step = min((b - a for a, b in zip(ys, ys[1:]) if b - a > 1), default=16.0)
        return ys[-1] + step
    if it['k'] == 'pic':
        return (it.get('y_last') or it['y']) + (it.get('shape_h') or 0)
    size = it.get('size') or 10
    if size > 200:                         # wdUndefined (mixed sizes)
        size = 10
    ls, lsr = it.get('ls') or 0, it.get('lsr')
    line = ls if lsr in (3, 4) and ls > size else size * 1.25
    return (it.get('y_last') or it['y']) + line + max(it.get('sa') or 0, 0)


def _w_page_ends(it, page_bottom):
    """``{page: y}`` — where the item's content ends on every page it occupies."""
    if it['k'] == 'table':
        by = {}
        for r in it['rows']:
            if r['page'] is not None and r['y'] is not None and r['y'] >= 0:
                by.setdefault(r['page'], []).append(r['y'])
        out = {}
        for p, ys in by.items():
            ys.sort()
            step = min((b - a for a, b in zip(ys, ys[1:]) if b - a > 1), default=16.0)
            out[p] = ys[-1] + step
        return out
    out = {p: page_bottom for p in range(it['page'], it['page_end'])}
    out[it['page_end']] = _w_end_y(it)
    return out


def _w_side_parts(rows):
    """Column ranges of the side-by-side parts of a Word table: two (or more) tables built as
    ONE table with a blank column between them (the Narrative's code-table pairs). Read from
    its last wide header row - runs of cells that hold text, split by empty cells."""
    hdr = [r for r in rows if r.get('hdr') and len(r.get('cells') or []) >= 4]
    if not hdr:
        return []
    parts, lo = [], None
    for i, full in enumerate(list(hdr[-1]['cells']) + [False]):
        if full and lo is None:
            lo = i
        elif not full and lo is not None:
            parts.append((lo, i))
            lo = None
    return parts if len(parts) >= 2 else []


def _w_side_part_flags(it, by, what, hdr_rep):
    """A side-by-side part of a table that runs over pages: its header repeated over no rows
    of it (the part ended on an earlier page), or the part split leaving 1-2 of its rows."""
    out = []
    for lo, hi in _w_side_parts(it['rows']):
        per = collections.OrderedDict(
            (p, sum(1 for r in rows if not r['hdr'] and any((r.get('cells') or [])[lo:hi])))
            for p, rows in by.items())
        on = [p for p, c in per.items() if c]
        if not on:
            continue
        name = f'columns {lo + 1}-{hi}'
        empty = [p for p, c in per.items() if not c and p > on[0]]
        if empty and hdr_rep:
            out.append(_flag('empty_repeated_header', empty[0],
                             f'{what}: its side-by-side part ({name}) repeats its header on page '
                             f'{empty[0]} over no rows (the part ends on page {on[-1]})'))
        if len(on) > 1 and min(per[p] for p in on) < MIN_ROWS:
            split = ', '.join(f'p{p}: {per[p]} row(s)' for p in on)
            out.append(_flag('table_split_few_rows', on[0],
                             f'{what}: its side-by-side part ({name}) is split {split} '
                             f'(the rules keep >= {MIN_ROWS})'))
    return out


def analyze_word_layout(layout):
    """Pagination defects from Word's own layout (``word_layout()`` output)."""
    pg = layout['page']
    top, bot = pg['top'], pg['h'] - pg['bottom']
    area = max(bot - top, 1.0)
    items = layout['items']
    npages = layout.get('pages') or max((it['page_end'] for it in items), default=0)
    flags, info = [], []

    def section_break_before(j):
        """True when a page break by design sits right before items[j]."""
        it = items[j]
        if it.get('pbb') or (it.get('brk') and not (it.get('t') or '').strip()):
            return True                   # page-break-before, or a paragraph holding only a break
        return j > 0 and bool(items[j - 1].get('brk'))

    for j, it in enumerate(items):
        k = _w_kind(it)
        nxt = items[j + 1] if j + 1 < len(items) else None
        if k == 'heading' and nxt is not None and nxt['page'] > it['page_end'] and not it.get('brk') \
                and not section_break_before(j + 1):
            kn = _w_kind(nxt)
            # (a heading TABLE — the Studio's numbered item title — opens a new item: it outranks
            # any paragraph heading before it)
            if not (kn == 'heading' and (nxt.get('size') or (99 if nxt['k'] == 'table' else 0))
                    >= (it.get('size') or 0)):
                kind = 'kpi_separated_from_heading' if kn == 'kpi' else 'orphaned_heading'
                flags.append(_flag(kind, it['page_end'], f"heading {it['t'][:60]!r} ends page {it['page_end']}; "
                                                         f"its {kn} starts page {nxt['page']}"))
        if k == 'heading' and nxt is not None and nxt['page'] == it['page_end'] and _w_kind(nxt) == 'text':
            n2 = items[j + 2] if j + 2 < len(items) else None
            if (n2 is not None and _w_kind(n2) in ('table', 'picture', 'kpi') and n2['page'] > nxt['page_end']
                    and not section_break_before(j + 2) and (nxt.get('len') or 0) <= 320):
                kn = _w_kind(n2)
                kind = 'kpi_separated_from_heading' if kn == 'kpi' else 'heading_separated_from_block'
                flags.append(_flag(kind, it['page_end'], f"heading {it['t'][:60]!r} + its intro end page "
                                                         f"{it['page_end']}; its {kn} starts page {n2['page']}"))
        if k == 'picture' and j > 0:
            prv = items[j - 1]
            if (prv['page_end'] < it['page'] and _w_kind(prv) == 'text' and len(prv.get('t') or '') <= 160
                    and not section_break_before(j)):
                flags.append(_flag('picture_separated_from_caption', it['page'],
                                   f"picture on page {it['page']} but its label / caption "
                                   f"{prv['t'][:50]!r} is on page {prv['page_end']}"))
        if k == 'picture' and (it.get('shape_h') or 0) > 0:
            end = (it.get('y_last') or it['y']) + it['shape_h']
            if it['shape_h'] > area + 2 or end > pg['h'] - 2:
                flags.append(_flag('graphic_cut', it['page'], f"picture of {it['shape_h']:.0f}pt does not fit "
                                                              f"the page ({area:.0f}pt) - it is clipped"))
        if it['k'] == 'table' and it['page_end'] > it['page'] and (it.get('ncols') or 1) >= 2:
            by = collections.OrderedDict()
            for r in it['rows']:
                if r['y'] is None or r['page'] is None:
                    continue
                by.setdefault(r['page'], []).append(r)
            hdr_rep = any(r['hdr'] for r in it['rows'])
            nh = 0
            while nh < len(it['rows']) and (it['rows'][nh]['hdr'] or (nh == 0 and it['rows'][0].get('bold')
                                                                         and len(it['rows']) > 1
                                                                         and not it['rows'][1].get('bold'))):
                nh += 1
                if not it['rows'][nh - 1]['hdr']:
                    break
            frags = []
            tail = {}              # per fragment: its last row's extra lines (a wrapped last row)
            first = True
            for p, rows in by.items():
                tail[p] = max(0.0, (rows[-1].get('y_last') or rows[-1]['y']) - rows[-1]['y'])
                body = len([r for r in rows if not r['hdr']]) - (nh if (first and not hdr_rep) else 0)
                ys = [r['y'] for r in rows]
                frags.append((p, max(body, 0), ys))
                first = False
            if len(frags) < 2:
                continue
            steps = [b - a for _, _, ys in frags for a, b in zip(ys, ys[1:]) if b - a > 1]
            step = sorted(steps)[len(steps) // 2] if steps else 16.0
            height = sum((ys[-1] - ys[0]) + tail.get(p, 0.0) + step for p, _, ys in frags)
            desc = ', '.join(f'p{p}: {b} body row(s)' for p, b, _ in frags)
            what = f"table {it['t'][:40]!r} ({len(it['rows'])} rows)"
            flags.extend(_w_side_part_flags(it, by, what, hdr_rep))
            small = [f for f in frags if f[1] < MIN_ROWS]
            if small:
                flags.append(_flag('table_split_few_rows', small[0][0], f'{what} split {desc} '
                                                                        f'(the rules keep >= {MIN_ROWS})'))
            elif height <= FIT * area:
                flags.append(_flag('small_table_split', frags[0][0], f'{what} of {height:.0f}pt '
                                                                      f'({height / area:.0%} of a page) fits on '
                                                                      f'one page but is split: {desc}'))
            if nh and not hdr_rep:
                flags.append(_flag('table_header_not_repeated', frags[1][0],
                                   f'{what} continues on page {frags[1][0]} without its header row'))
    # large blank at a page end followed by a pushed block; empty pages
    last_on, first_on, covered = {}, {}, set()
    for j, it in enumerate(items):
        # a text-less paragraph holding only a page break is not content: when it spills onto
        # a page of its own (the section before filled its page) that page is EMPTY - it was
        # counted as covered and read only as a 'section_break_blank' (SG Word p12, NARRFIX)
        if not (it['k'] == 'p' and it.get('brk') and not (it.get('t') or '').strip()
                and not it.get('shape_h')):
            covered.update(range(it['page'], it['page_end'] + 1))
        # an item that runs over several pages ends on EACH of them (a table: its last row
        # there; text: the page bottom) and is the first thing on each page it continues onto
        # — counting only its last page read a full page it filled as "ends early"
        for p, e in _w_page_ends(it, bot).items():
            if e is not None and (p not in last_on or e >= last_on[p][0]):
                last_on[p] = (e, it, j)
        for p in range(it['page'], it['page_end'] + 1):
            if p not in first_on:
                first_on[p] = (it, j)
    for p in range(1, (npages or 0) + 1):
        if p not in covered and npages > 1:
            flags.append(_flag('empty_page', p, 'the page has nothing on it'))
    for it in items:                          # text clipped inside an EXACT-height table row
        if it.get('k') != 'table':
            continue
        for rw in it.get('rows') or []:
            cut = rw.get('cut')
            if cut:
                flags.append(_flag('text_cut', rw.get('page') or it['page'],
                                   f"a table cell's text is cut off by its exact row height: "
                                   f"cell {cut['cell']} {cut['t']!r} needs ~{cut['need']:.0f} pt, "
                                   f"the row is {cut['have']:.0f} pt"))
    for p in sorted(last_on):
        e, it, j = last_on[p]
        if p + 1 not in first_on:
            continue
        fill = (e - top) / area
        if fill >= 1 - BLANK:
            continue
        nf, jn = first_on[p + 1]
        kn = _w_kind(nf)
        what = f"page {p} ends at ~{max(fill, 0):.0%} ({1 - max(fill, 0):.0%} blank); page {p + 1} starts with {kn} " \
               f"{(nf.get('t') or '')[:40]!r}"
        if section_break_before(jn) or it.get('brk'):
            info.append(_flag('section_break_blank', p, what + ' (a page break by design)'))
        else:
            flags.append(_flag('large_blank_then_continuation', p, what))
    return flags, info, npages


def _word_layout_subprocess(path, timeout):
    """Run the COM layout in its own process (a hung Word never hangs the caller)."""
    fd, out = tempfile.mkstemp(suffix='.json', prefix='cx_wl_')
    os.close(fd)
    pid_file = out + '.pid'
    try:
        try:
            cp = subprocess.run([sys.executable, '-m', 'p6_export.pagination_check', '--word-layout', path, out,
                                 pid_file],
                                timeout=timeout, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
                                cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        except subprocess.TimeoutExpired:
            # the child is gone but the Word it started keeps laying the document out (CPU +
            # RAM for hours): stop THAT Word only (its pid was recorded by the child) - never
            # a Word the user has open
            _stop_pid_in(pid_file)
            raise
        if cp.returncode:
            tail = (cp.stderr or b'').decode('utf-8', 'replace').strip().splitlines()[-1:] or ['']
            raise RuntimeError(f'Word layout exited {cp.returncode}: {tail[0][:200]}')
        with open(out, encoding='utf-8') as fh:
            return json.load(fh)
    finally:
        for f in (out, pid_file):
            try:
                os.remove(f)
            except OSError:
                pass


def _stop_pid_in(pid_file):
    """Force-stop the process whose id ``pid_file`` holds (nothing when it is missing)."""
    try:
        with open(pid_file, encoding='ascii') as fh:
            pid = int(fh.read().strip())
    except (OSError, ValueError):
        return False
    try:
        if os.name == 'nt':
            subprocess.run(['taskkill', '/PID', str(pid), '/F'], stdout=subprocess.DEVNULL,
                           stderr=subprocess.DEVNULL, timeout=60)
        else:
            os.kill(pid, 9)
        return True
    except Exception:
        return False


def _spire_to_pdf(path):
    from spire.doc import Document, FileFormat
    fd, out = tempfile.mkstemp(suffix='.pdf', prefix='cx_spire_')
    os.close(fd)
    d = Document()
    d.LoadFromFile(os.path.abspath(path))
    d.SaveToFile(out, FileFormat.PDF)
    d.Close()
    return out


def check_docx(path, engine='auto', timeout=1800):
    """Check a Word file (.docx or .doc). ``engine``: 'auto' (Word, then Spire.Doc), 'word',
    'spire'. Returns status 'skipped' with a clear message when no renderer is available."""
    path = os.path.abspath(path)
    fmt = os.path.splitext(path)[1].lstrip('.').lower() or 'docx'
    tried = []
    if engine in ('auto', 'word'):
        if word_available():
            try:
                if getattr(sys, 'frozen', False):
                    layout = word_layout(path)
                else:
                    layout = _word_layout_subprocess(path, timeout)
                flags, info, npages = analyze_word_layout(layout)
                if fmt == 'docx':             # the pictures themselves: cut / mostly blank
                    flags += docx_picture_flags(path, [it['page'] for it in layout.get('items') or []
                                                       if it.get('k') == 'pic'])
                return _result(path, fmt, 'word-com', npages, flags, info,
                               message="Word's own pagination (COM layout)")
            except Exception as exc:          # fall through to Spire
                tried.append(f'Word COM failed: {str(exc)[:200]}')
        else:
            tried.append('Microsoft Word (COM via pywin32) is not available')
    if engine in ('auto', 'spire'):
        if spire_available():
            pdf = None
            try:
                pdf = _spire_to_pdf(path)
                res = check_pdf(pdf, headings=docx_headings(path) if fmt == 'docx' else None)
                if fmt == 'docx':             # the pictures themselves: cut / mostly blank
                    extra = docx_picture_flags(path)
                    if extra:
                        res = _result(path, fmt, 'spire', res['pages'], res['flags'] + extra,
                                      res['info'], message=res.get('message', ''))
                res.update(file=path, format=fmt, engine='spire')
                res['message'] = ('rendered with Spire.Doc (its free edition converts only the first 10 '
                                  'pages; Word may paginate slightly differently)'
                                  + ('; ' + '; '.join(tried) if tried else ''))
                return res
            except Exception as exc:
                tried.append(f'Spire.Doc failed: {str(exc)[:200]}')
            finally:
                if pdf:
                    try:
                        os.remove(pdf)
                    except OSError:
                        pass
        else:
            tried.append('Spire.Doc (spire.doc) is not installed')
    return _result(path, fmt, 'none', 0, [], status='skipped',
                   message='Word file NOT checked - no local renderer: ' + '; '.join(tried)
                           + '. Install Microsoft Word + pywin32, or spire.doc, or check the PDF.')


def check_file(path, engine='auto', headings=None, timeout=1800):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.pdf':
        return check_pdf(path, headings=headings)
    if ext in ('.docx', '.doc'):
        return check_docx(path, engine=engine, timeout=timeout)
    raise ValueError(f'unsupported file type {ext!r} (expected .pdf, .docx or .doc)')


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv[:1] == ['--word-layout']:                    # internal: COM layout in a child process
        res = word_layout(argv[1], pid_file=argv[3] if len(argv) > 3 else None)
        with open(argv[2], 'w', encoding='utf-8') as fh:
            json.dump(res, fh, ensure_ascii=False)
        return EXIT_CLEAN
    ap = argparse.ArgumentParser(prog='python -m p6_export.pagination_check',
                                 description='Find page-composition defects in a report PDF / Word file.')
    ap.add_argument('file', help='report .pdf, .docx or .doc')
    ap.add_argument('--engine', choices=('auto', 'word', 'spire'), default='auto',
                    help='Word renderer (default: Word COM, then Spire.Doc)')
    ap.add_argument('--html', help='the report source HTML - its heading texts help find headings')
    ap.add_argument('--timeout', type=int, default=1800, help='seconds allowed for Word to lay out')
    ap.add_argument('--indent', type=int, default=1)
    a = ap.parse_args(argv)
    if not os.path.isfile(a.file):
        print(json.dumps({'file': a.file, 'status': 'error', 'message': 'file not found'}))
        return EXIT_ERROR
    hints = None
    if a.html:
        with open(a.html, encoding='utf-8', errors='replace') as fh:
            hints = html_headings(fh.read())
    try:
        res = check_file(a.file, engine=a.engine, headings=hints, timeout=a.timeout)
    except Exception as exc:
        print(json.dumps({'file': a.file, 'status': 'error', 'message': str(exc)[:500]}))
        return EXIT_ERROR
    sys.stdout.write(json.dumps(res, ensure_ascii=False, indent=a.indent) + '\n')
    if res['status'] == 'skipped':
        return EXIT_SKIPPED
    return EXIT_FLAGS if res['flag_count'] else EXIT_CLEAN


if __name__ == '__main__':
    sys.exit(main())
