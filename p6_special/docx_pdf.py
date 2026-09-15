"""Word ``.docx`` that is a pixel-exact copy of the Special Report PDF.

Ibrahim's standing requirement is that the Word export match the PDF *exactly*
— same page layout, same tables, same styling, same page numbers. A native
Word rebuild (see :mod:`p6_special.docx_report`) can only ever approximate the
PDF: Word's layout engine paginates differently, styles tables its own way, and
its contents-page numbers depend on a manual field refresh (which an unactivated
Word won't even run). The only way to guarantee "exactly the same" is to render
the *same HTML the PDF uses*, print it to a PDF with the *same Chrome flags*, and
drop each PDF page into the Word document as a full-page image.

The result is a ``.docx`` that is, page for page, the PDF — so page coordination,
the contents-page numbers, and every table's row/column heights, format, signs and
colours are identical by construction. The trade-off (chosen by Ibrahim, 2026-09):
the tables/text are images and so not editable inside Word.

Requires Chrome (to print the PDF) and PyMuPDF (to rasterise it). The caller
(:func:`p6_special.assemble.docx`) falls back to the native builder if either is
unavailable, so the export never hard-fails.
"""
import io
import os
import subprocess
import tempfile


# A4 at 210 × 297 mm. Rasterise at ~200 DPI — crisp text when printed/zoomed, and
# python-docx scales the image to the page from its pixel aspect ratio (A4), so we
# only ever set the width and the height follows exactly.
_A4_W_MM = 210.0
_A4_H_MM = 297.0
_DPI = 200.0


def _html_to_pdf(html, chrome, pdf_path):
    """Print ``html`` to ``pdf_path`` with the SAME Chrome invocation the PDF export
    uses (``--headless --print-to-pdf --no-pdf-header-footer``) so the bytes match."""
    html_path = None
    try:
        fd, html_path = tempfile.mkstemp(suffix='.html')
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write(html)
        subprocess.run(
            [chrome, '--headless', '--disable-gpu', '--no-sandbox',
             f'--print-to-pdf={pdf_path}', '--no-pdf-header-footer',
             f'file:///{html_path.replace(os.sep, "/")}'],
            check=True, capture_output=True, timeout=180)
    finally:
        if html_path and os.path.exists(html_path):
            try:
                os.remove(html_path)
            except OSError:
                pass


def build_docx_from_pdf(path, html, chrome):
    """Render ``html`` to a PDF via ``chrome`` and write ``path`` as a ``.docx`` whose
    every page is a full-bleed image of the matching PDF page (A4, zero margins).

    Raises on any failure (missing chrome, missing PyMuPDF, Chrome/print error) so the
    caller can fall back to the native Word builder. Returns ``str(path)`` on success.
    """
    if not chrome:
        raise RuntimeError('chrome executable required for the PDF-exact Word export')

    try:                       # PyMuPDF — raises ImportError if not bundled; caller falls back
        import pymupdf as fitz  # 1.24+ package name
    except ImportError:
        import fitz             # legacy import name
    from docx import Document
    from docx.shared import Mm, Pt

    pdf_path = None
    try:
        fd, pdf_path = tempfile.mkstemp(suffix='.pdf')
        os.close(fd)
        _html_to_pdf(html, chrome, pdf_path)

        pdf = fitz.open(pdf_path)
        try:
            if pdf.page_count == 0:
                raise RuntimeError('Chrome produced an empty PDF')

            document = Document()
            section = document.sections[0]
            # A4, no margins — the image is the page.
            section.page_width = Mm(_A4_W_MM)
            section.page_height = Mm(_A4_H_MM)
            section.left_margin = section.right_margin = Mm(0)
            section.top_margin = section.bottom_margin = Mm(0)
            section.header_distance = Mm(0)
            section.footer_distance = Mm(0)

            zoom = _DPI / 72.0
            matrix = fitz.Matrix(zoom, zoom)
            # Full page width, minus a 2mm safety margin so the image (A4 aspect) plus the
            # paragraph's own leading never tips onto a blank following page. The PDF's
            # frame is inset from the edge, so this sliver of white margin is invisible.
            img_w = Mm(_A4_W_MM - 2.0)

            for i, page in enumerate(pdf):
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                png = pix.tobytes('png')
                p = document.add_paragraph()
                pf = p.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                # Single/auto line spacing (float, lineRule="auto") — the line box GROWS to
                # contain the tall inline image. NEVER use an exact Length (e.g. Pt(1)) here:
                # Word honours "Exactly" literally and clips the image to that height, so the
                # page comes out blank. (Caught in review 2026-09-15.)
                pf.line_spacing = 1.0
                if i > 0:
                    pf.page_break_before = True   # each page image on its own page
                p.add_run().add_picture(io.BytesIO(png), width=img_w)

            document.save(str(path))
        finally:
            pdf.close()
    finally:
        if pdf_path and os.path.exists(pdf_path):
            try:
                os.remove(pdf_path)
            except OSError:
                pass
    return str(path)
